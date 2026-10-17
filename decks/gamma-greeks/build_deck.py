#!/usr/bin/env python3
"""Build the Gamma & the Greeks briefing (16:9, dark).

Writes gamma_greeks.pptx next to this file. With --render, also exports
previews/slide_XX.png via LibreOffice and pdftoppm.

The SPY levels are the 16 Oct 2026 close supplied for this briefing:
flip 761.19, call wall 764, put wall 763, spot about 763.25.
The strike chart is an EXAMPLE with made-up numbers.
"""

from __future__ import annotations

import argparse
import math
import os
import shutil
import subprocess
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from lxml import etree
from matplotlib import font_manager
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
from matplotlib.ticker import FuncFormatter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
BUILD = ROOT / "build"
CHARTS = BUILD / "charts"
PREVIEWS = ROOT / "previews"

# --- palette ---------------------------------------------------------------
# Hex without '#'. Matplotlib helpers add it.
BG = "0B1020"
CARD = "141B30"
CARD_2 = "1C2742"
LINE = "2C3B5A"
TEXT = "F3F6FB"
MUTED = "B7C3D6"
DIM = "93A0B8"
GOLD = "E4B15A"
GREEN = "3DDC97"
GREEN_DIM = "1C8F66"
RED = "FF6B81"
RED_DIM = "C2566A"
CYAN = "6EC8FF"
PURPLE = "C4B5FD"
WHITE = "FFFFFF"
POS_WASH = "10281E"
NEG_WASH = "2C1822"
PIN_WASH = "3A3018"

FONT = "Inter"
DPI = 200

# Black–Scholes inputs shared by every Greek chart. Illustrative, not a quote.
R = 0.045
Q = 0.013
IV = 0.20
S0 = 100.0

# Real SPY snapshot, 16 Oct 2026 close.
SPY_SPOT = 763.25
SPY_FLIP = 761.19
SPY_CALL = 764.0
SPY_PUT = 763.0

SW = 13.3333333333
SH = 7.5

for _font in Path("/usr/share/fonts/truetype/macos").glob("Inter*.ttf"):
    font_manager.fontManager.addfont(str(_font))

plt.rcParams.update(
    {
        "font.family": FONT,
        "font.size": 12,
        "text.color": f"#{TEXT}",
        "axes.labelcolor": f"#{MUTED}",
        "xtick.color": f"#{MUTED}",
        "ytick.color": f"#{MUTED}",
        "axes.edgecolor": f"#{LINE}",
        "figure.facecolor": f"#{CARD}",
        "savefig.facecolor": f"#{CARD}",
        "axes.facecolor": f"#{CARD}",
        "axes.unicode_minus": False,
    }
)


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.replace("#", ""))


def hx(value: str) -> str:
    return value if value.startswith("#") else f"#{value}"


# --- Black–Scholes ---------------------------------------------------------
_erf = np.vectorize(math.erf, otypes=[float])


def _norm_cdf(x: np.ndarray) -> np.ndarray:
    return 0.5 * (1.0 + _erf(np.asarray(x, dtype=float) / math.sqrt(2.0)))


def _norm_pdf(x: np.ndarray) -> np.ndarray:
    x = np.asarray(x, dtype=float)
    return np.exp(-0.5 * x * x) / math.sqrt(2.0 * math.pi)


def _d1_d2(S, K, T, sig, r=R, q=Q):
    S = np.asarray(S, dtype=float)
    K = np.asarray(K, dtype=float)
    T = np.asarray(T, dtype=float)
    sig = np.asarray(sig, dtype=float)
    sqrt_t = np.sqrt(T)
    d1 = (np.log(S / K) + (r - q + 0.5 * sig * sig) * T) / (sig * sqrt_t)
    d2 = d1 - sig * sqrt_t
    return d1, d2


def call_price(S, K, T, sig, r=R, q=Q) -> np.ndarray:
    d1, d2 = _d1_d2(S, K, T, sig, r, q)
    T = np.asarray(T, dtype=float)
    return np.exp(-q * T) * S * _norm_cdf(d1) - np.exp(-r * T) * K * _norm_cdf(d2)


def put_price(S, K, T, sig, r=R, q=Q) -> np.ndarray:
    d1, d2 = _d1_d2(S, K, T, sig, r, q)
    T = np.asarray(T, dtype=float)
    return np.exp(-r * T) * K * _norm_cdf(-d2) - np.exp(-q * T) * S * _norm_cdf(-d1)


def call_delta(S, K, T, sig, r=R, q=Q) -> np.ndarray:
    d1, _ = _d1_d2(S, K, T, sig, r, q)
    T = np.asarray(T, dtype=float)
    return np.exp(-q * T) * _norm_cdf(d1)


def call_gamma(S, K, T, sig, r=R, q=Q) -> np.ndarray:
    d1, _ = _d1_d2(S, K, T, sig, r, q)
    T = np.asarray(T, dtype=float)
    S = np.asarray(S, dtype=float)
    sig = np.asarray(sig, dtype=float)
    return np.exp(-q * T) * _norm_pdf(d1) / (S * sig * np.sqrt(T))


def call_theta_daily(S, K, T, sig, r=R, q=Q) -> np.ndarray:
    """Calendar-day theta of a long call. Negative when the option decays."""
    d1, d2 = _d1_d2(S, K, T, sig, r, q)
    T = np.asarray(T, dtype=float)
    S = np.asarray(S, dtype=float)
    K = np.asarray(K, dtype=float)
    sig = np.asarray(sig, dtype=float)
    sqrt_t = np.sqrt(T)
    yearly = (
        -np.exp(-q * T) * S * _norm_pdf(d1) * sig / (2.0 * sqrt_t)
        - r * K * np.exp(-r * T) * _norm_cdf(d2)
        + q * S * np.exp(-q * T) * _norm_cdf(d1)
    )
    return yearly / 365.0


def self_check() -> None:
    T = 30 / 365
    c = float(call_price(S0, S0, T, IV))
    p = float(put_price(S0, S0, T, IV))
    fwd = math.exp(-Q * T) * S0 - math.exp(-R * T) * S0
    if abs((c - p) - fwd) > 1e-8:
        raise SystemExit(f"put-call parity failed: {c - p} vs {fwd}")
    delta = float(call_delta(S0, S0, T, IV))
    if not 0.50 < delta < 0.60:
        raise SystemExit(f"ATM delta out of range: {delta}")
    if float(call_gamma(S0, S0, T, IV)) <= 0:
        raise SystemExit("gamma should be positive")
    if float(call_theta_daily(S0, S0, T, IV)) >= 0:
        raise SystemExit("ATM call theta should be negative")
    g7 = float(call_gamma(S0, S0, 7 / 365, IV))
    g90 = float(call_gamma(S0, S0, 90 / 365, IV))
    if g7 <= g90:
        raise SystemExit("near-dated ATM gamma should exceed 90-day gamma")
    # Decay should be larger (more negative theta, larger cost) at 7 days than 45.
    d7 = -float(call_theta_daily(S0, S0, 7 / 365, IV))
    d45 = -float(call_theta_daily(S0, S0, 45 / 365, IV))
    if d7 <= d45:
        raise SystemExit("decay should speed up as expiration approaches")
    print(
        f"BS check ok  ATM 30d call ${c:.2f}  delta {delta:.3f}  "
        f"gamma7 {g7:.3f}  gamma90 {g90:.3f}  decay7 ${d7:.3f}/sh"
    )


# --- charts ----------------------------------------------------------------
def _save(fig, name: str) -> Path:
    CHARTS.mkdir(parents=True, exist_ok=True)
    path = CHARTS / name
    fig.savefig(path, dpi=DPI, facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def _style_ax(ax, xlabel: str | None = None, ylabel: str | None = None) -> None:
    ax.set_facecolor(hx(CARD))
    ax.set_axisbelow(True)
    ax.grid(axis="y", color="#243049", lw=0.8, zorder=0)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color(hx(LINE))
    ax.spines["bottom"].set_color(hx(LINE))
    ax.tick_params(colors=hx(MUTED), labelsize=12, width=0.6, length=3, pad=3)
    if xlabel:
        ax.set_xlabel(xlabel, color=hx(MUTED), fontsize=13, labelpad=6)
    if ylabel:
        ax.set_ylabel(ylabel, color=hx(MUTED), fontsize=13, labelpad=8)


def _box():
    return dict(boxstyle="round,pad=0.32", fc=hx(CARD), ec=hx(LINE), lw=0.7)


def chart_positive(w: float, h: float) -> Path:
    """Range-bound path: dealers sell peaks and buy troughs."""
    fig = plt.figure(figsize=(w, h), dpi=DPI)
    fig.patch.set_facecolor(hx(CARD))
    ax = fig.add_axes([0.12, 0.14, 0.84, 0.74])
    _style_ax(ax, xlabel="time", ylabel="price")
    xs = np.linspace(0, 4 * math.pi, 700)
    ys = 100 + 1.15 * np.sin(xs)
    ax.fill_between(xs, 98.75, 101.25, color=hx(GOLD), alpha=0.08, zorder=0)
    ax.axhline(100, color=hx(GOLD), lw=1.15, ls=(0, (4, 3)), zorder=1)
    ax.plot(xs, ys, color=hx(CYAN), lw=2.5, zorder=3)
    peaks = [math.pi / 2, 5 * math.pi / 2]
    troughs = [3 * math.pi / 2, 7 * math.pi / 2]
    ax.scatter(peaks, [101.15, 101.15], s=36, color=hx(RED), zorder=4)
    ax.scatter(troughs, [98.85, 98.85], s=36, color=hx(GREEN), zorder=4)
    ax.annotate(
        "sell the rip",
        xy=(peaks[0], 101.15),
        xytext=(peaks[0] + 0.15, 103.15),
        ha="center",
        va="bottom",
        color=hx(RED),
        fontsize=13,
        fontweight="medium",
        arrowprops=dict(arrowstyle="-|>", color=hx(RED), lw=1.6),
        bbox=_box(),
        zorder=5,
    )
    ax.annotate(
        "buy the dip",
        xy=(troughs[0], 98.85),
        xytext=(troughs[0] + 0.2, 96.85),
        ha="center",
        va="top",
        color=hx(GREEN),
        fontsize=13,
        fontweight="medium",
        arrowprops=dict(arrowstyle="-|>", color=hx(GREEN), lw=1.6),
        bbox=_box(),
        zorder=5,
    )
    ax.text(4 * math.pi - 0.2, 100.32, "pin", color=hx(GOLD), ha="right", va="bottom", fontsize=13)
    # Keep "range" off the peak marker. The second crest is near x=7.9.
    ax.text(9.6, 102.35, "range", color=hx(GOLD), ha="left", va="bottom", fontsize=13)
    ax.set_xlim(-0.15, 4 * math.pi + 0.15)
    ax.set_ylim(96.2, 104.3)
    ax.set_xticks([])
    ax.set_yticks([99, 100, 101])
    fig.text(0.12, 0.955, "SCHEMATIC", color=hx(GOLD), fontsize=12, fontweight="bold", va="top")
    fig.text(
        0.98,
        0.955,
        "hedge fights the move",
        color=hx(MUTED),
        fontsize=12,
        ha="right",
        va="top",
    )
    return _save(fig, "positive.png")


def chart_negative(w: float, h: float) -> Path:
    """Two accelerating paths: the hedge points the same way as price."""
    fig = plt.figure(figsize=(w, h), dpi=DPI)
    fig.patch.set_facecolor(hx(CARD))
    x = np.linspace(0, 1, 500)
    bend = np.maximum(x - 0.36, 0.0) ** 1.45
    up_alone = 100 + 6.0 * x
    up_hedge = 100 + 6.0 * x + 9.0 * bend
    dn_alone = 100 - 6.0 * x
    dn_hedge = 100 - 6.0 * x - 9.0 * bend

    def _panel(rect, alone, chased, title, chase_color, label, label_dy):
        ax = fig.add_axes(rect)
        _style_ax(ax, xlabel=None, ylabel="price")
        ax.plot(x, alone, color=hx(DIM), lw=1.6, ls=(0, (4, 3)), zorder=2)
        ax.plot(x, chased, color=hx(chase_color), lw=2.6, zorder=3)
        # Mark where the paths split, and point at the chased path.
        split = int(np.argmax(x >= 0.72))
        ax.annotate(
            label,
            xy=(x[split], chased[split]),
            xytext=(0.18, chased[split] + label_dy),
            ha="left",
            va="center",
            color=hx(chase_color),
            fontsize=13,
            fontweight="medium",
            arrowprops=dict(arrowstyle="-|>", color=hx(chase_color), lw=1.5),
            bbox=_box(),
            zorder=5,
        )
        end = -1
        ax.text(x[end] - 0.02, alone[end], "price alone", color=hx(DIM), ha="right", va="center", fontsize=11)
        ax.set_xlim(-0.02, 1.02)
        pad = 1.6
        lo = min(alone.min(), chased.min()) - pad
        hi = max(alone.max(), chased.max()) + pad
        # Keep the callout inside the axes.
        hi = max(hi, chased[split] + label_dy + 1.3)
        lo = min(lo, chased[split] + label_dy - 1.3)
        ax.set_ylim(lo, hi)
        ax.set_xticks([])
        ax.set_title(title, loc="left", color=hx(TEXT), fontsize=14, pad=6, fontweight="medium")
        return ax

    _panel(
        [0.12, 0.54, 0.84, 0.34],
        up_alone,
        up_hedge,
        "Rally",
        RED,
        "dealers buy the rip",
        2.4,
    )
    _panel(
        [0.12, 0.09, 0.84, 0.36],
        dn_alone,
        dn_hedge,
        "Selloff",
        RED,
        "dealers sell the drop",
        -2.6,
    )
    # Right side only — a left-side tag collides with the Rally title.
    fig.text(
        0.98,
        0.975,
        "SCHEMATIC    ·    hedge and price point the same way",
        color=hx(MUTED),
        fontsize=12,
        ha="right",
        va="top",
    )
    return _save(fig, "negative.png")


def chart_gex(w: float, h: float) -> Path:
    strikes = [96, 97, 98, 99, 100, 101, 102, 103, 104, 105, 106, 107, 108]
    vals = [-8, -16, -44, -22, -10, 8, 16, 28, 36, 62, 26, 14, 6]
    put_i, call_i = 2, 9
    flip_x, spot_x = 4.5, 6.0
    colors = []
    for i, v in enumerate(vals):
        if i == call_i:
            colors.append(hx(GREEN))
        elif i == put_i:
            colors.append(hx(RED))
        elif v > 0:
            colors.append(hx(GREEN_DIM))
        else:
            colors.append(hx(RED_DIM))

    fig = plt.figure(figsize=(w, h), dpi=DPI)
    fig.patch.set_facecolor(hx(CARD))
    ax = fig.add_axes([0.07, 0.16, 0.90, 0.74])
    _style_ax(ax, xlabel="Strike", ylabel="Dealer gamma  ·  example units")
    idx = np.arange(len(strikes))
    ax.bar(idx, vals, color=colors, width=0.72, zorder=3)
    ax.axhline(0, color=hx(GOLD), lw=1.15, zorder=2)
    ax.axvline(flip_x, color=hx(CYAN), lw=1.25, ls=(0, (3, 2.5)), zorder=4)
    ax.axvline(spot_x, color=hx(GOLD), lw=1.25, ls=(0, (5, 3)), zorder=4)
    ax.scatter([spot_x], [0], s=46, marker="D", color=hx(GOLD), zorder=6)
    for i, v in enumerate(vals):
        if i in (put_i, call_i):
            continue
        if v >= 0:
            ax.text(i, v + 1.6, f"{v:+.0f}", ha="center", va="bottom", color=hx(GREEN), fontsize=11)
        else:
            ax.text(i, v - 1.6, f"{v:.0f}", ha="center", va="top", color=hx(RED), fontsize=11)
    ax.annotate(
        "GAMMA FLIP\nsign changes here",
        xy=(flip_x, 0),
        xytext=(0.15, 86),
        ha="left",
        va="top",
        color=hx(CYAN),
        fontsize=12.5,
        fontweight="medium",
        linespacing=1.35,
        arrowprops=dict(arrowstyle="-|>", color=hx(CYAN), lw=1.4),
        bbox=_box(),
        zorder=7,
    )
    ax.annotate(
        "CALL WALL   +62\nresistance or a magnet\na break can squeeze",
        xy=(call_i, 62),
        xytext=(call_i - 0.15, 96),
        ha="center",
        va="top",
        color=hx(GREEN),
        fontsize=12.5,
        fontweight="medium",
        linespacing=1.35,
        arrowprops=dict(arrowstyle="-|>", color=hx(GREEN), lw=1.4),
        bbox=_box(),
        zorder=7,
    )
    ax.annotate(
        "PUT WALL   −44\nsupport\na break speeds selling",
        xy=(put_i, -44),
        xytext=(-0.35, -58),
        ha="left",
        va="center",
        color=hx(RED),
        fontsize=12.5,
        fontweight="medium",
        linespacing=1.35,
        arrowprops=dict(arrowstyle="-|>", color=hx(RED), lw=1.4),
        bbox=_box(),
        zorder=7,
    )
    ax.text(spot_x + 0.35, 46, "SPOT  102", color=hx(GOLD), fontsize=12.5, fontweight="medium", ha="left", va="bottom")
    ax.set_xlim(-0.6, len(strikes) - 0.4)
    ax.set_ylim(-78, 112)
    ax.set_xticks(idx)
    ax.set_xticklabels([str(k) for k in strikes], fontsize=12)
    handles = [
        Patch(facecolor=hx(GREEN), edgecolor="none", label="Long gamma"),
        Patch(facecolor=hx(RED), edgecolor="none", label="Short gamma"),
        Line2D([0], [0], color=hx(CYAN), lw=1.4, ls=(0, (3, 2)), label="Flip"),
        Line2D([0], [0], color=hx(GOLD), lw=1.4, ls=(0, (5, 3)), label="Spot"),
    ]
    leg = ax.legend(
        handles=handles,
        loc="lower right",
        frameon=True,
        facecolor=hx(CARD),
        edgecolor=hx(LINE),
        fontsize=11,
        borderpad=0.6,
        labelspacing=0.4,
    )
    for t in leg.get_texts():
        t.set_color(hx(TEXT))
    fig.text(0.07, 0.955, "EXAMPLE", color=hx(GOLD), fontsize=14, fontweight="bold", va="top")
    fig.text(
        0.97,
        0.955,
        "Made-up book  ·  not SPY  ·  not a forecast",
        color=hx(MUTED),
        fontsize=12,
        ha="right",
        va="top",
    )
    return _save(fig, "gex_example.png")


def chart_delta(w: float, h: float) -> Path:
    strikes = np.linspace(70, 130, 240)
    T = 30 / 365
    delta = call_delta(S0, strikes, T, IV)
    fig = plt.figure(figsize=(w, h), dpi=DPI)
    fig.patch.set_facecolor(hx(CARD))
    ax = fig.add_axes([0.10, 0.16, 0.86, 0.76])
    _style_ax(ax, xlabel="Strike", ylabel="Call delta")
    ax.axhspan(0.40, 0.60, color=hx(GOLD), alpha=0.14, zorder=0)
    ax.axvline(100, color=hx(GOLD), lw=1.1, ls=(0, (4, 3)), zorder=1)
    ax.plot(strikes, delta, color=hx(CYAN), lw=2.6, zorder=3)
    k50 = float(strikes[np.argmin(np.abs(delta - 0.50))])
    d50 = float(call_delta(S0, k50, T, IV))
    ax.scatter([k50], [d50], s=42, color=hx(WHITE), zorder=4)
    ax.annotate(
        "delta 0.50\n≈ $0.50 per share\nrough coin-flip ITM",
        xy=(k50, d50),
        xytext=(112, 0.78),
        ha="left",
        va="center",
        color=hx(TEXT),
        fontsize=12.5,
        linespacing=1.35,
        arrowprops=dict(arrowstyle="-|>", color=hx(CYAN), lw=1.4),
        bbox=_box(),
        zorder=5,
    )
    ax.text(72, 0.70, "deep ITM", color=hx(MUTED), fontsize=12)
    ax.text(112, 0.50, "0.40–0.60 window", color=hx(GOLD), fontsize=12.5, fontweight="medium", va="center")
    ax.text(118, 0.10, "far OTM", color=hx(MUTED), fontsize=12, ha="right")
    ax.text(101.2, 0.97, "spot 100", color=hx(GOLD), fontsize=12)
    ax.set_xlim(70, 130)
    ax.set_ylim(0, 1.05)
    ax.set_yticks([0, 0.2, 0.4, 0.6, 0.8, 1.0])
    return _save(fig, "delta.png")


def chart_gamma(w: float, h: float) -> Path:
    prices = np.linspace(80, 120, 360)
    series = [
        (7, GOLD, "7 days"),
        (30, CYAN, "30 days"),
        (90, "8FA4C8", "90 days"),
    ]
    fig = plt.figure(figsize=(w, h), dpi=DPI)
    fig.patch.set_facecolor(hx(CARD))
    ax = fig.add_axes([0.11, 0.16, 0.85, 0.74])
    _style_ax(ax, xlabel="Underlying price", ylabel="Gamma   ·   delta change per $1")
    peak_xy = None
    for days, color, name in series:
        g = call_gamma(prices, S0, days / 365, IV)
        ax.plot(prices, g, color=hx(color), lw=2.5, label=name, zorder=3)
        if days == 7:
            i = int(np.argmax(g))
            peak_xy = (float(prices[i]), float(g[i]))
    ax.axvline(100, color=hx(GOLD), lw=1.05, ls=(0, (4, 3)), zorder=1)
    ax.scatter([peak_xy[0]], [peak_xy[1]], s=36, color=hx(GOLD), zorder=4)
    ax.annotate(
        "tallest at the money\nnear expiration",
        xy=peak_xy,
        xytext=(84, peak_xy[1] * 0.72),
        ha="left",
        va="center",
        color=hx(GOLD),
        fontsize=12.5,
        linespacing=1.35,
        arrowprops=dict(arrowstyle="-|>", color=hx(GOLD), lw=1.4),
        bbox=_box(),
        zorder=5,
    )
    ymax = peak_xy[1] * 1.32
    ax.set_ylim(0, ymax)
    ax.set_xlim(80, 120)
    ax.text(105.5, ymax * 0.94, "strike 100", color=hx(GOLD), fontsize=12, ha="left", va="top")
    leg = ax.legend(
        loc="upper right",
        frameon=True,
        facecolor=hx(CARD),
        edgecolor=hx(LINE),
        fontsize=12,
        borderpad=0.55,
    )
    for t in leg.get_texts():
        t.set_color(hx(TEXT))
    return _save(fig, "gamma.png")


def chart_theta(w: float, h: float) -> Path:
    dte = np.linspace(90, 1, 360)
    decay = -call_theta_daily(S0, S0, dte / 365, IV)  # positive cost
    fig = plt.figure(figsize=(w, h), dpi=DPI)
    fig.patch.set_facecolor(hx(CARD))
    ax = fig.add_axes([0.11, 0.16, 0.85, 0.74])
    _style_ax(ax, xlabel="Days to expiration    →    expiration", ylabel="Daily decay   ·   $ per share")
    ax.axvspan(1, 21, color=hx(GOLD), alpha=0.13, zorder=0)
    ax.plot(dte, decay, color=hx(GOLD), lw=2.6, zorder=3)
    ax.set_xlim(90, 1)  # time moves left to right, toward expiration
    ymax = float(decay.max()) * 1.22
    ax.set_ylim(0, ymax)
    # Point at the bend inside the shaded window, not only the final spike.
    i21 = int(np.argmin(np.abs(dte - 21)))
    i7 = int(np.argmin(np.abs(dte - 7)))
    ax.annotate(
        "speeds up in the\nlast 2–3 weeks",
        xy=(dte[i7], decay[i7]),
        xytext=(55, ymax * 0.62),
        ha="center",
        va="center",
        color=hx(GOLD),
        fontsize=12.5,
        linespacing=1.35,
        arrowprops=dict(arrowstyle="-|>", color=hx(GOLD), lw=1.4),
        bbox=_box(),
        zorder=5,
    )
    ax.text(21, ymax * 0.08, "21 days", color=hx(GOLD), fontsize=11, ha="right")
    ax.scatter([dte[i21]], [decay[i21]], s=28, color=hx(WHITE), zorder=4)
    return _save(fig, "theta.png")


def chart_vega(w: float, h: float) -> Path:
    ivs = np.linspace(0.10, 0.60, 200)
    p30 = call_price(S0, S0, 30 / 365, ivs)
    p90 = call_price(S0, S0, 90 / 365, ivs)
    fig = plt.figure(figsize=(w, h), dpi=DPI)
    fig.patch.set_facecolor(hx(CARD))
    ax = fig.add_axes([0.11, 0.16, 0.85, 0.74])
    _style_ax(ax, xlabel="Implied volatility", ylabel="Call premium   ·   $ per share")
    ax.plot(ivs * 100, p90, color=hx(CYAN), lw=2.4, label="90 days", zorder=3)
    ax.plot(ivs * 100, p30, color=hx(PURPLE), lw=2.6, label="30 days", zorder=3)
    y30 = float(call_price(S0, S0, 30 / 365, 0.30))
    y18 = float(call_price(S0, S0, 30 / 365, 0.18))
    ax.scatter([30, 18], [y30, y18], s=42, color=hx(GOLD), zorder=4)
    loss = y30 - y18
    # Under the 30-day curve, clear of both lines. Arrow points back at the crush.
    ax.annotate(
        f"IV crush   30% → 18%\n−${loss:.2f} per share\n−${loss * 100:.0f} per contract",
        xy=(30, y30),
        xytext=(55, 1.45),
        ha="center",
        va="center",
        color=hx(GOLD),
        fontsize=12.5,
        linespacing=1.35,
        arrowprops=dict(arrowstyle="-|>", color=hx(GOLD), lw=1.3),
        bbox=_box(),
        zorder=5,
    )
    ax.set_xlim(8, 68)
    ax.set_ylim(0, float(p90.max()) * 1.12)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:.0f}%"))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v:.0f}"))
    leg = ax.legend(
        loc="upper left",
        frameon=True,
        facecolor=hx(CARD),
        edgecolor=hx(LINE),
        fontsize=12,
        borderpad=0.55,
        title="ATM call",
    )
    leg.get_title().set_color(hx(MUTED))
    for t in leg.get_texts():
        t.set_color(hx(TEXT))
    return _save(fig, "vega.png")


# --- slide primitives ------------------------------------------------------
ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
}
ANCHOR = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}


def _apply_run(run, text, size, color, bold, font, tracking):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "en-US")
    if tracking:
        rpr.set("spc", str(int(tracking)))
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", font)


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=16,
    color=TEXT,
    bold=False,
    align="left",
    anchor="top",
    font=FONT,
    tracking=None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    tf.anchor = ANCHOR[anchor]
    p = tf.paragraphs[0]
    p.alignment = ALIGN[align]
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    p.line_spacing = 1.0
    run = p.add_run()
    _apply_run(run, text, size, color, bold, font, tracking)
    return box


def add_lines(slide, lines, x, y, w, h, *, anchor="top", align="left"):
    """lines: dicts with text, size, color, bold, before, after, spacing, align, font."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.anchor = ANCHOR[anchor]
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[item.get("align", align)]
        p.space_before = Pt(item.get("before", 0))
        p.space_after = Pt(item.get("after", 0))
        p.line_spacing = item.get("spacing", 1.02)
        run = p.add_run()
        _apply_run(
            run,
            item["text"],
            item.get("size", 16),
            item.get("color", TEXT),
            item.get("bold", False),
            item.get("font", FONT),
            item.get("tracking"),
        )
    return box


def add_shape(slide, kind, x, y, w, h, fill, line=None, line_w=1.0, radius=None):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
    if radius is not None:
        shp.adjustments[0] = radius
    return shp


def add_line(slide, x1, y1, x2, y2, color=LINE, pt=1.0):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(pt)
    return conn


def blank_slide(prs) -> object:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for shp in list(slide.placeholders):
        elm = shp._element
        elm.getparent().remove(elm)
    # Paint the master background and a full-bleed rect so the PDF has no white edge.
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(BG)
    add_shape(slide, MSO_SHAPE.RECTANGLE, -0.02, -0.02, SW + 0.04, SH + 0.04, BG)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, 0.045, GOLD)
    return slide


def content_header(slide, kicker: str, title: str, dek: str | None = None) -> float:
    add_text(slide, kicker, 0.42, 0.16, 12.4, 0.30, size=13, color=GOLD, bold=True, tracking=90)
    add_text(slide, title, 0.40, 0.44, 12.5, 0.52, size=30, color=TEXT, bold=True)
    if dek:
        add_text(slide, dek, 0.42, 1.00, 12.45, 0.36, size=15, color=MUTED)
        add_line(slide, 0.42, 1.44, 12.91, 1.44, LINE, 1.0)
        return 1.60
    add_line(slide, 0.42, 1.08, 12.91, 1.08, LINE, 1.0)
    return 1.24


def add_footer(slide, number: int, total: int) -> None:
    add_line(slide, 0.42, 7.06, 12.91, 7.06, LINE, 1.0)
    add_text(slide, "GAMMA  &  THE GREEKS", 0.42, 7.14, 4.6, 0.28, size=11, color=DIM, tracking=60)
    add_text(
        slide,
        "Educational briefing",
        4.7,
        7.14,
        3.9,
        0.28,
        size=11,
        color=DIM,
        align="center",
    )
    add_text(
        slide,
        f"{number:02d}   /   {total:02d}",
        9.2,
        7.14,
        3.7,
        0.28,
        size=11,
        color=DIM,
        align="right",
    )


def chart_panel(slide, path: Path, x, y, w, h) -> None:
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, CARD)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    border = add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, CARD, line=LINE, line_w=1.0)
    border.fill.background()


def card(slide, x, y, w, h, fill=CARD, line=LINE, radius=0.08):
    return add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line=line, line_w=1.0, radius=radius
    )


def accent_bar(slide, x, y, h, color, inset=0.16):
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y + inset, 0.055, max(0.2, h - 2 * inset), color)


# --- slides ----------------------------------------------------------------
def slide_title(slide) -> None:
    add_text(slide, "FOR DOMINIC", 0.55, 1.72, 8, 0.32, size=14, color=GOLD, bold=True, tracking=140)
    add_text(slide, "Gamma & the Greeks", 0.52, 2.10, 12, 0.86, size=52, color=TEXT, bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.58, 3.12, 2.15, 0.035, GOLD)
    add_text(
        slide,
        "How dealer hedging moves price, and how to trade the Greeks.",
        0.55,
        3.32,
        10.5,
        0.42,
        size=20,
        color=MUTED,
    )
    add_text(
        slide,
        "Retail cash account   ·   SPY snapshot 16 Oct 2026",
        0.55,
        3.90,
        10,
        0.34,
        size=15,
        color=DIM,
    )
    add_text(slide, "IN THIS BRIEFING", 0.55, 5.22, 6, 0.28, size=12, color=DIM, bold=True, tracking=80)
    items = [
        ("01", "The hedge"),
        ("02", "The flip"),
        ("03", "Walls and risk"),
        ("04", "The Greeks"),
    ]
    gap = 0.16
    x0, total = 0.55, 12.22
    cw = (total - 3 * gap) / 4
    for i, (num, label) in enumerate(items):
        x = x0 + i * (cw + gap)
        card(slide, x, 5.58, cw, 1.05, fill=CARD)
        add_text(slide, num, x + 0.16, 5.68, cw - 0.32, 0.30, size=13, color=GOLD, bold=True)
        add_text(slide, label, x + 0.16, 5.98, cw - 0.32, 0.42, size=18, color=TEXT, bold=True)
    add_text(
        slide,
        "Educational briefing. Not a recommendation to trade.",
        0.55,
        7.12,
        8,
        0.26,
        size=12,
        color=DIM,
    )


def slide_hedge(slide) -> None:
    y = content_header(
        slide,
        "01     DEALER HEDGING",
        "Dealers hedge with the underlying",
        "They trade shares against options they have sold. That flow is a push on price.",
    )
    steps = [
        ("01", "SOLD", "Customers buy the option.", "The dealer is short it."),
        ("02", "DELTA MOVES", "Spot changes.", "Delta changes with it."),
        ("03", "HEDGE", "They buy or sell shares", "to get flat again."),
        ("04", "PRICE IMPACT", "Those shares hit the tape", "and move the price."),
    ]
    gap = 0.34
    x0, total = 0.42, 12.49
    cw = (total - 3 * gap) / 4
    top, ch = y, 2.42
    for i, (num, title, a, b) in enumerate(steps):
        x = x0 + i * (cw + gap)
        card(slide, x, top, cw, ch)
        accent_bar(slide, x + 0.12, top, ch, GOLD, inset=0.18)
        add_text(slide, num, x + 0.28, top + 0.16, cw - 0.42, 0.28, size=13, color=GOLD, bold=True)
        add_text(slide, title, x + 0.26, top + 0.48, cw - 0.40, 0.64, size=16, color=TEXT, bold=True)
        add_lines(
            slide,
            [
                {"text": a, "size": 14, "color": MUTED, "after": 2},
                {"text": b, "size": 14, "color": MUTED},
            ],
            x + 0.26,
            top + 1.20,
            cw - 0.42,
            0.95,
        )
        if i < 3:
            ax0 = x + cw
            add_shape(
                slide,
                MSO_SHAPE.RIGHT_ARROW,
                ax0 + (gap - 0.20) / 2,
                top + ch / 2 - 0.08,
                0.20,
                0.16,
                GOLD,
            )
    add_text(
        slide,
        "GAMMA DECIDES WHICH WAY THAT HEDGE FACES",
        0.42,
        top + ch + 0.16,
        12,
        0.28,
        size=12,
        color=DIM,
        bold=True,
        tracking=70,
    )
    outcomes = [
        (GREEN, "LONG GAMMA", "Above the flip", "Buy dips, sell rips.", "Price ranges and pins."),
        (RED, "SHORT GAMMA", "Below the flip", "Sell drops, buy rips.", "Moves accelerate and trend."),
    ]
    oy = top + ch + 0.50
    oh = 6.90 - oy
    ow = (total - 0.18) / 2
    for i, (accent, title, where, a, b) in enumerate(outcomes):
        x = x0 + i * (ow + 0.18)
        card(slide, x, oy, ow, oh)
        accent_bar(slide, x + 0.14, oy, oh, accent, inset=0.16)
        add_text(slide, title, x + 0.36, oy + 0.14, ow - 0.52, 0.32, size=14, color=accent, bold=True, tracking=40)
        add_text(slide, where, x + 0.34, oy + 0.48, ow - 0.52, 0.36, size=20, color=TEXT, bold=True)
        add_lines(
            slide,
            [
                {"text": a, "size": 15, "color": MUTED, "after": 2},
                {"text": b, "size": 15, "color": MUTED},
            ],
            x + 0.34,
            oy + 0.96,
            ow - 0.55,
            0.80,
        )


def _hedge_numbers(slide, x, y, w, h, rows, caption):
    """rows: (label, value, note, value_color). Value is a single short token."""
    card(slide, x, y, w, h, fill=CARD_2, line=LINE, radius=0.1)
    n = len(rows)
    gap = 0.06
    cw = (w - 0.24 - gap * (n - 1)) / n
    for i, (label, value, note, color) in enumerate(rows):
        cx = x + 0.12 + i * (cw + gap)
        add_text(slide, label, cx, y + 0.12, cw, 0.24, size=12, color=DIM, bold=True, tracking=40)
        add_text(slide, value, cx, y + 0.38, cw, 0.48, size=26, color=color, bold=True)
        add_text(slide, note, cx, y + 0.92, cw, 0.30, size=13, color=MUTED)
    add_text(slide, caption, x + 0.12, y + h - 0.34, w - 0.24, 0.26, size=12, color=DIM)


def slide_positive(slide) -> None:
    y = content_header(
        slide,
        "02     POSITIVE GAMMA",
        "Buy dips, sell rips",
        "Above the flip. Dealers are long gamma, so the hedge fades the move.",
    )
    left_w, gap = 5.05, 0.18
    card(slide, 0.42, y, left_w, 6.90 - y)
    accent_bar(slide, 0.56, y, 6.90 - y, GREEN, inset=0.18)
    add_text(slide, "ABOVE THE FLIP", 0.78, y + 0.16, 4.4, 0.26, size=13, color=GREEN, bold=True, tracking=70)
    add_text(slide, "Price ranges and pins", 0.76, y + 0.46, 4.5, 0.40, size=20, color=TEXT, bold=True)
    add_lines(
        slide,
        [
            {"text": "Dealers are net long options.", "size": 15, "color": MUTED, "after": 6},
            {"text": "They buy dips and sell rips.", "size": 15, "color": MUTED, "after": 6},
            {"text": "The hedge fights the move.", "size": 15, "color": MUTED},
        ],
        0.76,
        y + 1.00,
        4.45,
        1.35,
    )
    _hedge_numbers(
        slide,
        0.68,
        y + 2.50,
        left_w - 0.48,
        2.15,
        [
            ("START", "500", "short shares", TEXT),
            ("UP  $1", "100", "sell the rip", RED),
            ("DOWN  $1", "100", "buy the dip", GREEN),
        ],
        "Example · long 10 calls · long puts fade too",
    )
    cw = 12.91 - (0.42 + left_w + gap)
    ch = 6.90 - y
    chart_panel(slide, chart_positive(cw, ch), 0.42 + left_w + gap, y, cw, ch)


def slide_negative(slide) -> None:
    y = content_header(
        slide,
        "03     NEGATIVE GAMMA",
        "Sell drops, buy rips",
        "Below the flip. Dealers are short gamma, so the hedge chases the move.",
    )
    left_w, gap = 5.05, 0.18
    card(slide, 0.42, y, left_w, 6.90 - y)
    accent_bar(slide, 0.56, y, 6.90 - y, RED, inset=0.18)
    add_text(slide, "BELOW THE FLIP", 0.78, y + 0.16, 4.4, 0.26, size=13, color=RED, bold=True, tracking=70)
    add_text(slide, "Moves accelerate", 0.76, y + 0.46, 4.5, 0.40, size=20, color=TEXT, bold=True)
    add_lines(
        slide,
        [
            {"text": "Dealers are net short options.", "size": 15, "color": MUTED, "after": 6},
            {"text": "They sell drops and buy rips.", "size": 15, "color": MUTED, "after": 6},
            {"text": "The hedge feeds the move.", "size": 15, "color": MUTED},
        ],
        0.76,
        y + 1.00,
        4.45,
        1.35,
    )
    _hedge_numbers(
        slide,
        0.68,
        y + 2.50,
        left_w - 0.48,
        2.15,
        [
            ("START", "500", "long shares", TEXT),
            ("UP  $1", "100", "buy the rip", GREEN),
            ("DOWN  $1", "100", "sell the drop", RED),
        ],
        "Example · short 10 calls · short puts chase too",
    )
    cw = 12.91 - (0.42 + left_w + gap)
    ch = 6.90 - y
    chart_panel(slide, chart_negative(cw, ch), 0.42 + left_w + gap, y, cw, ch)


def slide_levels(slide) -> None:
    y = content_header(
        slide,
        "04     THREE LEVELS",
        "Flip, call wall, put wall",
        "Three strikes that change the speed of the tape.",
    )
    specs = [
        (
            GOLD,
            "THE BOUNDARY",
            "Gamma flip",
            [
                "Dealer gamma changes sign here.",
                "Above it, they fade the move.",
                "Below it, they chase the move.",
            ],
            "A loss of the flip is a regime change.",
        ),
        (
            GREEN,
            "ABOVE SPOT",
            "Call wall",
            [
                "The large call-gamma strike above spot.",
                "It can resist, or act as a magnet.",
                "Price often slows as it gets there.",
            ],
            "A break above it can squeeze higher.",
        ),
        (
            RED,
            "BELOW SPOT",
            "Put wall",
            [
                "The large put-gamma strike below spot.",
                "It is support while it holds.",
                "Dealers are leaning against a drop.",
            ],
            "A break below it speeds up selling.",
        ),
    ]
    gap = 0.16
    cw = (12.49 - 2 * gap) / 3
    ch = 4.55
    for i, (accent, tag, title, lines, foot) in enumerate(specs):
        x = 0.42 + i * (cw + gap)
        card(slide, x, y, cw, ch)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, cw, 0.06, accent)
        add_text(slide, tag, x + 0.22, y + 0.24, cw - 0.40, 0.28, size=13, color=accent, bold=True, tracking=70)
        add_text(slide, title, x + 0.20, y + 0.58, cw - 0.40, 0.48, size=26, color=TEXT, bold=True)
        add_lines(
            slide,
            [{"text": line, "size": 15, "color": MUTED, "before": 8} for line in lines],
            x + 0.20,
            y + 1.30,
            cw - 0.42,
            1.85,
        )
        add_line(slide, x + 0.22, y + ch - 1.15, x + cw - 0.22, y + ch - 1.15, LINE, 1.0)
        add_text(slide, foot, x + 0.20, y + ch - 0.95, cw - 0.42, 0.72, size=15, color=TEXT, bold=True)
    add_text(
        slide,
        "Walls slow price. They are not floors or ceilings. A break changes the speed.",
        0.42,
        y + ch + 0.16,
        12.4,
        0.36,
        size=15,
        color=MUTED,
    )


def slide_example(slide) -> None:
    y = content_header(
        slide,
        "05     EXAMPLE BOOK",
        "Gamma by strike",
        "Illustrative numbers only, so the labels are easy to see. The SPY snapshot is next.",
    )
    ch = 6.90 - y
    chart_panel(slide, chart_gex(12.49, ch), 0.42, y, 12.49, ch)


def slide_spy(slide) -> None:
    y = content_header(
        slide,
        "06     SPY SNAPSHOT",
        "SPY at the 16 Oct 2026 close",
        "Real levels. Spot about 763.25, above the flip, inside a one-point pocket.",
    )
    # True price scale. Color lives in a narrow strip so labels sit on the card, not on the wash.
    px, py, pw, ph = 0.42, y, 7.55, 6.90 - y
    card(slide, px, py, pw, ph, radius=0.06)
    p_hi, p_lo = 766.4, 759.4
    y_top = py + 0.48
    y_bot = py + ph - 0.46
    strip_x = px + 2.22
    strip_w = 0.36
    rail = strip_x + strip_w + 0.10

    def price_y(p):
        return y_top + (p_hi - p) / (p_hi - p_lo) * (y_bot - y_top)

    y_call, y_spot, y_put, y_flip = map(price_y, (SPY_CALL, SPY_SPOT, SPY_PUT, SPY_FLIP))
    add_shape(slide, MSO_SHAPE.RECTANGLE, strip_x, y_top, strip_w, max(0.2, y_flip - y_top), POS_WASH)
    add_shape(slide, MSO_SHAPE.RECTANGLE, strip_x, y_flip, strip_w, max(0.2, y_bot - y_flip), NEG_WASH)
    add_shape(slide, MSO_SHAPE.RECTANGLE, strip_x, y_call, strip_w, max(0.12, y_put - y_call), PIN_WASH)
    add_line(slide, rail, y_top, rail, y_bot, DIM, 1.35)
    for yy, col in ((y_call, GREEN), (y_spot, GOLD), (y_put, RED), (y_flip, CYAN)):
        add_line(slide, rail - 0.02, yy, rail + 0.14, yy, col, 1.75)

    # Spot stays on the true tick, left of the strip, so no line runs through the words.
    add_text(
        slide,
        "SPOT  ≈ 763.25",
        px + 0.14,
        y_spot - 0.18,
        strip_x - px - 0.28,
        0.36,
        size=14,
        color=GOLD,
        bold=True,
        align="right",
        anchor="middle",
    )
    label_x = rail + 0.26
    label_w = px + pw - label_x - 0.16
    add_text(slide, "POSITIVE GAMMA", label_x, y_top + 0.02, label_w, 0.28, size=13, color=GREEN, bold=True, tracking=40)
    add_text(slide, "CALL WALL    764", label_x, y_call - 0.17, label_w, 0.34, size=16, color=GREEN, bold=True)
    # 763 and 763.25 are $0.25 apart — nudge the put-wall label down and lead it back to the tick.
    put_label = y_put + 0.30
    add_line(slide, rail + 0.14, y_put, label_x - 0.04, put_label + 0.17, RED, 1.15)
    add_text(slide, "PUT WALL    763", label_x, put_label, label_w, 0.34, size=16, color=RED, bold=True)
    add_text(slide, "GAMMA FLIP    761.19", label_x, y_flip - 0.17, label_w, 0.34, size=16, color=CYAN, bold=True)
    add_text(slide, "NEGATIVE GAMMA", label_x, y_flip + 0.36, label_w, 0.28, size=13, color=RED, bold=True, tracking=40)
    add_text(slide, "True price scale", px + 0.18, py + ph - 0.32, 2.6, 0.24, size=12, color=DIM)

    rx = px + pw + 0.16
    rw = 12.91 - rx
    cards = [
        (GREEN, "REGIME", "Above the flip by $2.06.\nDealers fade. A range\nis the base case."),
        (GOLD, "THE POCKET", "763 under spot, 764 over it.\nFade into the walls.\nTake profit at the walls."),
        (RED, "IF IT BREAKS", "Above 764 can squeeze.\nBelow 763, selling speeds up.\n4h close under 761.19\nchanges the regime."),
    ]
    # The third card has four lines; give the stack even heights and let text sit inside.
    gap = 0.12
    ch = (ph - 2 * gap) / 3
    for i, (accent, title, body) in enumerate(cards):
        cy = py + i * (ch + gap)
        card(slide, rx, cy, rw, ch)
        accent_bar(slide, rx + 0.12, cy, ch, accent, inset=0.14)
        add_text(slide, title, rx + 0.30, cy + 0.10, rw - 0.46, 0.28, size=13, color=accent, bold=True, tracking=50)
        parts = body.split("\n")
        add_lines(
            slide,
            [{"text": part, "size": 14, "color": TEXT, "before": 3, "spacing": 1.02} for part in parts],
            rx + 0.28,
            cy + 0.42,
            rw - 0.44,
            ch - 0.52,
        )


def slide_risk(slide) -> None:
    y = content_header(
        slide,
        "07     RISK",
        "Trade the regime you are in",
        "Same levels. Opposite trades. The flip picks the playbook.",
    )
    panels = [
        (
            GREEN,
            "ABOVE THE FLIP",
            "Fade",
            [
                "Fade into the walls.",
                "Take profit at the walls.",
                "The base case is a range.",
            ],
        ),
        (
            RED,
            "BELOW THE FLIP",
            "Go with it",
            [
                "Trade with the move.",
                "Size down. Use wider stops.",
                "Do not fade.",
            ],
        ),
    ]
    gap = 0.16
    pw = (12.49 - gap) / 2
    ph = 3.05
    for i, (accent, tag, title, lines) in enumerate(panels):
        x = 0.42 + i * (pw + gap)
        card(slide, x, y, pw, ph)
        accent_bar(slide, x + 0.16, y, ph, accent, inset=0.18)
        add_text(slide, tag, x + 0.40, y + 0.16, pw - 0.58, 0.26, size=13, color=accent, bold=True, tracking=60)
        add_text(slide, title, x + 0.38, y + 0.46, pw - 0.56, 0.42, size=26, color=TEXT, bold=True)
        add_lines(
            slide,
            [{"text": line, "size": 16, "color": MUTED, "before": 6} for line in lines],
            x + 0.38,
            y + 1.05,
            pw - 0.58,
            1.75,
        )
    notes = [
        (GOLD, "STOPS", "Put them beyond the wall, not on it. A tag of the wall is normal. A stop sitting on it gets run."),
        (CYAN, "REGIME CHANGE", "Losing the flip ends the fade. Switch playbooks. Do not keep fading a trend."),
        (GOLD, "4-HOUR CLOSE", "Confirm on a 4h candle close. A wick through the level is not a break."),
    ]
    ny = y + ph + 0.14
    nh = 6.90 - ny
    nw = (12.49 - 2 * gap) / 3
    for i, (accent, title, body) in enumerate(notes):
        x = 0.42 + i * (nw + gap)
        card(slide, x, ny, nw, nh)
        add_text(slide, title, x + 0.18, ny + 0.12, nw - 0.34, 0.28, size=13, color=accent, bold=True, tracking=50)
        add_text(slide, body, x + 0.16, ny + 0.46, nw - 0.32, nh - 0.58, size=14, color=TEXT)


def slide_time(slide) -> None:
    y = content_header(
        slide,
        "08     TIMEFRAMES",
        "Which clock you are on",
        "The same levels do a different job on each clock.",
    )
    items = [
        ("INTRADAY", "Levels matter most", "The flip and the walls\nguide the session."),
        ("EXPIRATION WEEK", "Price tends to pin", "Large strikes act\nas magnets."),
        ("THE WEEK AFTER", "Ranges expand", "The pin releases once\nexpiration passes."),
        ("OVER A MONTH", "Regime sets size", "Above or below the flip\ndecides how big."),
    ]
    gap = 0.14
    cw = (12.49 - 3 * gap) / 4
    ch = 3.85
    top = y + (6.90 - y - ch) / 2
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.42 + 0.3, top + 0.36, 12.49 - 0.6, 0.015, LINE)
    for i, (tag, title, body) in enumerate(items):
        x = 0.42 + i * (cw + gap)
        card(slide, x, top, cw, ch)
        add_shape(slide, MSO_SHAPE.OVAL, x + 0.22, top + 0.28, 0.18, 0.18, GOLD)
        add_text(slide, tag, x + 0.18, top + 0.58, cw - 0.36, 0.48, size=13, color=GOLD, bold=True)
        add_text(slide, title, x + 0.18, top + 1.15, cw - 0.36, 0.78, size=20, color=TEXT, bold=True)
        parts = body.split("\n")
        add_lines(
            slide,
            [{"text": part, "size": 15, "color": MUTED, "before": 3} for part in parts],
            x + 0.16,
            top + 2.05,
            cw - 0.32,
            1.2,
        )


def _greek_cards(slide, y, accent, cards_spec):
    x, w = 0.42, 4.55
    h = 6.90 - y
    gap = 0.12
    ch = (h - 2 * gap) / 3
    for i, (title, lines) in enumerate(cards_spec):
        cy = y + i * (ch + gap)
        card(slide, x, cy, w, ch)
        accent_bar(slide, x + 0.12, cy, ch, accent, inset=0.14)
        add_text(slide, title, x + 0.30, cy + 0.10, w - 0.46, 0.40, size=14, color=accent, bold=True, tracking=40)
        add_lines(
            slide,
            [{"text": line, "size": 14, "color": TEXT, "before": 3} for line in lines],
            x + 0.28,
            cy + 0.52,
            w - 0.44,
            ch - 0.62,
        )
    return 5.15, y, 12.91 - 5.15, h


def _foot_note(slide, x, y, w, text):
    add_text(slide, text, x, y, w, 0.28, size=12, color=DIM)


def slide_delta(slide) -> None:
    y = content_header(
        slide,
        "09     DELTA",
        "The move, and the odds",
        "Per $1 of the underlying, and a rough read on finishing in the money.",
    )
    cx, cy, cw, ch = _greek_cards(
        slide,
        y,
        CYAN,
        [
            ("PER $1", ["0.40 delta ≈ $0.40 a share.", "About $40 a contract,", "before gamma updates."]),
            ("ROUGH ODDS", ["Traders read |delta| as the", "chance of finishing ITM.", "Close to the model. Not equal."]),
            ("THE WINDOW", ["The band is 0.40–0.60.", "That is the zone for a single", "when the move should be fast."]),
        ],
    )
    # Leave a strip under the chart for the model note.
    chart_h = ch - 0.36
    chart_panel(slide, chart_delta(cw, chart_h), cx, cy, cw, chart_h)
    _foot_note(
        slide,
        cx,
        cy + chart_h + 0.06,
        cw,
        "Black–Scholes call  ·  spot 100  ·  IV 20%  ·  30 days  ·  illustrative",
    )


def slide_gamma(slide) -> None:
    y = content_header(
        slide,
        "10     GAMMA",
        "How fast delta changes",
        "Highest at the money, and tallest when expiration is close.",
    )
    cx, cy, cw, ch = _greek_cards(
        slide,
        y,
        GREEN,
        [
            ("PER $1", ["Gamma of 0.05 means delta", "changes by about 0.05", "if the underlying moves $1."]),
            ("THREE CURVES", ["Same strike, three expirations.", "7 days spikes.", "90 days is a low hill."]),
            ("YOU PAY FOR IT", ["Long gamma wants movement.", "A quiet tape hands it back", "as theta."]),
        ],
    )
    chart_h = ch - 0.36
    chart_panel(slide, chart_gamma(cw, chart_h), cx, cy, cw, chart_h)
    _foot_note(
        slide,
        cx,
        cy + chart_h + 0.06,
        cw,
        "Black–Scholes  ·  strike 100  ·  IV 20%  ·  7, 30, and 90 days  ·  illustrative",
    )


def slide_theta(slide) -> None:
    y = content_header(
        slide,
        "11     THETA",
        "What a day costs",
        "Daily decay of a long option. It speeds up in the last 2–3 weeks.",
    )
    cx, cy, cw, ch = _greek_cards(
        slide,
        y,
        GOLD,
        [
            ("PER SHARE", ["The curve is the daily cost", "of holding, per share.", "Multiply by 100 for one lot."]),
            ("THE BEND", ["Shading is the last 21 days.", "That is where decay", "gets expensive."]),
            ("QUIET TAPE", ["Short-dated longs bleed", "when price does nothing.", "That is the theta bill."]),
        ],
    )
    chart_h = ch - 0.36
    chart_panel(slide, chart_theta(cw, chart_h), cx, cy, cw, chart_h)
    _foot_note(
        slide,
        cx,
        cy + chart_h + 0.06,
        cw,
        "ATM call  ·  spot = strike = 100  ·  IV 20%  ·  calendar-day decay  ·  illustrative",
    )


def slide_vega(slide) -> None:
    y = content_header(
        slide,
        "12     VEGA",
        "What IV can take back",
        "The slope is vega. A drop in IV is a crush — even if direction was right.",
    )
    cx, cy, cw, ch = _greek_cards(
        slide,
        y,
        PURPLE,
        [
            ("THE SLOPE", ["Steeper means more vega.", "90-day premium moves more", "than 30-day premium."]),
            ("THE ARROW", ["30% IV to 18% on the", "30-day call. Premium falls", "with no move in spot."]),
            ("BEFORE YOU BUY", ["Rich IV favors a spread.", "Cheap IV can justify a single.", "Check IV first."]),
        ],
    )
    chart_h = ch - 0.36
    chart_panel(slide, chart_vega(cw, chart_h), cx, cy, cw, chart_h)
    _foot_note(
        slide,
        cx,
        cy + chart_h + 0.06,
        cw,
        "ATM calls  ·  spot = strike = 100  ·  crush marked 30% → 18%  ·  illustrative",
    )


def slide_rules(slide) -> None:
    y = content_header(
        slide,
        "13     PRACTICAL RULES",
        "How to buy it",
        "Cash account. The contract has to match the speed of the move and the IV.",
    )
    rules = [
        ("01", "Single long options", "0.40–0.60 delta, for a fast move or cheap IV."),
        ("02", "Debit spreads", "For a slow move or rich IV. You sell the expensive wing."),
        ("03", "Quiet stretches", "Do not hold short-dated options when nothing is happening."),
        ("04", "Check IV", "Know what you are paying. Crush can give the gain back."),
        ("05", "Cash account only", "No margin. The most you lose is the debit you paid."),
    ]
    n = len(rules)
    gap = 0.10
    rh = (6.90 - y - gap * (n - 1)) / n
    for i, (num, title, body) in enumerate(rules):
        ry = y + i * (rh + gap)
        card(slide, 0.42, ry, 12.49, rh, radius=0.08)
        add_text(slide, num, 0.60, ry, 0.70, rh, size=18, color=GOLD, bold=True, anchor="middle")
        add_text(slide, title, 1.45, ry + 0.10, 10.8, 0.36, size=18, color=TEXT, bold=True)
        add_text(slide, body, 1.45, ry + 0.46, 11.1, 0.36, size=15, color=MUTED)


def slide_caveats(slide) -> None:
    y = content_header(
        slide,
        "14     CAVEATS",
        "Context, not a trigger",
        "Use the map. Do not let the map take the trade.",
    )
    items = [
        (
            "01",
            "Estimates",
            "Gamma numbers are a model. They shift through the day as price, time, and IV change.",
        ),
        (
            "02",
            "Next to the tape",
            "Read them with moving averages and price action. They map positioning, not a signal.",
        ),
        (
            "03",
            "Never alone",
            "If the level and the tape disagree, trust the tape. Do not trade gamma by itself.",
        ),
    ]
    gap = 0.12
    rh = 1.38
    for i, (num, title, body) in enumerate(items):
        ry = y + i * (rh + gap)
        card(slide, 0.42, ry, 12.49, rh)
        add_text(slide, num, 0.62, ry, 0.80, rh, size=22, color=GOLD, bold=True, anchor="middle")
        add_text(slide, title, 1.60, ry + 0.16, 10.6, 0.40, size=22, color=TEXT, bold=True)
        add_text(slide, body, 1.60, ry + 0.64, 10.8, 0.58, size=16, color=MUTED)
    by = y + 3 * rh + 2 * gap + 0.14
    bh = 6.90 - by
    card(slide, 0.42, by, 12.49, bh, fill=CARD_2)
    add_text(
        slide,
        "The example book is not SPY. These levels are one close, and they move.",
        0.70,
        by,
        11.9,
        bh,
        size=20,
        color=GOLD,
        bold=True,
        anchor="middle",
    )


def slide_summary(slide) -> None:
    y = content_header(
        slide,
        "15     SUMMARY",
        "The whole map",
        "One page. Price on the left, the contract on the right.",
    )
    cols = [
        (
            "THE TAPE",
            GREEN,
            [
                "Dealers hedge sold options by trading shares.",
                "Above the flip: buy dips, sell rips, pin.",
                "Below the flip: sell drops, buy rips, trend.",
                "Call wall resists or pins. A break can squeeze.",
                "Put wall is support. A break speeds selling.",
                "Fade above. Go with it below. Stops past the walls.",
            ],
        ),
        (
            "THE CONTRACT",
            CYAN,
            [
                "Delta: the move per $1, and rough ITM odds.",
                "Gamma peaks at the money, near expiration.",
                "Theta speeds up in the last 2–3 weeks.",
                "Vega is IV risk. Crush can take premium back.",
                "0.40–0.60 singles if fast or cheap IV.",
                "Spreads if slow or rich. Check IV. Cash only.",
            ],
        ),
    ]
    gap = 0.18
    cw = (12.49 - gap) / 2
    # Bottom bar reserved for clocks + the caveat.
    bar_h = 1.15
    col_h = 6.90 - y - bar_h - 0.12
    for i, (title, accent, lines) in enumerate(cols):
        x = 0.42 + i * (cw + gap)
        card(slide, x, y, cw, col_h)
        add_text(slide, title, x + 0.22, y + 0.12, cw - 0.40, 0.32, size=14, color=accent, bold=True, tracking=60)
        add_line(slide, x + 0.22, y + 0.52, x + cw - 0.22, y + 0.52, LINE, 1.0)
        rh = (col_h - 0.70) / len(lines)
        for j, line in enumerate(lines):
            ry = y + 0.62 + j * rh
            add_text(slide, f"{j + 1:02d}", x + 0.18, ry, 0.46, rh, size=13, color=GOLD, bold=True, anchor="middle")
            add_text(slide, line, x + 0.68, ry, cw - 0.88, rh, size=14, color=TEXT, anchor="middle")
    by = y + col_h + 0.12
    card(slide, 0.42, by, 12.49, bar_h, fill=CARD_2)
    add_text(
        slide,
        "Intraday, levels matter. Expiration week, price pins. The week after, ranges expand. Over a month, the regime sets size.",
        0.62,
        by + 0.10,
        12.1,
        0.46,
        size=14,
        color=TEXT,
        anchor="middle",
    )
    add_text(
        slide,
        "Gamma is an estimate and it moves all day. Use it with moving averages and price action — never alone.",
        0.62,
        by + 0.54,
        12.1,
        0.46,
        size=14,
        color=GOLD,
        anchor="middle",
    )


def build() -> Path:
    self_check()
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = "Gamma & the Greeks"
    prs.core_properties.subject = "Educational briefing for Dominic"
    prs.core_properties.category = "Education"
    builders = [
        (slide_title, False),
        (slide_hedge, True),
        (slide_positive, True),
        (slide_negative, True),
        (slide_levels, True),
        (slide_example, True),
        (slide_spy, True),
        (slide_risk, True),
        (slide_time, True),
        (slide_delta, True),
        (slide_gamma, True),
        (slide_theta, True),
        (slide_vega, True),
        (slide_rules, True),
        (slide_caveats, True),
        (slide_summary, True),
    ]
    made = []
    for fn, numbered in builders:
        slide = blank_slide(prs)
        fn(slide)
        made.append((slide, numbered))
    total = len(made)
    for i, (slide, numbered) in enumerate(made, start=1):
        if numbered:
            add_footer(slide, i, total)
    out = ROOT / "gamma_greeks.pptx"
    prs.save(out)
    print(f"wrote {out}  ({total} slides)")
    return out


def render(pptx: Path) -> None:
    PREVIEWS.mkdir(parents=True, exist_ok=True)
    for old in PREVIEWS.glob("*.png"):
        old.unlink()
    pdf_dir = BUILD / "pdf"
    pdf_dir.mkdir(parents=True, exist_ok=True)
    profile = Path(f"/tmp/lo-gamma-{os.getpid()}")
    if profile.exists():
        shutil.rmtree(profile)
    profile.mkdir(parents=True)
    env = os.environ.copy()
    env["SAL_USE_VCLPLUGIN"] = "svp"
    subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--norestore",
            "--convert-to",
            "pdf",
            "--outdir",
            str(pdf_dir),
            str(pptx),
        ],
        check=True,
        env=env,
    )
    pdf = pdf_dir / "gamma_greeks.pdf"
    if not pdf.exists():
        raise SystemExit(f"PDF was not created at {pdf}")
    subprocess.run(
        ["pdftoppm", "-png", "-r", "140", str(pdf), str(PREVIEWS / "slide")],
        check=True,
    )
    for png in sorted(PREVIEWS.glob("slide-*.png")):
        n = int(png.stem.split("-")[1])
        png.rename(PREVIEWS / f"slide_{n:02d}.png")
    pages = sorted(PREVIEWS.glob("slide_*.png"))
    print(f"rendered {len(pages)} previews in {PREVIEWS}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the Gamma & the Greeks deck.")
    parser.add_argument("--render", action="store_true", help="Also export preview PNGs.")
    args = parser.parse_args()
    path = build()
    if args.render:
        render(path)


if __name__ == "__main__":
    main()
