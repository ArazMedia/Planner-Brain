#!/usr/bin/env python3
"""Build the Gamma Swing Trading deck.

    python3 decks/gamma-swing/build_deck.py

Writes gamma_swing.pptx next to this file.
Charts are Black-Scholes illustrations with the inputs printed on the slide.
"""

from __future__ import annotations

import math
import os
import shutil
import subprocess
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
CHART_DIR = Path("/tmp/gamma-charts")
CHART_DIR.mkdir(parents=True, exist_ok=True)

W = 13.333333
H = 7.5
ML = 0.42
CONTENT_W = 12.49
TOP = 1.08
BOTTOM = 7.02

FONT = "Liberation Sans"
FONT_REG = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"

# Dark, quiet, high contrast. Regime colors are used as labels, not as body text.
BG = "0E1218"
CARD = "171D28"
NEST = "222A38"
LINE = "313A4B"
TEXT = "F4F7FB"
BODY = "D5DDE8"
DIM = "B4BECC"
FOOT = "A8B3C2"
TEAL = "3DDCB0"
TEAL_DIM = "14352F"
CORAL = "FF7676"
CORAL_DIM = "3A2226"
AMBER = "E6B54A"
AMBER_DIM = "332C16"
BLUE = "8EBEFF"
BLUE_DIM = "1A2838"
PURPLE = "C9B2FF"
PURPLE_DIM = "2A2340"
CYAN = "7AD7E4"
WHITE = "F4F7FB"

font_manager.fontManager.addfont(FONT_REG)
font_manager.fontManager.addfont(FONT_BOLD)
plt.rcParams["font.family"] = "Liberation Sans"
plt.rcParams["axes.unicode_minus"] = False


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def check_ascii(text: str) -> None:
    bad = [(i, ch, hex(ord(ch))) for i, ch in enumerate(text) if ord(ch) > 126]
    if bad:
        raise SystemExit(f"Non-ASCII {bad} in {text!r}")
    if "|" in text:
        raise SystemExit(f"Absolute-value bar in {text!r}")


# --- Black-Scholes (European, no dividend). Teaching charts only. ---

def _ncdf(x: float) -> float:
    return 0.5 * (1.0 + math.erf(x / math.sqrt(2.0)))


def _npdf(x: float) -> float:
    return math.exp(-0.5 * x * x) / math.sqrt(2.0 * math.pi)


def _d1(s: float, k: float, t: float, r: float, sig: float) -> float:
    return (math.log(s / k) + (r + 0.5 * sig * sig) * t) / (sig * math.sqrt(t))


def _d2(s: float, k: float, t: float, r: float, sig: float) -> float:
    return _d1(s, k, t, r, sig) - sig * math.sqrt(t)


def call_price(s: float, k: float, t: float, r: float, sig: float) -> float:
    if t <= 1e-8:
        return max(s - k, 0.0)
    return s * _ncdf(_d1(s, k, t, r, sig)) - math.exp(-r * t) * k * _ncdf(_d2(s, k, t, r, sig))


def call_delta(s: float, k: float, t: float, r: float, sig: float) -> float:
    if t <= 1e-8:
        return 1.0 if s > k else 0.0
    return _ncdf(_d1(s, k, t, r, sig))


def call_gamma(s: float, k: float, t: float, r: float, sig: float) -> float:
    return _npdf(_d1(s, k, t, r, sig)) / (s * sig * math.sqrt(t))


def bs_facts() -> dict:
    """Numbers printed on slides. All of them come from this function."""
    r = 0.04
    g30 = call_gamma(100, 100, 30 / 365, r, 0.20)
    g7 = call_gamma(100, 100, 7 / 365, r, 0.20)
    g1 = call_gamma(100, 100, 1 / 365, r, 0.20)
    shares = []
    for days in (1, 7, 30):
        t = days / 365
        shares.append((call_delta(101, 100, t, r, 0.20) - call_delta(100, 100, t, r, 0.20)) * 100)
    # 14-day ATM crush, IV 50% to 25%.
    t14 = 14 / 365
    crush_start = call_price(100, 100, t14, r, 0.50)
    crush_flat = call_price(100, 100, t14, r, 0.25)
    crush_up = call_price(102, 100, t14, r, 0.25)
    # 21-day delta bands and payoff debits. IV 25% unless noted.
    t21 = 21 / 365
    bands = []
    for strike in (106, 100, 95):
        delta = call_delta(100, strike, t21, r, 0.25)
        price = call_price(100, strike, t21, r, 0.25)
        bands.append({"k": strike, "delta": delta, "price": price, "shares": delta * 100})
    call_25 = call_price(100, 100, t21, r, 0.25)
    short_25 = call_price(100, 105, t21, r, 0.25)
    spread_25 = call_25 - short_25
    rich = {}
    for sig in (0.18, 0.40):
        atm = call_price(100, 100, t21, r, sig)
        spr = atm - call_price(100, 105, t21, r, sig)
        rich[sig] = {"call": atm, "spread": spr, "max_gain": 5.0 - spr}
    decay = []
    for days in range(45, 0, -1):
        t = days / 365
        single = call_price(100, 100, t, r, 0.25)
        spread = single - call_price(100, 105, t, r, 0.25)
        decay.append((days, single, spread))
    return {
        "g_ratio_1": g1 / g30,
        "g_ratio_7": g7 / g30,
        "shares": shares,
        "crush_start": crush_start,
        "crush_flat": crush_flat,
        "crush_up": crush_up,
        "bands": bands,
        "call_25": call_25,
        "short_25": short_25,
        "spread_25": spread_25,
        "max_gain_25": 5.0 - spread_25,
        "rich": rich,
        "decay": decay,
    }


FACTS = bs_facts()


def money(value: float) -> str:
    return f"{value:.2f}"


def _style_ax(ax, ylabel: str, xlabel: str) -> None:
    ax.set_facecolor("#171D28")
    ax.tick_params(colors="#C5CEDA", labelsize=11, width=0)
    ax.xaxis.label.set_color("#D5DDE8")
    ax.yaxis.label.set_color("#D5DDE8")
    ax.xaxis.label.set_size(12)
    ax.yaxis.label.set_size(12)
    for spine in ax.spines.values():
        spine.set_color("#313A4B")
    ax.grid(True, color="#313A4B", linewidth=0.7)
    ax.set_axisbelow(True)
    ax.set_ylabel(ylabel)
    ax.set_xlabel(xlabel)


def _legend(ax, loc="upper right") -> None:
    leg = ax.legend(loc=loc, frameon=True, fontsize=11)
    frame = leg.get_frame()
    frame.set_facecolor("#171D28")
    frame.set_edgecolor("#313A4B")
    for text in leg.get_texts():
        text.set_color("#F4F7FB")


def chart_gamma() -> Path:
    r = 0.04
    strikes = [80 + i * 0.1 for i in range(401)]
    g30 = call_gamma(100, 100, 30 / 365, r, 0.20)
    fig, ax = plt.subplots(figsize=(7.70, 5.35), dpi=160)
    fig.patch.set_facecolor("#0E1218")
    for days, color, label in (
        (30, "#8EBEFF", "30 days"),
        (7, "#E6B54A", "7 days"),
        (1, "#FF7676", "1 day"),
    ):
        ys = [call_gamma(100, k, days / 365, r, 0.20) / g30 for k in strikes]
        ax.plot(strikes, ys, color=color, linewidth=2.4, label=label)
    ax.axvline(100, color="#A8B3C2", linewidth=1.0, linestyle=(0, (3, 3)))
    _style_ax(ax, "Gamma, scaled so the 30-day peak is 1", "Strike. Spot is 100")
    _legend(ax, "upper left")
    ax.set_xlim(80, 120)
    ax.set_ylim(0, FACTS["g_ratio_1"] * 1.14)
    fig.subplots_adjust(left=0.11, right=0.97, top=0.94, bottom=0.16)
    path = CHART_DIR / "gamma.png"
    fig.savefig(path, dpi=160)
    plt.close(fig)
    return path


def chart_hedge_shares() -> Path:
    labels = ["1 day", "7 days", "30 days"]
    values = FACTS["shares"]
    fig, ax = plt.subplots(figsize=(6.20, 5.28), dpi=160)
    fig.patch.set_facecolor("#0E1218")
    ax.bar(labels, values, color=["#FF7676", "#E6B54A", "#8EBEFF"], width=0.62)
    for i, value in enumerate(values):
        ax.text(i, value + 0.6, f"{value:.1f}", ha="center", va="bottom", color="#F4F7FB", fontsize=13)
    _style_ax(ax, "Shares bought, per short call", "Time left. Same 1-point rally")
    ax.set_ylim(0, max(values) * 1.22)
    fig.subplots_adjust(left=0.16, right=0.96, top=0.92, bottom=0.16)
    path = CHART_DIR / "hedge.png"
    fig.savefig(path, dpi=160)
    plt.close(fig)
    return path


def chart_delta() -> Path:
    r = 0.04
    strikes = [85 + i * 0.1 for i in range(301)]
    deltas = [call_delta(100, k, 21 / 365, r, 0.25) for k in strikes]
    fig, ax = plt.subplots(figsize=(7.35, 5.20), dpi=160)
    fig.patch.set_facecolor("#0E1218")
    ax.axhline(0.40, color="#E6B54A", linewidth=1.1, linestyle=(0, (4, 3)), zorder=2)
    ax.axhline(0.60, color="#E6B54A", linewidth=1.1, linestyle=(0, (4, 3)), zorder=2)
    ax.text(85.4, 0.625, "Mid Band  0.40 to 0.60", color="#E6B54A", fontsize=11)
    ax.plot(strikes, deltas, color="#8EBEFF", linewidth=2.6, zorder=3)
    offsets = {95: (12, 14), 100: (22, -16), 106: (12, -16)}
    for band in FACTS["bands"]:
        ax.scatter([band["k"]], [band["delta"]], s=42, color="#F4F7FB", zorder=4)
        ax.annotate(
            f"{band['k']:.0f} Strike   {band['delta']:.2f}",
            (band["k"], band["delta"]),
            textcoords="offset points",
            xytext=offsets[band["k"]],
            color="#F4F7FB",
            fontsize=11,
        )
    _style_ax(ax, "Call Delta", "Strike. Spot is 100")
    ax.set_xlim(85, 115)
    ax.set_ylim(0, 1.02)
    fig.subplots_adjust(left=0.10, right=0.97, top=0.93, bottom=0.16)
    path = CHART_DIR / "delta.png"
    fig.savefig(path, dpi=160)
    plt.close(fig)
    return path


def chart_decay() -> Path:
    days = [row[0] for row in FACTS["decay"]]
    single = [row[1] for row in FACTS["decay"]]
    spread = [row[2] for row in FACTS["decay"]]
    fig, ax = plt.subplots(figsize=(7.55, 5.15), dpi=160)
    fig.patch.set_facecolor("#0E1218")
    ax.plot(days, single, color="#8EBEFF", linewidth=2.5, label="Single ATM call")
    ax.plot(days, spread, color="#3DDCB0", linewidth=2.5, label="100 / 105 call spread")
    _style_ax(ax, "Option price, spot held at 100", "Days left")
    ax.invert_xaxis()
    _legend(ax, "upper right")
    fig.subplots_adjust(left=0.10, right=0.97, top=0.92, bottom=0.16)
    path = CHART_DIR / "decay.png"
    fig.savefig(path, dpi=160)
    plt.close(fig)
    return path


def chart_crush() -> Path:
    labels = ["Start\nSpot 100\nIV 50%", "Spot Flat\nSpot 100\nIV 25%", "Spot Up 2\nSpot 102\nIV 25%"]
    values = [FACTS["crush_start"], FACTS["crush_flat"], FACTS["crush_up"]]
    colors = ["#8EBEFF", "#FF7676", "#E6B54A"]
    fig, ax = plt.subplots(figsize=(7.20, 5.15), dpi=160)
    fig.patch.set_facecolor("#0E1218")
    ax.bar(labels, values, color=colors, width=0.66)
    for i, value in enumerate(values):
        ax.text(i, value + 0.06, money(value), ha="center", va="bottom", color="#F4F7FB", fontsize=14)
    _style_ax(ax, "Call Price", "14-day ATM call. Spot starts at 100")
    ax.set_ylim(0, max(values) * 1.22)
    fig.subplots_adjust(left=0.14, right=0.97, top=0.92, bottom=0.24)
    path = CHART_DIR / "crush.png"
    fig.savefig(path, dpi=160)
    plt.close(fig)
    return path


def chart_payoff() -> Path:
    debit_call = FACTS["call_25"]
    debit_spread = FACTS["spread_25"]
    spots = [90 + i * 0.25 for i in range(101)]
    single = [max(s - 100, 0) - debit_call for s in spots]
    spread = [min(max(s - 100, 0), 5) - debit_spread for s in spots]
    fig, ax = plt.subplots(figsize=(7.45, 5.20), dpi=160)
    fig.patch.set_facecolor("#0E1218")
    ax.axhline(0, color="#A8B3C2", linewidth=1.0)
    ax.plot(spots, single, color="#8EBEFF", linewidth=2.5, label="Long 100 call")
    ax.plot(spots, spread, color="#3DDCB0", linewidth=2.5, label="Long 100 / short 105")
    _style_ax(ax, "Profit at expiry, per share", "Spot at expiry")
    _legend(ax, "upper left")
    ax.set_xlim(90, 115)
    ax.set_ylim(-3.6, 14.2)
    fig.subplots_adjust(left=0.12, right=0.97, top=0.92, bottom=0.16)
    path = CHART_DIR / "payoff.png"
    fig.savefig(path, dpi=160)
    plt.close(fig)
    return path


def build_charts() -> dict[str, Path]:
    return {
        "gamma": chart_gamma(),
        "hedge": chart_hedge_shares(),
        "delta": chart_delta(),
        "decay": chart_decay(),
        "crush": chart_crush(),
        "payoff": chart_payoff(),
    }


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Emu(12192000)
        self.prs.slide_height = Emu(6858000)
        self.prs.core_properties.title = "Gamma Swing Trading"
        self.prs.core_properties.subject = "Dealer gamma, walls, and candle-close swings"
        self.prs.core_properties.author = "Gamma Swing"
        self.blank = self._blank()
        self.n = 0
        self.titles: list[str] = []
        self.records: list[dict] = []
        self.warnings: list[str] = []
        self._fonts = {
            False: ImageFont.truetype(FONT_REG, 10),
            True: ImageFont.truetype(FONT_BOLD, 10),
        }

    def _blank(self):
        for layout in self.prs.slide_layouts:
            if layout.name.lower() == "blank":
                return layout
        return self.prs.slide_layouts[-1]

    def _font(self, size: float, bold: bool) -> ImageFont.FreeTypeFont:
        key = (round(size, 2), bold)
        cached = getattr(self, "_font_cache", None)
        if cached is None:
            self._font_cache = {}
            cached = self._font_cache
        if key not in cached:
            path = FONT_BOLD if bold else FONT_REG
            cached[key] = ImageFont.truetype(path, size)
        return cached[key]

    def _audit(self, text: str, x, y, w, h, size, bold, ml, mr, mt, mb, spacing) -> None:
        check_ascii(text)
        usable_w = (w - ml - mr) * 72
        font = self._font(size, bold)
        lines = 0
        for para in text.split("\n"):
            if para == "":
                lines += 1
                continue
            words = para.split(" ")
            cur = ""
            lines += 1
            for word in words:
                trial = word if cur == "" else cur + " " + word
                width = font.getlength(trial)
                if width <= usable_w:
                    cur = trial
                    continue
                if cur == "":
                    self.warnings.append(
                        f"slide {self.n:02d} word wider than box ({word!r} {width:.0f}>{usable_w:.0f}) {text[:80]!r}"
                    )
                    cur = word
                    continue
                lines += 1
                if font.getlength(word) > usable_w:
                    self.warnings.append(
                        f"slide {self.n:02d} word wider than box ({word!r}) {text[:80]!r}"
                    )
                cur = word
        line_h = size * spacing
        need = lines * line_h / 72 + mt + mb
        if need > h + 0.02:
            self.warnings.append(
                f"slide {self.n:02d} text height {need:.2f} > box {h:.2f} size {size} lines {lines}: {text[:90]!r}"
            )
        if y + h > 7.12 and y < 7.05:
            self.warnings.append(f"slide {self.n:02d} text runs into footer y {y:.2f} h {h:.2f}: {text[:60]!r}")

    def _apply_run(self, run, text: str, size: float, bold: bool, color: str) -> None:
        check_ascii(text)
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = False
        run.font.color.rgb = rgb(color)
        r_pr = run._r.get_or_add_rPr()
        r_pr.set("lang", "en-US")
        for tag in ("a:latin", "a:ea", "a:cs"):
            el = r_pr.find(qn(tag))
            if el is None:
                el = r_pr.makeelement(qn(tag), {})
                r_pr.append(el)
            el.set("typeface", FONT)

    def _style_paragraph(self, paragraph, align: str, size: float, spacing: float) -> None:
        paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = Pt(size * spacing)

    def text(
        self,
        slide,
        x,
        y,
        w,
        h,
        content,
        size=16,
        bold=False,
        color=BODY,
        align="left",
        anchor="top",
        spacing=1.18,
        ml=0.04,
        mr=0.04,
        mt=0.02,
        mb=0.02,
    ):
        if content is None or content == "":
            return None
        self._audit(content, x, y, w, h, size, bold, ml, mr, mt, mb, spacing)
        self.records.append({"slide": self.n, "text": content})
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.auto_size = None
        frame.margin_left = Inches(ml)
        frame.margin_right = Inches(mr)
        frame.margin_top = Inches(mt)
        frame.margin_bottom = Inches(mb)
        frame.anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        lines = content.split("\n")
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            self._style_paragraph(paragraph, align, size, spacing)
            run = paragraph.add_run()
            self._apply_run(run, line, size, bold, color)
        if os.environ.get("DEBUG_BOXES") == "1":
            self.rect(slide, x, y, w, h, None, line="FF00AA", line_w=0.75)
        return box

    def rich(self, slide, x, y, w, h, paragraphs, align="left", anchor="top", ml=0.06, mr=0.06, mt=0.04, mb=0.04):
        """paragraphs: list of list of (text, size, bold, color)."""
        plain = "\n".join("".join(run[0] for run in para) for para in paragraphs)
        max_size = max(run[1] for para in paragraphs for run in para)
        any_bold = any(run[2] for para in paragraphs for run in para)
        self._audit(plain, x, y, w, h, max_size, any_bold, ml, mr, mt, mb, 1.18)
        self.records.append({"slide": self.n, "text": plain})
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.auto_size = None
        frame.margin_left = Inches(ml)
        frame.margin_right = Inches(mr)
        frame.margin_top = Inches(mt)
        frame.margin_bottom = Inches(mb)
        frame.anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        for i, runs in enumerate(paragraphs):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            size = max(run[1] for run in runs)
            self._style_paragraph(paragraph, align, size, 1.18)
            for text, rsize, bold, color in runs:
                run = paragraph.add_run()
                self._apply_run(run, text, rsize, bold, color)
        return box

    def rect(self, slide, x, y, w, h, fill, line=None, line_w=1.0):
        if x < -0.02 or y < -0.02 or x + w > W + 0.02 or y + h > H + 0.02:
            raise SystemExit(f"Shape off slide {self.n}: {x:.2f},{y:.2f} {w:.2f}x{h:.2f}")
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.shadow.inherit = False
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(line_w)
        return shape

    def round(self, slide, x, y, w, h, fill, line=None, line_w=1.0, adj=0.15):
        if x < -0.02 or y < -0.02 or x + w > W + 0.02 or y + h > H + 0.02:
            raise SystemExit(f"Shape off slide {self.n}: {x:.2f},{y:.2f} {w:.2f}x{h:.2f}")
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.shadow.inherit = False
        try:
            shape.adjustments[0] = adj
        except Exception:
            pass
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(line_w)
        return shape

    def oval(self, slide, x, y, w, h, fill, line=None, line_w=1.0):
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.shadow.inherit = False
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(line_w)
        return shape

    def arrow(self, slide, x, y, w, h, fill, direction="right"):
        kind = {
            "right": MSO_SHAPE.RIGHT_ARROW,
            "left": MSO_SHAPE.LEFT_ARROW,
            "up": MSO_SHAPE.UP_ARROW,
            "down": MSO_SHAPE.DOWN_ARROW,
        }[direction]
        shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.shadow.inherit = False
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.fill.background()
        return shape

    def panel(self, slide, x, y, w, h, accent=None, fill=CARD):
        self.rect(slide, x, y, w, h, fill, line=LINE, line_w=1.0)
        if accent:
            self.rect(slide, x, y, w, 0.055, accent)

    def chip(self, slide, x, y, w, h, label, fill, fg, size=12):
        self.round(slide, x, y, w, h, fill, adj=0.5)
        self.text(slide, x, y, w, h, label, size=size, bold=True, color=fg, align="center", anchor="middle", ml=0.06, mr=0.06, mt=0.0, mb=0.0)

    def hline(self, slide, x, y, w, color, thickness=0.018):
        self.rect(slide, x, y, w, thickness, color)

    def vline(self, slide, x, y, h, color, thickness=0.018):
        self.rect(slide, x, y, thickness, h, color)

    def footer(self, slide, cover=False):
        self.text(slide, 0.42, 7.16, 8.2, 0.26, "Gamma Swing", size=12, color=FOOT, anchor="middle", ml=0.0, mr=0.0, mt=0.0, mb=0.0)
        self.text(
            slide,
            10.4,
            7.16,
            2.5,
            0.26,
            f"{self.n:02d}",
            size=12,
            color=FOOT,
            align="right",
            anchor="middle",
            ml=0.0,
            mr=0.0,
            mt=0.0,
            mb=0.0,
        )

    def chrome(self, kicker: str, title: str):
        kicker = title_case(kicker)
        title = title_case(title)
        self.n += 1
        self.titles.append(title)
        slide = self.prs.slides.add_slide(self.blank)
        self.rect(slide, 0, 0, W, H, BG)
        self.rect(slide, 0, 0, W, 0.055, TEAL)
        self.text(slide, ML, 0.14, 12.4, 0.24, kicker, size=12, bold=True, color=AMBER, ml=0.0, mr=0.0, mt=0.0, mb=0.0)
        self.text(slide, ML, 0.36, 12.5, 0.48, title, size=28, bold=True, color=TEXT, ml=0.0, mr=0.0, mt=0.0, mb=0.0, spacing=1.05)
        self.rect(slide, ML, 0.90, 0.50, 0.028, TEAL)
        self.rect(slide, ML + 0.50, 0.906, CONTENT_W - 0.50, 0.012, LINE)
        self.footer(slide)
        return slide

    def dashed_h(self, slide, x, y, w, color, thickness=0.016, dash=0.14, gap=0.09):
        cursor = x
        end = x + w
        while cursor < end - 0.01:
            seg = min(dash, end - cursor)
            self.rect(slide, cursor, y, seg, thickness, color)
            cursor += dash + gap

    def candle(self, slide, cx, y_high, y_low, y_open, y_close, color, body=0.20):
        top = min(y_high, y_low)
        bot = max(y_high, y_low)
        self.rect(slide, cx - 0.012, top, 0.024, max(0.02, bot - top), color)
        body_top = min(y_open, y_close)
        body_bot = max(y_open, y_close)
        if body_bot - body_top < 0.10:
            body_bot = body_top + 0.10
        self.rect(slide, cx - body / 2, body_top, body, body_bot - body_top, color)

    def stamp(self, slide, x, y, word="ILLUSTRATIVE"):
        width = 1.92 if word == "ILLUSTRATIVE" else 1.22
        self.chip(slide, x, y, width, 0.32, word, AMBER_DIM, AMBER, size=12)

    def picture(self, slide, path: Path, x, y, w):
        image = Image.open(path)
        h = w * image.height / image.width
        if y + h > 7.08:
            self.warnings.append(f"slide {self.n:02d} picture overflows y {y:.2f} h {h:.2f}")
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        return h

    def picture_in(self, slide, path: Path, x, y, w, h):
        """Scale a diagram to fit a panel and center it in that panel."""
        image = Image.open(path)
        aspect = image.width / image.height
        if w / h > aspect:
            draw_h = h
            draw_w = h * aspect
        else:
            draw_w = w
            draw_h = w / aspect
        px = x + (w - draw_w) / 2
        py = y + (h - draw_h) / 2
        slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(draw_w), Inches(draw_h))
        return draw_h

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)


def cols(n, gap=0.18, x=ML, width=CONTENT_W):
    cw = (width - gap * (n - 1)) / n
    return [(x + i * (cw + gap), cw) for i in range(n)]


_SMALL = {
    "a", "an", "the", "and", "or", "but", "nor", "for", "on", "at", "to",
    "from", "by", "of", "in", "with", "vs", "via", "per", "as",
}
_KEEP = {
    "iv": "IV",
    "spy": "SPY",
    "oi": "OI",
    "opex": "OPEX",
    "atm": "ATM",
    "itm": "ITM",
    "otm": "OTM",
    "0dte": "0DTE",
    "gex": "GEX",
    "vix": "VIX",
    "spx": "SPX",
}


def _cap_token(token: str, force: bool) -> str:
    if token == "":
        return token
    lead = ""
    core = token
    tail = ""
    while core and core[0] in "\"'(":
        lead += core[0]
        core = core[1:]
    while core and core[-1] in ".,:;\"')":
        tail = core[-1] + tail
        core = core[:-1]
    if core == "":
        return token
    if any(ch.isdigit() for ch in core):
        shown = core
    elif "-" in core and core.lower() not in _KEEP:
        parts = core.split("-")
        shown = "-".join(_cap_token(part, True).strip(".,:;") for part in parts)
    else:
        key = core.lower()
        if key in _KEEP:
            shown = _KEEP[key]
        elif key in _SMALL and not force:
            shown = key
        else:
            shown = key[:1].upper() + key[1:]
    return lead + shown + tail


def title_case(text: str) -> str:
    """Title Case for headings. Leaves numbers, 4h, 15m, IV, SPY, and 0DTE alone."""
    if text is None:
        return text
    lines = []
    for line in text.split("\n"):
        words = line.split(" ")
        useful = [i for i, word in enumerate(words) if word.strip("\"'().,:;") != ""]
        out = []
        previous_break = True
        for i, word in enumerate(words):
            if word.strip("\"'().,:;") == "":
                out.append(word)
                continue
            is_edge = i == useful[0] or i == useful[-1]
            shown = _cap_token(word, force=is_edge or previous_break)
            out.append(shown)
            previous_break = word.endswith((".", ":"))
        lines.append(" ".join(out))
    return "\n".join(lines)


# Slides. ASCII only. No absolute-value bars.


def _wrap_lines(d, text, width_in, size, bold, ml=0.0, mr=0.0):
    font = d._font(size, bold)
    usable = max(12, (width_in - ml - mr) * 72)
    total = 0
    for para in text.split("\n"):
        if para == "":
            total += 1
            continue
        words = para.split(" ")
        cur = ""
        total += 1
        for word in words:
            trial = word if cur == "" else cur + " " + word
            if font.getlength(trial) <= usable:
                cur = trial
            else:
                if cur != "":
                    total += 1
                cur = word
    return max(1, total)


def bands(y, h, n, gap=0.14):
    """Equal row bands inside a shared vertical span."""
    bh = (h - gap * (n - 1)) / n
    return [(y + i * (bh + gap), bh) for i in range(n)]


def in_box(
    d, s, x, y, w, h, content,
    size=16, bold=False, color=BODY, align="center",
    ml=0.08, mr=0.08, lo=1.12, hi=1.46,
):
    """The one vertical-alignment helper.

    Text sits in the middle of the box. Line spacing opens just enough to
    fill a short box, and stops before the lines turn into stripes of empty air.
    """
    if content is None or content == "":
        return
    pad = 0.06
    inner_h = max(0.2, h - pad * 2)
    use = size
    lines = _wrap_lines(d, content, w, use, bold, ml, mr)
    fit = (inner_h * 72) / max(1, lines * use)
    while use > 12 and fit < 1.05:
        use -= 1
        lines = _wrap_lines(d, content, w, use, bold, ml, mr)
        fit = (inner_h * 72) / max(1, lines * use)
    spacing = max(lo, min(fit, hi))
    used = lines * use * spacing / 72
    if used > inner_h:
        spacing = max(1.0, (inner_h * 72) / max(1, lines * use))
        used = lines * use * spacing / 72
    top = y + (h - used) / 2
    d.text(
        s, x, top, w, used + 0.04, content,
        size=use, bold=bold, color=color, align=align, anchor="top",
        spacing=spacing, ml=ml, mr=mr, mt=0, mb=0,
    )


def stack_box(d, s, x, y, w, h, parts, align="center", gap=0.08, pad_x=0.14):
    """Center a group of text runs inside one box. parts: (text, size, bold, color)."""
    inner_w = w - pad_x * 2
    heights = []
    for text, size, bold, _color in parts:
        lines = _wrap_lines(d, text, inner_w, size, bold)
        heights.append(lines * size * 1.16 / 72)
    total = sum(heights) + gap * (len(parts) - 1)
    top = y + max(0.04, (h - total) / 2)
    for (text, size, bold, color), ph in zip(parts, heights):
        d.text(
            s, x + pad_x, top, inner_w, ph + 0.03, text,
            size=size, bold=bold, color=color, align=align, anchor="top",
            spacing=1.16, ml=0, mr=0, mt=0, mb=0,
        )
        top += ph + gap


def block(d, s, x, y, w, h, title, body, accent=None, ts=18, bs=16, align="left"):
    """Card. Title stays under the accent. Body uses in_box so the copy is centered in what remains."""
    title = title_case(title)
    pad_x = 0.22
    inner_w = w - pad_x * 2
    if h >= 4.6 and inner_w >= 4.2:
        ts = max(ts, 20)
    elif h >= 3.2:
        ts = max(ts, 18)
    title_lines = _wrap_lines(d, title, inner_w, ts, True)
    title_h = max(0.38, title_lines * ts * 1.08 / 72 + 0.04)
    body_top = 0.14 + title_h + 0.04
    body_h = h - body_top - 0.14
    chosen = bs
    ceiling = 22 if inner_w >= 4.0 else 18
    for candidate in range(ceiling, 13, -1):
        lines = _wrap_lines(d, body, inner_w, candidate, False)
        if lines * candidate * 1.12 / 72 <= body_h:
            chosen = candidate
            break
    d.panel(s, x, y, w, h, accent=accent)
    d.text(
        s, x + pad_x, y + 0.14, inner_w, title_h, title,
        size=ts, bold=True, color=TEXT, align=align, ml=0, mr=0, mt=0, mb=0, spacing=1.05,
    )
    in_box(
        d, s, x + pad_x, y + body_top, inner_w, body_h, body,
        size=chosen, color=BODY, align=align, ml=0, mr=0,
    )


def ymap(price, pmin, pmax, top, bot):
    return bot - (price - pmin) / (pmax - pmin) * (bot - top)


def slide_cover(d: Deck) -> None:
    d.n += 1
    d.titles.append("Gamma Swing Trading")
    s = d.prs.slides.add_slide(d.blank)
    d.rect(s, 0, 0, W, H, BG)
    d.rect(s, 0, 0, W, 0.055, TEAL)
    d.text(s, 0.55, 1.35, 7.2, 0.28, "Dealer Gamma", size=14, bold=True, color=AMBER, ml=0, mr=0, mt=0, mb=0)
    d.text(s, 0.52, 1.70, 7.4, 0.92, "Gamma", size=58, bold=True, color=TEXT, ml=0, mr=0, mt=0, mb=0, spacing=1.0)
    d.text(s, 0.52, 2.58, 7.5, 0.66, "Swing Trading", size=36, bold=True, color=TEXT, ml=0, mr=0, mt=0, mb=0, spacing=1.0)
    d.rect(s, 0.55, 3.36, 2.05, 0.032, TEAL)
    d.text(
        s, 0.55, 3.55, 7.15, 1.20,
        "How the hedge pushes price.\nLong premium, short premium, and defined-risk spreads.\nThe walls, the flip, and the limits of the model.",
        size=18, color=BODY, ml=0, mr=0, mt=0, mb=0, spacing=1.18,
    )
    row_w = 7.15
    gap = 0.12
    chip_h = 0.40
    row1 = ["Open Interest", "Call Wall", "Put Wall", "The Flip"]
    cw = (row_w - gap * (len(row1) - 1)) / len(row1)
    for i, label in enumerate(row1):
        d.chip(s, 0.55 + i * (cw + gap), 5.00, cw, chip_h, label, NEST, BODY, size=12)
    row2 = ["Long Gamma", "Short Gamma", "Spreads"]
    cw2 = cw
    row2_w = cw2 * len(row2) + gap * (len(row2) - 1)
    x0 = 0.55 + (row_w - row2_w) / 2
    for i, label in enumerate(row2):
        d.chip(s, x0 + i * (cw2 + gap), 5.52, cw2, chip_h, label, NEST, BODY, size=12)
    d.text(s, 0.55, 6.12, row_w, 0.32, "October 2026", size=14, color=DIM, align="left", ml=0, mr=0, mt=0, mb=0)

    px, py, pw, ph = 8.22, 0.72, 4.66, 6.05
    band = 0.52
    half = (ph - band) / 2
    d.rect(s, px, py, pw, half, TEAL_DIM)
    d.text(s, px, py + 0.16, pw, 0.28, "Positive Gamma", size=14, bold=True, color=TEAL, align="center", ml=0, mr=0, mt=0, mb=0)
    d.text(s, px + 0.3, py + 0.48, pw - 0.6, 0.36, "Sell rallies. Buy dips.", size=16, color=TEXT, align="center", ml=0, mr=0, mt=0, mb=0)
    d.arrow(s, px + 1.35, py + 1.05, 0.34, 0.42, TEAL, "down")
    d.arrow(s, px + 2.85, py + 1.05, 0.34, 0.42, TEAL, "up")
    d.text(s, px + 1.05, py + 1.52, 0.95, 0.26, "Sell", size=13, bold=True, color=TEAL, align="center", ml=0, mr=0, mt=0, mb=0)
    d.text(s, px + 2.55, py + 1.52, 0.95, 0.26, "Buy", size=13, bold=True, color=TEAL, align="center", ml=0, mr=0, mt=0, mb=0)
    d.text(s, px + 0.3, py + half - 0.48, pw - 0.6, 0.32, "Price ranges and pins.", size=15, color=BODY, align="center", ml=0, mr=0, mt=0, mb=0)

    d.rect(s, px, py + half, pw, band, AMBER_DIM)
    d.text(s, px, py + half, pw, band, "The Flip: Behavior Changes Here", size=13, bold=True, color=AMBER, align="center", anchor="middle", ml=0.12, mr=0.12, mt=0, mb=0)

    ny = py + half + band
    d.rect(s, px, ny, pw, half, CORAL_DIM)
    d.text(s, px, ny + 0.16, pw, 0.28, "Negative Gamma", size=14, bold=True, color=CORAL, align="center", ml=0, mr=0, mt=0, mb=0)
    d.text(s, px + 0.3, ny + 0.48, pw - 0.6, 0.36, "Buy rallies. Sell dips.", size=16, color=TEXT, align="center", ml=0, mr=0, mt=0, mb=0)
    d.arrow(s, px + 1.35, ny + 1.05, 0.34, 0.42, CORAL, "up")
    d.arrow(s, px + 2.85, ny + 1.05, 0.34, 0.42, CORAL, "down")
    d.text(s, px + 1.05, ny + 1.52, 0.95, 0.26, "Buy", size=13, bold=True, color=CORAL, align="center", ml=0, mr=0, mt=0, mb=0)
    d.text(s, px + 2.55, ny + 1.52, 0.95, 0.26, "Sell", size=13, bold=True, color=CORAL, align="center", ml=0, mr=0, mt=0, mb=0)
    d.text(s, px + 0.3, ny + half - 0.48, pw - 0.6, 0.32, "Price trends and accelerates.", size=15, color=BODY, align="center", ml=0, mr=0, mt=0, mb=0)
    d.rect(s, px, py, pw, ph, None, line=LINE, line_w=1.25)
    d.footer(s, cover=True)


def slide_map(d: Deck) -> None:
    s = d.chrome("Start here", "Read it as a map")
    items = [
        (TEAL, "What you get", "The dealer hedge, the flip,\nand the walls.\n\nHow long gamma and short\ngamma differ.\n\nWhere the model stops."),
        (AMBER, "What this is not", "No win rate.\nNo backtest.\n\nNo invented track record.\nExample numbers are marked."),
        (CORAL, "What stays true", "Dealers hedge with the stock.\nGamma is how fast delta changes.\n\nA wick and a close\nare different prints."),
    ]
    for (x, w), (accent, title, body) in zip(cols(3), items):
        block(d, s, x, 1.18, w, 3.15, title, body, accent=accent, ts=18, bs=15)
    steps = [("1", "Regime"), ("2", "Walls"), ("3", "Flip"), ("4", "Structure"), ("5", "The Model")]
    xs = cols(5, gap=0.28)
    for i, ((x, w), (num, label)) in enumerate(zip(xs, steps)):
        d.panel(s, x, 4.52, w, 1.15, accent=TEAL)
        stack_box(
            d, s, x, 4.52, w, 1.15,
            [(num, 13, True, TEAL), (label, 16, True, TEXT)],
            gap=0.04,
        )
        if i < 4:
            d.arrow(s, x + w + 0.02, 4.96, 0.24, 0.24, TEAL, "right")
    d.panel(s, ML, 5.86, CONTENT_W, 0.98, accent=BLUE)
    d.text(
        s, ML + 0.22, 6.00, CONTENT_W - 0.4, 0.72,
        "Numbers that are not the October 16, 2026 SPY book are marked EXAMPLE.\nCharts are Black-Scholes. The inputs are printed on the slide.",
        size=15, color=BODY, ml=0, mr=0, mt=0, mb=0,
    )


def slide_tools(d: Deck) -> None:
    s = d.chrome("The read", "What a gamma read uses")
    top, height = 1.16, 5.76
    left_w, gap = 8.10, 0.18
    right_x = ML + left_w + gap
    right_w = CONTENT_W - left_w - gap
    rows = [
        ("Book", "Open interest by strike. That is the swing map."),
        ("Spot", "Where price sits against that book."),
        ("Flip", "Where net gamma changes sign."),
        ("Walls", "The heavy call strike above. The heavy put strike below."),
        ("Check", "A second public source, when you have one."),
        ("Chart", "Price, so a close can be told from a wick."),
    ]
    row_gap = 0.10
    row_h = (height - row_gap * (len(rows) - 1)) / len(rows)
    for i, (label, line) in enumerate(rows):
        y = top + i * (row_h + row_gap)
        d.panel(s, ML, y, left_w, row_h)
        d.rect(s, ML, y, 1.70, row_h, NEST)
        d.text(s, ML, y, 1.70, row_h, label, size=14, bold=True, color=AMBER, align="center", anchor="middle", ml=0.06, mr=0.06, mt=0, mb=0)
        d.text(s, ML + 1.86, y, left_w - 2.02, row_h, line, size=16, color=TEXT, anchor="middle", ml=0.1, mr=0.1, mt=0, mb=0)
    d.panel(s, right_x, top, right_w, height, accent=TEAL)
    mid = top + height / 2
    d.text(s, right_x + 0.2, top + 0.22, right_w - 0.4, 0.34, "Dealer Gamma", size=16, bold=True, color=TEAL, align="center", ml=0, mr=0, mt=0, mb=0)
    d.text(s, right_x + 0.2, top + 0.62, right_w - 0.4, 1.85, "Their hedge\nin the stock.\n\nPositive: they fade.\nNegative: they chase.", size=16, color=TEXT, align="center", anchor="middle", ml=0, mr=0, mt=0, mb=0, spacing=1.2)
    d.hline(s, right_x + 0.28, mid, right_w - 0.56, LINE, 0.015)
    d.text(s, right_x + 0.2, mid + 0.16, right_w - 0.4, 0.34, "Your Gamma", size=16, bold=True, color=AMBER, align="center", ml=0, mr=0, mt=0, mb=0)
    d.text(s, right_x + 0.2, mid + 0.54, right_w - 0.4, 1.90, "Long options: long gamma.\nShort options: short gamma.\n\nThe two books can match,\nor they can oppose.", size=16, color=BODY, align="center", anchor="middle", ml=0.08, mr=0.08, mt=0, mb=0, spacing=1.2)


def slide_core(d: Deck) -> None:
    s = d.chrome("The point", "Gamma shows where the hedge pushes")
    steps = [
        ("1", "Customers trade\noptions"),
        ("2", "Dealers take\nthe other side"),
        ("3", "Dealers hedge\nwith the stock"),
        ("4", "Gamma says how\nhard that pushes"),
    ]
    xs = cols(4, gap=0.36)
    for i, ((x, w), (num, label)) in enumerate(zip(xs, steps)):
        d.panel(s, x, 1.28, w, 1.95, accent=TEAL)
        stack_box(
            d, s, x, 1.28, w, 1.95,
            [(num, 14, True, TEAL), (label, 16, False, TEXT)],
            gap=0.08,
        )
        if i < 3:
            d.arrow(s, x + w + 0.04, 2.10, 0.28, 0.26, TEAL, "right")
    d.panel(s, ML, 3.46, CONTENT_W, 1.05, accent=AMBER)
    d.text(
        s, ML + 0.24, 3.58, CONTENT_W - 0.48, 0.80,
        "Dealers hedge. Gamma shows where that hedging pushes price.",
        size=22, bold=True, color=TEXT, anchor="middle", ml=0, mr=0, mt=0, mb=0,
    )
    notes = [
        (TEAL, "The push is pressure from the hedge. It is not a promise."),
        (AMBER, "A wick through a level and a close beyond it are different prints."),
        (CORAL, "Your long or short gamma is a separate fact from the dealer book."),
    ]
    for (x, w), (accent, line) in zip(cols(3), notes):
        d.panel(s, x, 4.72, w, 1.55)
        d.rect(s, x, 4.72, 0.08, 1.55, accent)
        in_box(d, s, x + 0.22, 4.72, w - 0.40, 1.55, line, size=16, color=BODY, align="left", ml=0.06, mr=0.08)


def slide_dealers(d: Deck) -> None:
    s = d.chrome("Dealers", "Three lines on dealers")
    rows = [
        (
            "Dealers hold large options positions from institutional and retail trading",
            "Their book is the other side of that flow. You read the net book, not one order.",
        ),
        (
            "Dealers use the underlying stock to hedge their options positions",
            "They buy or sell shares against the option. That hedge is in the stock you chart.",
        ),
        (
            "Delta tells them how many shares they need to stay neutral",
            "Delta is the share count. When delta changes, the count is wrong, so they trade again.",
        ),
    ]
    for i, (quote, gloss) in enumerate(rows):
        y = 1.18 + i * 1.88
        d.panel(s, ML, y, CONTENT_W, 1.74, accent=AMBER)
        d.text(s, ML + 0.24, y + 0.18, 1.1, 0.32, f"0{i+1}", size=14, bold=True, color=AMBER, ml=0, mr=0, mt=0, mb=0)
        d.text(s, ML + 0.24, y + 0.50, CONTENT_W - 0.5, 0.55, '"' + quote + '"', size=16, bold=True, color=TEXT, ml=0, mr=0, mt=0, mb=0)
        d.text(s, ML + 0.24, y + 1.10, CONTENT_W - 0.5, 0.46, gloss, size=15, color=BODY, ml=0, mr=0, mt=0, mb=0)


def slide_delta(d: Deck) -> None:
    s = d.chrome("Delta", "Delta is their share count")
    d.panel(s, ML, 1.18, 7.55, 4.05, accent=BLUE)
    d.text(s, 0.64, 1.38, 7.1, 0.70, "One contract stands for 100 shares.\nExposure is delta times those 100 shares.", size=18, color=TEXT, ml=0, mr=0, mt=0, mb=0)
    rows = [("Call Delta 0.20", "About 20 Shares"), ("Call Delta 0.50", "About 50 Shares"), ("Call Delta 0.80", "About 80 Shares")]
    for i, (left, right) in enumerate(rows):
        y = 2.30 + i * 0.85
        d.rect(s, 0.68, y, 3.15, 0.70, NEST)
        d.rect(s, 3.95, y, 3.55, 0.70, NEST)
        d.text(s, 0.68, y, 3.15, 0.70, left, size=16, color=BODY, align="center", anchor="middle", ml=0.08, mr=0.08, mt=0, mb=0)
        d.arrow(s, 3.68, y + 0.20, 0.22, 0.28, BLUE, "right")
        d.text(s, 3.95, y, 3.55, 0.70, right, size=16, bold=True, color=TEXT, align="center", anchor="middle", ml=0.08, mr=0.08, mt=0, mb=0)
    d.panel(s, 8.15, 1.18, 4.76, 4.05, accent=CORAL)
    d.text(s, 8.35, 1.40, 4.36, 0.42, "You Are Not the Dealer", size=18, bold=True, color=TEXT, align="center", ml=0, mr=0, mt=0, mb=0)
    in_box(
        d, s, 8.35, 1.90, 4.36, 3.10,
        "The dealer hedge is in the stock.\nYour position is the option.\n\nOne contract is exposure to 100 shares, times delta.\nThat exposure is not the same as holding the shares.",
        size=17, color=BODY, align="center", ml=0.08, mr=0.08,
    )
    d.panel(s, ML, 5.40, CONTENT_W, 1.42, accent=AMBER)
    d.text(
        s, ML + 0.24, 5.55, CONTENT_W - 0.48, 1.12,
        "A long put with delta -0.50 acts like short 50 shares. You still only own the put.\nYou did not borrow stock. The sign is direction. The number is size. Later you will ignore the sign on purpose.",
        size=16, color=BODY, ml=0, mr=0, mt=0, mb=0,
    )


def slide_gamma_def(d: Deck) -> None:
    s = d.chrome("Gamma", "Gamma is how fast delta changes")
    d.panel(s, ML, 1.18, CONTENT_W, 1.15, accent=AMBER)
    d.text(
        s, ML + 0.24, 1.32, CONTENT_W - 0.48, 0.88,
        "Dealers re-hedge when delta changes. Gamma sets how much, and how often.",
        size=22, bold=True, color=TEXT, anchor="middle", ml=0, mr=0, mt=0, mb=0,
    )
    block(
        d, s, ML, 2.52, 6.15, 2.95, "High gamma",
        "A small move changes delta a lot.\nThe share hedge is rebuilt in size.\nDealers are back in the stock often.\n\nThis is the strike, near expiry.",
        accent=CORAL, ts=18, bs=16,
    )
    block(
        d, s, 6.76, 2.52, 6.15, 2.95, "Low gamma",
        "The same move barely changes delta.\nThe hedge is a small update.\nThey can sit.\n\nThis is far from the strike, or far from expiry.",
        accent=BLUE, ts=18, bs=16,
    )
    d.text(
        s, ML, 5.65, CONTENT_W, 1.15,
        "Calls and puts have the same gamma. The sign of the book is a separate question.\nThe next slide is the shape, from the model. The one after that is the share count.",
        size=16, color=BODY, ml=0, mr=0, mt=0, mb=0,
    )


def slide_gamma_chart(d: Deck) -> None:
    s = d.chrome("Black-Scholes", "Gamma peaks at the strike, near expiry")
    top, height = 1.14, 5.28
    left_w, gap = 7.55, 0.18
    right_x = ML + left_w + gap
    right_w = CONTENT_W - left_w - gap
    d.picture_in(s, d.charts["gamma"], ML, top, left_w, height)
    r1 = f"{FACTS['g_ratio_1']:.1f}"
    r7 = f"{FACTS['g_ratio_7']:.1f}"
    notes = [
        (TEAL, "The peak is at the money. Here, strike 100."),
        (CORAL, f"1-day gamma at the peak is about {r1} times the 30-day peak."),
        (AMBER, f"7-day gamma at the peak is about {r7} times the 30-day peak."),
        (BLUE, "Far from the strike the line is flat. Little re-hedging."),
    ]
    for (y, bh), (accent, line) in zip(bands(top, height, 4, gap=0.12), notes):
        d.panel(s, right_x, y, right_w, bh)
        d.rect(s, right_x, y, 0.08, bh, accent)
        in_box(d, s, right_x + 0.20, y, right_w - 0.34, bh, line, size=16, color=BODY, align="left", ml=0.04, mr=0.08)
    d.text(
        s, ML, 6.52, CONTENT_W, 0.48,
        "Scaled so the 30-day peak is 1.\nEXAMPLE. Black-Scholes, European, no dividend. Spot 100, IV 20%, rate 4%. Not a forecast.",
        size=13, color=DIM, ml=0, mr=0, mt=0, mb=0,
    )


def slide_hedge_chart(d: Deck) -> None:
    s = d.chrome("Re-hedge", "One point of price, three hedges")
    top, height = 1.14, 5.28
    (lx, lw), (rx, rw) = cols(2, gap=0.18)
    d.picture_in(s, d.charts["hedge"], lx, top, lw, height)
    sh = FACTS["shares"]
    upper, lower = bands(top, height, 2, gap=0.14)
    block(
        d, s, rx, upper[0], rw, upper[1], "What the Bars Are",
        "EXAMPLE. The dealer is short one ATM call. Strike 100.\n\n"
        "Spot goes from 100 to 101.\n\n"
        "The bars are the shares they must buy to get back to neutral.\n\n"
        + f"{sh[0]:.1f}, {sh[1]:.1f}, and {sh[2]:.1f} shares.",
        accent=CORAL, ts=18, bs=16,
    )
    d.panel(s, rx, lower[0], rw, lower[1], accent=TEAL)
    in_box(
        d, s, rx + 0.22, lower[0], rw - 0.44, lower[1],
        "Flip the book. If the dealer is long that call, the same sizes are shares they sell. Same gamma. Opposite trade.",
        size=16, color=BODY, align="left", ml=0.04, mr=0.06,
    )
    d.text(
        s, ML, 6.52, CONTENT_W, 0.42,
        "EXAMPLE. Black-Scholes, European, no dividend. IV 20%, rate 4%. Share counts match the bars. Not a forecast.",
        size=13, color=DIM, ml=0, mr=0, mt=0, mb=0,
    )


def _regime_panel(d, s, x, y, w, h, fill, accent, title, lead, left_word, right_word, left_dir, right_dir, foot):
    d.rect(s, x, y, w, h, fill, line=LINE)
    head, mid, tail = bands(y + 0.16, h - 0.32, 3, gap=0.08)
    d.text(s, x, head[0], w, 0.34, title, size=15, bold=True, color=accent, align="center", ml=0, mr=0, mt=0, mb=0)
    in_box(d, s, x + 0.28, head[0] + 0.36, w - 0.56, head[1] - 0.40, lead, size=18, color=TEXT, align="center", ml=0.04, mr=0.04)
    pair = 2.4
    ax = x + (w - pair) / 2
    d.arrow(s, ax + 0.15, mid[0] + 0.08, 0.42, 0.52, accent, left_dir)
    d.arrow(s, ax + 1.55, mid[0] + 0.08, 0.42, 0.52, accent, right_dir)
    d.text(s, ax, mid[0] + mid[1] - 0.36, 0.85, 0.30, left_word, size=14, bold=True, color=accent, align="center", ml=0, mr=0, mt=0, mb=0)
    d.text(s, ax + 1.35, mid[0] + mid[1] - 0.36, 0.95, 0.30, right_word, size=14, bold=True, color=accent, align="center", ml=0, mr=0, mt=0, mb=0)
    in_box(d, s, x + 0.28, tail[0], w - 0.56, tail[1], foot, size=16, color=BODY, align="center", ml=0.04, mr=0.04)


def slide_regimes(d: Deck) -> None:
    s = d.chrome("Regimes", "Two ways the hedge can push")
    top, height = 1.16, 4.62
    (lx, lw), (rx, rw) = cols(2)
    _regime_panel(
        d, s, lx, top, lw, height, TEAL_DIM, TEAL,
        "Positive Gamma", "Above the flip.\nDealers sell rallies and buy dips.",
        "Sell", "Buy", "down", "up",
        "Price ranges, and it can pin.\nPin means price sticks near a heavy strike.",
    )
    _regime_panel(
        d, s, rx, top, rw, height, CORAL_DIM, CORAL,
        "Negative Gamma", "Below the flip.\nDealers chase the move.",
        "Buy", "Sell", "up", "down",
        "Price can trend, and it can accelerate.\nFades get run over.",
    )
    d.panel(s, ML, 5.94, CONTENT_W, 0.92, accent=AMBER)
    in_box(
        d, s, ML + 0.22, 5.94, CONTENT_W - 0.44, 0.92,
        "This deck reads dealer gamma. Above the flip is positive. Below the flip is negative.",
        size=16, color=TEXT, align="center", ml=0.08, mr=0.08,
    )


def slide_positive(d: Deck) -> None:
    s = d.chrome("Positive gamma", "They sell the rally and buy the dip")
    top, height = 1.16, 5.70
    (lx, lw), (rx, rw) = cols(2, gap=0.18, width=CONTENT_W)
    # keep the diagram wider than the note: 8.05 / 4.26 on the same edges
    lw, rw = 8.05, CONTENT_W - 8.05 - 0.18
    rx = ML + lw + 0.18
    d.panel(s, ML, top, lw, height)
    d.stamp(s, ML + 0.18, top + 0.16, "ILLUSTRATIVE")
    line_x = ML + 0.85
    line_w = lw - 1.70
    call_y = top + 1.20
    put_y = top + height - 1.15
    d.hline(s, line_x, call_y, line_w, CYAN, 0.03)
    d.hline(s, line_x, put_y, line_w, PURPLE, 0.03)
    d.text(s, line_x, call_y - 0.40, 2.4, 0.32, "Call Wall", size=14, bold=True, color=CYAN, ml=0, mr=0, mt=0, mb=0)
    d.text(s, line_x, put_y + 0.16, 2.4, 0.32, "Put Wall", size=14, bold=True, color=PURPLE, ml=0, mr=0, mt=0, mb=0)
    mid = (call_y + put_y) / 2
    ax = line_x + line_w * 0.42
    d.arrow(s, ax, call_y + 0.18, 0.42, 0.52, TEAL, "down")
    in_box(d, s, ax + 0.55, call_y + 0.08, 3.15, 0.85, "Rally into the call wall.\nDealers sell stock.", size=15, color=BODY, align="left", ml=0, mr=0)
    d.arrow(s, ax, put_y - 0.78, 0.42, 0.52, TEAL, "up")
    in_box(d, s, ax + 0.55, put_y - 0.88, 3.15, 0.85, "Dip into the put wall.\nDealers buy stock.", size=15, color=BODY, align="left", ml=0, mr=0)
    in_box(d, s, line_x, mid - 0.45, 2.15, 0.90, "Price stays\nin the range.", size=15, color=TEXT, align="center", ml=0, mr=0)
    for i, yy in enumerate((mid - 0.55, mid - 0.18, mid + 0.18, mid + 0.55)):
        xx = line_x + 2.35 + (i % 2) * 0.38
        d.oval(s, xx, yy, 0.14, 0.14, TEAL)
    block(
        d, s, rx, top, rw, height, "Two different books",
        "Price tends to range, and it can pin near a heavy strike.\n\nA long option is long gamma. A credit spread, condor, or butterfly is short gamma.\n\nThose structures are later slides. They are not the dealer hedge.",
        accent=TEAL, ts=18, bs=16, align="center",
    )


def slide_negative(d: Deck) -> None:
    s = d.chrome("Negative gamma", "They chase, so the move can run")
    top, height = 1.16, 5.70
    lw = 8.05
    rw = CONTENT_W - lw - 0.18
    rx = ML + lw + 0.18
    d.panel(s, ML, top, lw, height)
    d.stamp(s, ML + 0.18, top + 0.16, "ILLUSTRATIVE")
    line_x = ML + 0.85
    line_w = lw - 1.70
    call_y = top + 1.35
    put_y = top + height - 1.15
    d.hline(s, line_x, call_y, line_w, CYAN, 0.03)
    d.hline(s, line_x, put_y, line_w, PURPLE, 0.03)
    d.text(s, line_x, call_y - 0.38, 2.4, 0.32, "Call Wall", size=14, bold=True, color=CYAN, ml=0, mr=0, mt=0, mb=0)
    d.text(s, line_x, put_y + 0.10, 2.4, 0.32, "Put Wall", size=14, bold=True, color=PURPLE, ml=0, mr=0, mt=0, mb=0)
    ax = line_x + line_w * 0.38
    d.arrow(s, ax, top + 0.42, 0.42, 0.52, CORAL, "up")
    in_box(d, s, ax + 0.58, top + 0.28, 3.3, 0.85, "Rally. Dealers buy.\nThe rally gets fuel.", size=15, color=BODY, align="left", ml=0, mr=0)
    d.arrow(s, ax, put_y + 0.22, 0.42, 0.52, CORAL, "down")
    in_box(d, s, ax + 0.58, put_y + 0.16, 3.3, 0.85, "Drop. Dealers sell.\nThe drop gets fuel.", size=15, color=BODY, align="left", ml=0, mr=0)
    in_box(d, s, line_x, (call_y + put_y) / 2 - 0.45, 2.4, 0.90, "A break can\nrun.", size=16, color=TEXT, align="center", ml=0, mr=0)
    block(
        d, s, rx, top, rw, height, "If the move runs",
        "The hedge can add fuel, so a move can extend.\n\nLong gamma gains delta as the move runs. Short gamma loses as it runs.\n\nA naked short has that short gamma, and the loss is not capped.",
        accent=CORAL, ts=18, bs=16, align="center",
    )


def slide_hedge_sign(d: Deck) -> None:
    s = d.chrome("Example", "Same calls, opposite hedge")
    d.text(
        s, ML, 1.12, CONTENT_W, 0.55,
        "EXAMPLE. 10 contracts. Deltas are chosen so the share count is obvious. This is one call book, not every book.",
        size=15, color=DIM, ml=0, mr=0, mt=0, mb=0,
    )
    (lx, lw), (rx, rw) = cols(2)
    card_y, card_h = 1.72, 5.05
    d.panel(s, lx, card_y, lw, card_h, accent=TEAL)
    d.text(s, lx, card_y + 0.16, lw, 0.70, "Positive Book\nDealer Is Long the Calls", size=18, bold=True, color=TEAL, align="center", ml=0.2, mr=0.2, mt=0, mb=0)
    in_box(
        d, s, lx + 0.28, card_y + 0.95, lw - 0.56, card_h - 1.15,
        "Start. Delta +0.50.\nHedge: short 500 shares.\n\nLater. Delta +0.60.\nHedge: short 600 shares.\n\nThey sell 100 shares\ninto the rally.",
        size=18, color=BODY, align="center", ml=0.06, mr=0.06,
    )
    d.panel(s, rx, card_y, rw, card_h, accent=CORAL)
    d.text(s, rx, card_y + 0.16, rw, 0.70, "Negative Book\nDealer Is Short the Calls", size=18, bold=True, color=CORAL, align="center", ml=0.2, mr=0.2, mt=0, mb=0)
    in_box(
        d, s, rx + 0.28, card_y + 0.95, rw - 0.56, card_h - 1.15,
        "Start. Delta -0.50.\nHedge: long 500 shares.\n\nLater. Delta -0.60.\nHedge: long 600 shares.\n\nThey buy 100 shares\ninto the rally.",
        size=18, color=BODY, align="center", ml=0.06, mr=0.06,
    )


def slide_flip(d: Deck) -> None:
    s = d.chrome("The flip", "The flip is a change in behavior")
    top, height = 1.16, 5.70
    lw, rw = 7.35, CONTENT_W - 7.35 - 0.18
    rx = ML + lw + 0.18
    d.panel(s, ML, top, lw, height)
    d.text(s, ML, top + 0.12, lw, 0.28, "Price", size=13, bold=True, color=DIM, align="center", ml=0, mr=0, mt=0, mb=0)
    inner_top = top + 0.50
    inner_h = height - 0.70
    box_h = (inner_h - 0.70) / 2
    box_w = 4.55
    box_x = ML + 0.45
    d.rect(s, box_x, inner_top, box_w, box_h, TEAL_DIM)
    in_box(d, s, box_x, inner_top, box_w - 0.7, box_h, "Above the flip\nPositive gamma. Fade.", size=16, color=TEXT, align="center", ml=0.12, mr=0.08)
    d.arrow(s, box_x + box_w - 0.55, inner_top + box_h * 0.22, 0.36, 0.40, TEAL, "down")
    d.arrow(s, box_x + box_w - 0.55, inner_top + box_h * 0.58, 0.36, 0.40, TEAL, "up")
    flip_y = inner_top + box_h + 0.16
    d.rect(s, ML + 0.28, flip_y, lw - 0.56, 0.08, AMBER)
    d.text(s, ML + 0.28, flip_y + 0.12, lw - 0.56, 0.32, "Flip: Net Gamma Changes Sign", size=14, bold=True, color=AMBER, align="center", ml=0, mr=0, mt=0, mb=0)
    low_y = flip_y + 0.50
    d.rect(s, box_x, low_y, box_w, box_h, CORAL_DIM)
    in_box(d, s, box_x, low_y, box_w - 0.7, box_h, "Below the flip\nNegative gamma. Chase.", size=16, color=TEXT, align="center", ml=0.12, mr=0.08)
    d.arrow(s, box_x + box_w - 0.55, low_y + box_h * 0.22, 0.36, 0.40, CORAL, "up")
    d.arrow(s, box_x + box_w - 0.55, low_y + box_h * 0.58, 0.36, 0.40, CORAL, "down")
    upper, lower = bands(top, height, 2, gap=0.16)
    block(
        d, s, rx, upper[0], rw, upper[1], "Not a Level to Lean On",
        "The flip marks a change in behavior.\n\nSupport and resistance are the walls.",
        accent=AMBER, ts=16, bs=16, align="center",
    )
    block(
        d, s, rx, lower[0], rw, lower[1], "It Can Sit on a Wall",
        "The flip can fall between two strikes.\n\nIf it sits on a wall, that edge is the regime line and the level, together.",
        accent=BLUE, ts=16, bs=16, align="center",
    )


def slide_flip_quiet(d: Deck) -> None:
    s = d.chrome("The flip", "On the flip, the hedge goes quiet")
    d.panel(s, ML, 1.18, CONTENT_W, 2.55, accent=AMBER)
    d.text(
        s, ML + 0.28, 1.40, CONTENT_W - 0.55, 2.10,
        "There is near-zero hedging on the flip itself.\n\nPrice often consolidates there, until it picks a side.\nA tight chop on the line is not a signal. It is the wait.",
        size=20, color=TEXT, ml=0, mr=0, mt=0, mb=0,
    )
    steps = [
        (AMBER, "On the line", "Little stock to buy or sell for the gamma hedge. Price often sits."),
        (TEAL, "Accepts above", "A close holds above the flip. Expect fade behavior. Read it as positive."),
        (CORAL, "Accepts below", "A close holds below the flip. Expect chase behavior. Read it as negative."),
    ]
    for (x, w), (accent, title, body) in zip(cols(3), steps):
        block(d, s, x, 3.95, w, 2.75, title, body, accent=accent, ts=18, bs=16)


def _candle_case(d, s, x, y, w, h, tone, title, note, high, low, open_, close):
    d.panel(s, x, y, w, h)
    d.stamp(s, x + 0.16, y + 0.14, "ILLUSTRATIVE")
    flip_y = y + h * 0.46
    d.hline(s, x + 0.30, flip_y, w * 0.40, AMBER, 0.028)
    d.text(s, x + 0.36, flip_y - 0.34, 1.2, 0.28, "Flip", size=13, bold=True, color=AMBER, ml=0, mr=0, mt=0, mb=0)
    cx = x + w * 0.32
    d.candle(s, cx, y + high, y + low, y + open_, y + close, tone, body=0.28)
    in_box(d, s, x + w * 0.50, y + 0.58, w * 0.46, h - 1.25, note, size=16, color=BODY, align="left", ml=0.08, mr=0.1)
    d.text(s, x + 0.2, y + h - 0.58, w - 0.4, 0.40, title, size=16, bold=True, color=TEXT, align="center", ml=0, mr=0, mt=0, mb=0)


def slide_flip_close(d: Deck) -> None:
    s = d.chrome("The flip", "A wick is not a new side")
    top, height = 1.16, 4.40
    (lx, lw), (rx, rw) = cols(2)
    _candle_case(
        d, s, lx, top, lw, height, DIM, "Still the Old Regime.",
        "Wick through.\nClose back on\nthe old side.",
        0.70, 3.55, 2.55, 2.85,
    )
    _candle_case(
        d, s, rx, top, rw, height, TEAL, "Possible New Side.",
        "The close is through\nthe flip. It still\nhas to hold.",
        0.55, 3.35, 1.15, 1.45,
    )
    d.panel(s, ML, 5.72, CONTENT_W, 1.12, accent=BLUE)
    in_box(
        d, s, ML + 0.24, 5.72, CONTENT_W - 0.48, 1.12,
        "A wick through the line is not a held close. A close through still has to hold on a later bar.\nIf later candles hold on the new side, the regime may have changed. Read the book again.",
        size=16, color=BODY, align="left", ml=0.04, mr=0.06,
    )


def slide_walls(d: Deck) -> None:
    s = d.chrome("Levels", "The walls are support and resistance")
    d.panel(s, ML, 1.18, 7.45, 5.55)
    d.stamp(s, 0.60, 1.34, "ILLUSTRATIVE")
    # schematic ladder, not to a real ticker
    levels = [
        (1.95, CYAN, "115    Call Wall    Resistance"),
        (3.15, AMBER, "108    Flip    Behavior, Not a Wall"),
        (4.35, TEXT, "100    Spot"),
        (5.55, PURPLE, "90    Put Wall    Support"),
    ]
    d.vline(s, 1.55, 1.70, 4.55, LINE, 0.02)
    for y, color, label in levels:
        if color == AMBER:
            d.dashed_h(s, 1.15, y, 3.3, color, 0.02, 0.12, 0.08)
        else:
            d.hline(s, 1.15, y, 3.3, color, 0.028)
        d.oval(s, 1.47, y - 0.07, 0.16, 0.16, color)
        d.text(s, 4.55, y - 0.16, 3.05, 0.40, label, size=14, color=color, anchor="middle", ml=0, mr=0, mt=0, mb=0)
    block(
        d, s, 8.05, 1.18, 4.86, 5.55, "How to use them",
        "Call wall: the heavy call strike above. Often resistance.\n\nPut wall: the heavy put strike below. Often support.\n\nThey are open-interest levels.\n\nThe flip is drawn here so it is not read as a third wall.\n\nThe dealer regime says how the hedge behaves at the level.",
        accent=CYAN, ts=18, bs=15,
    )


def slide_walls_regime(d: Deck) -> None:
    s = d.chrome("Levels", "The hedge at a wall depends on the regime")
    block(
        d, s, ML, 1.18, 6.15, 5.55, "Positive gamma at a wall",
        "The hedge leans back into the range.\n\nA tag of the wall and a close back inside are the same side of the flip.\n\nA close that holds outside can be a different regime.\n\nThe wall is still just the heavy strike.",
        accent=TEAL, ts=20, bs=18, align="center",
    )
    block(
        d, s, 6.76, 1.18, 6.15, 5.55, "Negative gamma at a wall",
        "Once price goes through, the hedge can chase.\n\nA wick through the wall is not a held close.\n\nA close that stays outside is a move the hedge may fuel.\n\nThe wall does not stop being a strike just because price tagged it.",
        accent=CORAL, ts=20, bs=18, align="center",
    )


def slide_opex(d: Deck) -> None:
    s = d.chrome("Levels", "Walls move. Big expiries reset them")
    d.panel(s, ML, 1.18, 6.15, 5.55, accent=BLUE)
    d.text(s, 0.64, 1.40, 5.7, 0.40, "The Wall Is This Book", size=20, bold=True, color=TEXT, ml=0, mr=0, mt=0, mb=0)
    # moving wall diagram
    d.text(s, 0.70, 2.05, 2.2, 0.30, "Earlier", size=13, bold=True, color=DIM, ml=0, mr=0, mt=0, mb=0)
    d.hline(s, 0.70, 2.55, 2.6, CYAN, 0.03)
    d.text(s, 3.45, 2.35, 2.2, 0.36, "Call Wall", size=14, bold=True, color=CYAN, ml=0, mr=0, mt=0, mb=0)
    d.arrow(s, 2.70, 3.15, 0.36, 0.46, AMBER, "down")
    d.text(s, 3.20, 3.20, 2.6, 0.40, "Open Interest Shifts", size=14, bold=True, color=AMBER, ml=0, mr=0, mt=0, mb=0)
    d.text(s, 0.70, 3.90, 2.2, 0.30, "Later", size=13, bold=True, color=DIM, ml=0, mr=0, mt=0, mb=0)
    d.hline(s, 0.70, 4.45, 2.6, CYAN, 0.03)
    d.text(s, 3.45, 4.25, 2.4, 0.36, "New Call Wall", size=14, bold=True, color=CYAN, ml=0, mr=0, mt=0, mb=0)
    in_box(d, s, 0.70, 5.00, 5.5, 1.45, "Read them again.\nYesterday's line can be dead.", size=18, color=TEXT, align="left", ml=0, mr=0)
    block(
        d, s, 6.76, 1.18, 6.15, 5.55, "OPEX",
        "Open interest is what\nbuilds the wall.\n\nWhen that expiry passes,\nthose contracts are gone.\n\nThe monthly expiry is\nthe big reset.\n\nSPY also expires through\nthe week. After a large\none, rebuild the map\nbefore you trust the\nold levels.",
        accent=AMBER, ts=20, bs=18,
    )


def slide_charm(d: Deck) -> None:
    s = d.chrome("Textbook only", "Charm and vanna, in one pass")
    block(
        d, s, ML, 1.18, 6.15, 4.15, "Charm",
        "The change in delta as time\npasses, with price and IV\nheld still.\n\nAlso called delta decay.\n\nA hedge can drift on a flat\nday because the clock moved.",
        accent=BLUE, ts=22, bs=18,
    )
    block(
        d, s, 6.76, 1.18, 6.15, 4.15, "Vanna",
        "The change in delta when IV\nchanges, with price held still.\n\nIt is also the change in vega\nwhen price changes.\n\nA hedge can shift on a\nvolatility move, even if the\nstock has not.",
        accent=PURPLE, ts=22, bs=18,
    )
    d.panel(s, ML, 5.52, CONTENT_W, 1.28, accent=AMBER)
    d.text(
        s, ML + 0.24, 5.70, CONTENT_W - 0.48, 0.95,
        "Textbook definitions. This deck does not trade charm or vanna.\nYou trade gamma, the walls, and the close. These two only explain a hedge that moves on a quiet tape.",
        size=16, color=BODY, ml=0, mr=0, mt=0, mb=0,
    )


def slide_uw_rules(d: Deck) -> None:
    s = d.chrome("Unusual Whales", "Four rules on the gamma page")
    rules = [
        ("01", "Heatmap colors are net gamma.", "Not put gamma. Not call gamma by itself. The combined number."),
        ("02", "By-strike draws two lines.", "The call wall, and spot. It does not draw the put wall. It does not draw the flip."),
        ("03", "You find the put wall.", "It is the longest purple bar below spot."),
        ("04", "You find the flip.", "It is where the heatmap color changes. Not a printed line."),
    ]
    for i, (num, title, body) in enumerate(rules):
        col = i % 2
        row = i // 2
        x = ML + col * 6.34
        y = 1.18 + row * 2.35
        d.panel(s, x, y, 6.15, 2.18, accent=PURPLE)
        d.text(s, x + 0.2, y + 0.16, 1.0, 0.36, num, size=14, bold=True, color=PURPLE, ml=0, mr=0, mt=0, mb=0)
        d.text(s, x + 0.2, y + 0.52, 5.75, 0.45, title_case(title), size=18, bold=True, color=TEXT, ml=0, mr=0, mt=0, mb=0)
        d.text(s, x + 0.2, y + 1.10, 5.75, 0.85, body, size=15, color=BODY, ml=0, mr=0, mt=0, mb=0)
    d.text(
        s, ML, 5.95, CONTENT_W, 0.85,
        "Open interest is the swing. Directionalized Volume is the day. The next drawings are schematics.\nThey are marked ILLUSTRATIVE. They are not screenshots, and the heatmap colors are stand-ins.",
        size=15, color=DIM, ml=0, mr=0, mt=0, mb=0,
    )


def slide_heatmap(d: Deck) -> None:
    s = d.chrome("Heatmap", "The color is net gamma")
    d.stamp(s, ML, 1.12, "ILLUSTRATIVE")
    d.text(s, 2.50, 1.14, 8.5, 0.30, "Stand-in colors. Not the Unusual Whales palette.", size=14, color=AMBER, ml=0, mr=0, mt=0, mb=0)
    headers = ["0DTE", "+2d", "+7d", "+14d", "Month", "+2mo"]
    pos = [800, 790, 780, 770]
    neg = [760, 750, 740, 730]
    gx = 1.35
    header_y = 1.58
    cell_w, cell_h, gap = 0.92, 0.46, 0.06
    for i, label in enumerate(headers):
        d.text(s, gx + i * (cell_w + gap), header_y, cell_w, 0.28, label, size=12, bold=True, color=DIM, align="center", ml=0, mr=0, mt=0, mb=0)
    def draw_row(y, strike, fill, mark):
        d.text(s, 0.42, y, 0.85, cell_h, str(strike), size=13, bold=True, color=DIM, align="right", anchor="middle", ml=0, mr=0.06, mt=0, mb=0)
        for i in range(6):
            x = gx + i * (cell_w + gap)
            d.rect(s, x, y, cell_w, cell_h, fill)
            d.text(s, x, y, cell_w, cell_h, mark, size=13, bold=True, color=TEXT, align="center", anchor="middle", ml=0, mr=0, mt=0, mb=0)
    y = 1.95
    for strike in pos:
        draw_row(y, strike, "1C6B58", "+")
        y += cell_h + gap
    d.rect(s, gx, y + 0.04, 6 * cell_w + 5 * gap, 0.040, AMBER)
    d.text(s, gx, y + 0.14, 5.8, 0.28, "Flip = Where the Color Changes", size=13, bold=True, color=AMBER, ml=0, mr=0, mt=0, mb=0)
    y += 0.50
    for strike in neg:
        draw_row(y, strike, "7A3038", "-")
        y += cell_h + gap
    d.panel(s, 8.55, 1.58, 4.35, 5.05, accent=PURPLE)
    d.text(s, 8.75, 1.78, 3.95, 0.42, "Read It This Way", size=18, bold=True, color=TEXT, align="center", ml=0, mr=0, mt=0, mb=0)
    in_box(
        d, s, 8.75, 2.28, 3.95, 4.10,
        "Plus cells: positive net gamma.\n\nMinus cells: negative net gamma.\n\nThe sign is the lesson. The shade is not live data.\n\nOn the real page, match the color to the legend. Do not memorize a hue.",
        size=16, color=BODY, align="center", ml=0.06, mr=0.06,
    )


def slide_bystrike(d: Deck) -> None:
    s = d.chrome("By strike", "This chart does not draw every level")
    d.stamp(s, ML, 1.12, "ILLUSTRATIVE")
    strikes = [
        (790, 36), (785, 58), (780, 128), (775, 84), (770, 48),
        (765, 40), (760, 160), (755, 72), (750, 108), (745, 30),
    ]
    top = 1.58
    pitch = 0.50
    label_x = 0.42
    bar_x = 1.25
    max_w = 4.70
    scale = max_w / 160
    d.panel(s, 0.42, 1.50, 6.15, 5.25)
    for i, (strike, length) in enumerate(strikes):
        y = top + i * pitch
        d.text(s, label_x, y, 0.72, 0.36, str(strike), size=12, bold=True, color=DIM, align="right", anchor="middle", ml=0, mr=0.04, mt=0, mb=0)
        d.rect(s, bar_x, y + 0.04, length * scale, 0.28, "A78BFA")
    # spot between 765 (index 5) and 760 (index 6)
    y765 = top + 5 * pitch + 0.18
    y760 = top + 6 * pitch + 0.18
    spot_y = (y765 + y760) / 2
    d.hline(s, bar_x, spot_y, max_w, WHITE, 0.022)
    d.text(s, bar_x + max_w - 0.85, spot_y - 0.30, 0.80, 0.26, "Spot", size=12, bold=True, color=WHITE, align="right", ml=0, mr=0, mt=0, mb=0)
    # call wall marker starts after the 775 bar, so the label is not on the bar
    call_y = top + 3 * pitch + 0.18
    bar_end = bar_x + 84 * scale
    d.hline(s, bar_end + 0.10, call_y, max_w - (84 * scale) - 0.10, CYAN, 0.028)
    d.text(s, bar_end + 0.12, call_y - 0.30, 1.45, 0.26, "Call Wall", size=12, bold=True, color=CYAN, ml=0, mr=0, mt=0, mb=0)

    notes = [
        (CYAN, "Drawn for you", "Spot, and the call wall. Those two lines are on the chart."),
        (PURPLE, "You mark the put wall", "Longest purple bar below spot. Here that is 760. It is not drawn as a line."),
        (AMBER, "Ignore the trap", "The long bar at 780 is above spot. Longest bar on the wrong side is not the put wall."),
        (CORAL, "Flip is not here", "Do not look for it on this chart. The color change is on the heatmap."),
    ]
    for i, (accent, title, body) in enumerate(notes):
        y = 1.18 + i * 1.38
        d.panel(s, 6.78, y, 6.13, 1.26)
        d.rect(s, 6.78, y, 0.08, 1.26, accent)
        d.text(s, 7.04, y + 0.08, 5.7, 0.32, title_case(title), size=15, bold=True, color=TEXT, ml=0, mr=0, mt=0, mb=0)
        d.text(s, 7.04, y + 0.42, 5.7, 0.74, body, size=14, color=BODY, ml=0, mr=0, mt=0, mb=0)


def slide_oi_volume(d: Deck) -> None:
    s = d.chrome("Two clocks", "Open interest is the swing. Volume is the day")
    block(
        d, s, ML, 1.18, 6.15, 4.35, "Open interest",
        "Contracts that are still open.\n\nThis is the book already on.\n\nIt builds the walls and\nthe flip.\n\nUse it for swing levels\nand for the regime.",
        accent=TEAL, ts=22, bs=18,
    )
    block(
        d, s, 6.76, 1.18, 6.15, 4.35, "Directionalized Volume",
        "What traded in the session,\nwith a direction on the flow.\n\nThis is the intraday view.\n\nIt does not set the swing\nwall, and it does not set\nthe monthly regime.",
        accent=AMBER, ts=22, bs=18,
    )
    d.panel(s, ML, 5.72, CONTENT_W, 1.10, accent=CORAL)
    d.text(
        s, ML + 0.24, 5.88, CONTENT_W - 0.48, 0.80,
        "A midday push on Directionalized Volume does not move the open-interest wall.",
        size=18, color=TEXT, anchor="middle", ml=0, mr=0, mt=0, mb=0,
    )


def slide_crosscheck(d: Deck) -> None:
    s = d.chrome("Check", "Check the three levels in public")
    d.panel(s, ML, 1.18, 5.9, 5.55, accent=BLUE)
    d.text(s, 0.66, 1.40, 5.4, 0.40, "Write These Down", size=20, bold=True, color=TEXT, ml=0, mr=0, mt=0, mb=0)
    for i, (name, gloss) in enumerate((("Put Wall", "Support. From open interest."), ("Call Wall", "Resistance. From open interest."), ("Flip", "Where net gamma changes sign."))):
        y = 2.05 + i * 1.35
        d.rect(s, 0.70, y, 0.56, 0.56, BLUE_DIM, line=BLUE, line_w=1.25)
        d.text(s, 0.70, y, 0.56, 0.56, str(i + 1), size=18, bold=True, color=BLUE, align="center", anchor="middle", ml=0, mr=0, mt=0, mb=0)
        d.text(s, 1.45, y, 4.4, 0.36, title_case(name), size=18, bold=True, color=TEXT, ml=0, mr=0, mt=0, mb=0)
        d.text(s, 1.45, y + 0.40, 4.4, 0.40, gloss, size=15, color=BODY, ml=0, mr=0, mt=0, mb=0)
    block(
        d, s, 6.52, 1.18, 6.39, 5.55, "If the second source disagrees",
        "The same story, not the same pixel.\n\nSame expiry?\n\nOpen interest, not today's volume?\n\nDealer gamma, or the customer sign? Customer gamma flips the colors.\n\nIf they still disagree, the level is not confirmed. The sign convention is a common cause.",
        accent=AMBER, ts=18, bs=16,
    )


def slide_month(d: Deck) -> None:
    s = d.chrome("Time", "The month and the day are two books")
    block(
        d, s, ML, 1.18, 6.15, 3.35, "Today's open interest",
        "The nearby book.\n\nSpot above the flip:\ndealer gamma is positive.\nThey tend to fade.\n\nSpot below the flip:\ndealer gamma is negative.\nThey can chase.",
        accent=TEAL, ts=20, bs=18, align="center",
    )
    block(
        d, s, 6.76, 1.18, 6.15, 3.35, "The longer book",
        "Further expiries carry their own walls.\n\nA positive longer book is a pinning backdrop.\n\nA negative longer book is a backdrop where breaks can run.\n\nIt does not replace today's strikes.",
        accent=AMBER, ts=20, bs=18, align="center",
    )
    d.panel(s, ML, 4.72, CONTENT_W, 2.05, accent=BLUE)
    in_box(
        d, s, ML + 0.28, 4.72, CONTENT_W - 0.55, 2.05,
        "If today and the longer book disagree, they are two facts.\nDo not average them into a new price.\nToday's open interest is the nearby hedge. The longer book is the backdrop.",
        size=16, color=BODY, align="center", ml=0.08, mr=0.08,
    )


def slide_clocks(d: Deck) -> None:
    s = d.chrome("Time", "Where each clock sits")
    clocks = [
        ("1", "0DTE", "The one-day gamma spike. Not the swing open interest.", AMBER),
        ("2", "Weekly book", "The near walls and the near flip.", TEAL),
        ("3", "Monthly book", "The larger strikes, and the longer regime.", BLUE),
        ("4", "The chart", "Where a close is visible against those levels.", PURPLE),
    ]
    top, height = 1.16, 5.70
    left_w = 8.05
    right_x = ML + left_w + 0.18
    right_w = CONTENT_W - left_w - 0.18
    for i, ((y, bh), (num, title, body, accent)) in enumerate(zip(bands(top, height, 4, gap=0.12), clocks)):
        d.panel(s, ML, y, left_w, bh, accent=accent)
        d.text(s, ML, y + 0.18, 0.70, 0.40, num, size=20, bold=True, color=accent, align="center", anchor="middle", ml=0, mr=0, mt=0, mb=0)
        d.text(s, ML + 0.78, y + 0.18, left_w - 0.98, 0.40, title_case(title), size=18, bold=True, color=TEXT, anchor="middle", ml=0.04, mr=0.08, mt=0, mb=0)
        in_box(d, s, ML + 0.78, y + 0.62, left_w - 0.98, bh - 0.74, body, size=16, color=BODY, align="left", ml=0.04, mr=0.08)
    block(
        d, s, right_x, top, right_w, height, "Separate books",
        "Each expiry is its own book.\n\nThe 0DTE column on a heatmap is today's gamma, not the monthly wall.\n\nA price chart shows whether a close held beyond a level. The level still comes from open interest.",
        accent=BLUE, ts=18, bs=16, align="center",
    )


def slide_align(d: Deck) -> None:
    s = d.chrome("Price and gamma", "Trade the range and the walls together")
    # aligned
    d.panel(s, ML, 1.18, 6.15, 4.20)
    d.stamp(s, 0.58, 1.32, "ILLUSTRATIVE")
    d.text(s, 0.62, 1.72, 5.6, 0.32, "They Line Up", size=18, bold=True, color=TEAL, align="center", ml=0, mr=0, mt=0, mb=0)
    d.rect(s, 1.35, 2.35, 2.55, 2.45, "1A2836", line=TEAL, line_w=1.5)
    in_box(d, s, 4.10, 2.25, 2.15, 0.85, "Call Wall\n= Box Top", size=15, color=CYAN, align="center", ml=0, mr=0)
    in_box(d, s, 4.10, 4.05, 2.15, 0.85, "Put Wall\n= Box Bottom", size=15, color=PURPLE, align="center", ml=0, mr=0)
    # conflict
    d.panel(s, 6.76, 1.18, 6.15, 4.20)
    d.stamp(s, 6.92, 1.32, "ILLUSTRATIVE")
    d.text(s, 6.96, 1.72, 5.6, 0.32, "They Do Not", size=18, bold=True, color=CORAL, align="center", ml=0, mr=0, mt=0, mb=0)
    d.hline(s, 7.15, 2.40, 2.15, CYAN, 0.03)
    d.rect(s, 7.35, 2.95, 1.75, 1.20, "1A2836", line=AMBER, line_w=1.5)
    d.hline(s, 7.15, 4.65, 2.15, PURPLE, 0.03)
    in_box(d, s, 9.30, 2.15, 3.2, 0.55, "Walls Outside", size=15, color=CYAN, align="center", ml=0, mr=0)
    in_box(d, s, 9.30, 3.15, 3.2, 0.80, "4h Box Is\nOnly the Middle", size=15, color=AMBER, align="center", ml=0, mr=0)
    d.panel(s, ML, 5.55, CONTENT_W, 1.25, accent=TEAL)
    d.text(
        s, ML + 0.24, 5.72, CONTENT_W - 0.48, 0.95,
        "It works best when they line up. Then the gamma levels and the chart levels are the same prices.\nWhen they are not, you have two stories. Do not average them into a new price.",
        size=16, color=BODY, ml=0, mr=0, mt=0, mb=0,
    )


def slide_disagree(d: Deck) -> None:
    s = d.chrome("Price and gamma", "If the maps disagree, wait")
    items = [
        (AMBER, "Box inside the walls", "The edges of the box are price action only. Do not call them walls. The walls are still outside."),
        (BLUE, "Through a wall, no close", "Price has poked a wall and the 4h has not closed out. Wait for that close. A wick is not a side."),
        (CORAL, "Wall inside the box", "The wall is the level. The rest of the box is range. Do not split the difference and invent a third price."),
    ]
    for i, (accent, title, body) in enumerate(items):
        y = 1.18 + i * 1.85
        d.panel(s, ML, y, CONTENT_W, 1.72, accent=accent)
        d.text(s, ML + 0.24, y + 0.16, CONTENT_W - 0.5, 0.40, title_case(title), size=20, bold=True, color=TEXT, ml=0, mr=0, mt=0, mb=0)
        d.text(s, ML + 0.24, y + 0.68, CONTENT_W - 0.5, 0.85, body, size=16, color=BODY, ml=0, mr=0, mt=0, mb=0)


def slide_spy_numbers(d: Deck) -> None:
    s = d.chrome("October 16, 2026", "The SPY book from that close")
    cards = [
        (PURPLE, "Put Wall", "760", "Support. Bottom of the box."),
        (CYAN, "Call Wall", "775", "Resistance. Top of the box."),
        (AMBER, "Flip", "about 774.5", "Behavior change. On the ceiling."),
        (TEXT, "Spot", "about 763.2", "Under the flip. Inside the walls."),
    ]
    for (x, w), (accent, name, value, gloss) in zip(cols(4, gap=0.16), cards):
        d.panel(s, x, 1.16, w, 3.55, accent=accent)
        d.text(s, x + 0.12, 1.40, w - 0.24, 0.36, title_case(name), size=15, bold=True, color=accent, align="center", ml=0, mr=0, mt=0, mb=0)
        d.text(s, x + 0.12, 1.90, w - 0.24, 1.15, value, size=32, bold=True, color=TEXT, align="center", anchor="middle", ml=0, mr=0, mt=0, mb=0)
        d.text(s, x + 0.16, 3.20, w - 0.32, 1.20, gloss, size=16, color=BODY, align="center", ml=0, mr=0, mt=0, mb=0)
    d.rect(s, ML, 4.90, CONTENT_W, 1.05, CORAL_DIM, line=CORAL, line_w=1.25)
    d.text(
        s, ML + 0.24, 5.02, CONTENT_W - 0.48, 0.80,
        "Regime: negative gamma. Spot is under the flip. Most of 760 to 775 is under it too.",
        size=18, color=TEXT, align="center", anchor="middle", ml=0, mr=0, mt=0, mb=0,
    )
    d.text(
        s, ML, 6.12, CONTENT_W, 0.80,
        "These are the open-interest levels from the October 16, 2026 SPY close, as used for this example.\n775 minus 760 is 15 points. On spot about 763, that box is about 2% wide.",
        size=15, color=DIM, ml=0, mr=0, mt=0, mb=0,
    )


def slide_spy_box(d: Deck) -> None:
    s = d.chrome("October 16, 2026", "The 4h range was the walls")
    d.panel(s, ML, 1.16, 7.85, 5.60)
    pmin, pmax = 754.0, 782.0
    top, bot = 1.70, 6.35
    axis_x = 2.15
    d.vline(s, axis_x, top, bot - top, LINE, 0.02)
    # range shade
    y_call = ymap(775, pmin, pmax, top, bot)
    y_put = ymap(760, pmin, pmax, top, bot)
    y_flip = ymap(774.5, pmin, pmax, top, bot)
    y_spot = ymap(763.2, pmin, pmax, top, bot)
    d.rect(s, 2.45, y_call, 3.55, y_put - y_call, "243044")
    in_box(d, s, 2.45, y_call, 3.55, y_put - y_call, "4h Range\n760 to 775", size=16, bold=True, color=TEXT, align="center", ml=0.1, mr=0.1)
    d.hline(s, 1.85, y_call, 4.35, CYAN, 0.03)
    d.dashed_h(s, 1.85, y_flip, 4.35, AMBER, 0.018, 0.10, 0.07)
    d.hline(s, 1.85, y_put, 4.35, PURPLE, 0.03)
    d.oval(s, 2.45 + 1.68, y_spot - 0.09, 0.18, 0.18, WHITE)
    d.text(s, 6.40, y_call - 0.16, 1.70, 0.32, "775 Call Wall", size=13, bold=True, color=CYAN, anchor="middle", ml=0, mr=0, mt=0, mb=0)
    d.text(s, 0.50, y_flip - 0.16, 1.25, 0.32, "Flip", size=13, bold=True, color=AMBER, anchor="middle", align="right", ml=0, mr=0.06, mt=0, mb=0)
    d.text(s, 6.40, y_spot - 0.16, 1.70, 0.32, "763.2 Spot", size=13, bold=True, color=WHITE, anchor="middle", ml=0, mr=0, mt=0, mb=0)
    d.text(s, 6.40, y_put - 0.16, 1.70, 0.32, "760 Put Wall", size=13, bold=True, color=PURPLE, anchor="middle", ml=0, mr=0, mt=0, mb=0)
    # zoom
    d.panel(s, 8.45, 1.16, 4.46, 2.55, accent=AMBER)
    d.text(s, 8.65, 1.32, 4.1, 0.30, "Top of the Box, Opened Up", size=14, bold=True, color=AMBER, ml=0, mr=0, mt=0, mb=0)
    ztop, zbot = 1.80, 3.40
    zy_call = ymap(775, 773.5, 776.2, ztop, zbot)
    zy_flip = ymap(774.5, 773.5, 776.2, ztop, zbot)
    d.hline(s, 8.75, zy_call, 2.3, CYAN, 0.03)
    d.dashed_h(s, 8.75, zy_flip, 2.3, AMBER, 0.018, 0.10, 0.06)
    d.text(s, 11.15, zy_call - 0.16, 1.55, 0.32, "775", size=14, bold=True, color=CYAN, anchor="middle", ml=0, mr=0, mt=0, mb=0)
    d.text(s, 11.15, zy_flip - 0.16, 1.55, 0.32, "774.5", size=14, bold=True, color=AMBER, anchor="middle", ml=0, mr=0, mt=0, mb=0)
    d.text(s, 8.65, 3.22, 4.05, 0.36, "Half a point apart. Linear scale.", size=12, color=DIM, ml=0, mr=0, mt=0, mb=0)
    block(
        d, s, 8.45, 3.88, 4.46, 2.88, "What the Picture Says",
        "The box and the walls\nare the same prices.\n\nThe flip is glued to\nthe call wall.\n\nAlmost the whole box\nis negative gamma.",
        accent=CORAL, ts=16, bs=15,
    )


def slide_spy_plan(d: Deck) -> None:
    s = d.chrome("October 16, 2026", "What that book was saying")
    cards = [
        (CORAL, "Close through 760", "Still under the flip.\nNegative gamma stands.\nThe downside hedge\ncan add fuel."),
        (AMBER, "Poke, then back in", "The close is inside\nthe box.\nThe outside print\ndid not hold."),
        (TEAL, "Holds above 775", "Spot crossed the flip\nand the call wall.\nThe regime may\nhave changed."),
    ]
    for (x, w), (accent, title, body) in zip(cols(3), cards):
        block(d, s, x, 1.20, w, 3.55, title, body, accent=accent, ts=18, bs=16)
    d.panel(s, ML, 4.95, CONTENT_W, 1.82, accent=BLUE)
    d.text(
        s, ML + 0.24, 5.12, CONTENT_W - 0.48, 1.48,
        "Inside 760 to 775, spot was under the flip, so most of the box was negative gamma.\nThe flip was about 774.5, half a point under the call wall. A hold above 775 is a cross of both.\nThat is a new read of the book. The old sign does not carry across the flip by itself.",
        size=16, color=BODY, ml=0, mr=0, mt=0, mb=0,
    )

def slide_pos_play(d: Deck) -> None:
    s = d.chrome("Structures", "Buying and selling in positive gamma")
    (lx, lw), (rx, rw) = cols(2)
    block(
        d, s, lx, 1.16, lw, 5.70, "Long premium",
        "A long call or a long put is long gamma.\n\nDelta moves with the price.\nTheta is the cost, every day.\n\nA debit spread is the same idea with less gamma, less vega, and a capped gain.\n\nMax loss on both is the debit.",
        accent=TEAL, ts=20, bs=18, align="center",
    )
    block(
        d, s, rx, 1.16, rw, 5.70, "Credit spreads",
        "A credit spread sells premium, and the loss is capped.\n\nA call credit spread sits above a call wall.\nA put credit spread sits below a put wall.\n\nMax loss is the width minus the credit.\nThe short strike is where that loss starts.\n\nThe position is short gamma.",
        accent=AMBER, ts=20, bs=18, align="center",
    )


def slide_pos_candle(d: Deck) -> None:
    s = d.chrome("Structures", "Condors and butterflies around the walls")
    cards = [
        (TEAL, "Iron condor", "A short put spread below the put wall, and a short call spread above the call wall.\n\nProfit if price finishes between the short strikes.\nThe loss is defined. It is larger than the credit."),
        (AMBER, "Butterfly", "Long one wing, short two at the body, long the other wing.\n\nProfit if price finishes near the body.\nThe body is often a heavy strike, where a pin is possible."),
        (CORAL, "What they share", "Both are short gamma next to a long option.\n\nA quiet range is the gain.\nA move through a wing is the loss.\nThe wings are the defined-risk lines."),
    ]
    for (x, w), (accent, title, body) in zip(cols(3), cards):
        block(d, s, x, 1.16, w, 5.70, title, body, accent=accent, ts=18, bs=16, align="center")


def slide_neg_play(d: Deck) -> None:
    s = d.chrome("Structures", "Buying and selling in negative gamma")
    (lx, lw), (rx, rw) = cols(2)
    block(
        d, s, lx, 1.16, lw, 4.40, "Long options and debit spreads",
        "Negative dealer gamma means the hedge can chase a move.\n\nA long option is long gamma, so its delta grows with the run.\nA debit spread does the same with less gamma and a capped gain.\n\nTheta is still paid.",
        accent=CORAL, ts=20, bs=17, align="center",
    )
    block(
        d, s, rx, 1.16, rw, 4.40, "Short premium",
        "A credit spread, condor, or butterfly is short gamma.\n\nThe chase is the loss, up to the defined wing.\nA naked short has the same short gamma and no wing.\n\nThe credit is the most that position can make.",
        accent=AMBER, ts=20, bs=17, align="center",
    )
    d.panel(s, ML, 5.72, CONTENT_W, 1.12, accent=BLUE)
    in_box(
        d, s, ML + 0.24, 5.72, CONTENT_W - 0.48, 1.12,
        "Dealer gamma describes their hedge. These structures are ways to be long or short gamma yourself. They can match the dealer book, or oppose it.",
        size=16, color=BODY, align="center", ml=0.08, mr=0.08,
    )


def slide_neg_candle(d: Deck) -> None:
    s = d.chrome("Your book", "Long gamma and short gamma")
    (lx, lw), (rx, rw) = cols(2)
    block(
        d, s, lx, 1.16, lw, 4.40, "Long gamma",
        "Long calls and long puts.\n\nAs price moves, delta increases in that direction.\nA further move can add to the gain.\n\nTheta is the cost.\nA debit spread is long gamma in a smaller size. The short leg cuts the gamma and caps the gain.",
        accent=TEAL, ts=20, bs=17, align="center",
    )
    block(
        d, s, rx, 1.16, rw, 4.40, "Short gamma",
        "Short calls, short puts, credit spreads, condors, and butterflies.\n\nAs price moves, delta increases against the position.\nA quiet underlying is when theta is collected.\n\nA large move is when the loss shows up.\nDefined-risk shorts stop at the long strike. Naked shorts do not.",
        accent=CORAL, ts=20, bs=17, align="center",
    )
    d.panel(s, ML, 5.72, CONTENT_W, 1.12, accent=AMBER)
    in_box(
        d, s, ML + 0.24, 5.72, CONTENT_W - 0.48, 1.12,
        "This is your position. Dealer gamma is their hedge. You can be long gamma while dealers are short, or the other way around.",
        size=16, color=TEXT, align="center", ml=0.08, mr=0.08,
    )


def slide_poke(d: Deck) -> None:
    s = d.chrome("Breaks", "A poke is not a break")
    d.panel(s, ML, 1.18, 6.15, 4.35)
    d.stamp(s, 0.58, 1.32, "ILLUSTRATIVE")
    d.text(s, 0.64, 1.72, 5.5, 0.34, "Comes Back", size=20, bold=True, color=AMBER, align="center", ml=0, mr=0, mt=0, mb=0)
    d.hline(s, 0.75, 2.85, 5.4, LINE, 0.02)
    d.text(s, 0.80, 2.52, 1.5, 0.28, "Range", size=12, color=DIM, ml=0, mr=0, mt=0, mb=0)
    d.candle(s, 2.35, 2.15, 4.15, 3.15, 3.45, DIM)
    d.candle(s, 3.55, 3.05, 4.35, 3.35, 3.85, DIM)
    d.text(s, 4.15, 3.20, 2.1, 1.15, "Pushes out.\nCloses back in.\nNext bar stays in.", size=14, color=BODY, ml=0, mr=0, mt=0, mb=0)
    d.text(s, 0.70, 4.75, 5.5, 0.50, "The Close Came Back In.", size=16, bold=True, color=TEXT, align="center", ml=0, mr=0, mt=0, mb=0)

    d.panel(s, 6.76, 1.18, 6.15, 4.35)
    d.stamp(s, 6.94, 1.32, "ILLUSTRATIVE")
    d.text(s, 7.00, 1.72, 5.5, 0.34, "Holds", size=20, bold=True, color=TEAL, align="center", ml=0, mr=0, mt=0, mb=0)
    d.hline(s, 7.05, 3.35, 5.4, LINE, 0.02)
    d.text(s, 7.10, 3.02, 1.5, 0.28, "Range", size=12, color=DIM, ml=0, mr=0, mt=0, mb=0)
    d.candle(s, 8.35, 1.95, 3.85, 2.25, 2.65, TEAL)
    d.candle(s, 9.70, 1.85, 3.15, 2.15, 2.55, TEAL)
    d.text(s, 10.15, 3.55, 2.4, 1.05, "Closes out.\nThe next close\nstays out.", size=14, color=BODY, ml=0, mr=0, mt=0, mb=0)
    d.text(s, 7.05, 4.75, 5.5, 0.50, "A Later Close Held Outside.", size=16, bold=True, color=TEXT, align="center", ml=0, mr=0, mt=0, mb=0)

    d.panel(s, ML, 5.70, CONTENT_W, 1.10, accent=AMBER)
    in_box(
        d, s, ML + 0.22, 5.70, CONTENT_W - 0.44, 1.10,
        "A wick, or one close outside, is not the same as a close that stays outside. If price closes back inside, the outside close did not hold.",
        size=16, color=BODY, align="center", ml=0.08, mr=0.08,
    )


def slide_rebuild(d: Deck) -> None:
    s = d.chrome("Breaks", "When a close holds, the book can change")
    steps = [
        ("1", "Close outside", "A candle finishes outside the old range."),
        ("2", "It holds", "A later candle is still closed outside. It did not come back in."),
        ("3", "The book can change", "Spot may have crossed the flip. The walls may be the same, or not."),
        ("4", "Read it again", "New open interest. New flip. New spot. The old map was the old book."),
    ]
    xs = cols(4, gap=0.16)
    for (x, w), (num, title, body) in zip(xs, steps):
        d.panel(s, x, 1.16, w, 3.30, accent=AMBER)
        stack_box(
            d, s, x, 1.16, w, 3.30,
            [(num, 20, True, AMBER), (title_case(title), 16, True, TEXT), (body, 14, False, BODY)],
            gap=0.08,
        )
    d.panel(s, ML, 4.64, CONTENT_W, 2.16, accent=CORAL)
    in_box(
        d, s, ML + 0.24, 4.64, CONTENT_W - 0.48, 2.16,
        "A close back inside means the outside close did not hold.\n\nIf the close also crosses the flip, the dealer hedge may have changed sign. That is a new read of the book, not a label carried over from the old one.\n\nThe walls are still the open-interest strikes. They move when the book moves.",
        size=16, color=BODY, align="center", ml=0.1, mr=0.1,
    )


def slide_alerts(d: Deck) -> None:
    s = d.chrome("Squeeze", "A gamma squeeze is the hedge feeding itself")
    (lx, lw), (rx, rw) = cols(2)
    block(
        d, s, lx, 1.16, lw, 5.70, "Calls, and a rising price",
        "Dealers who are short calls are short gamma.\n\nPrice rises. Call delta rises. They buy stock to get back to neutral.\n\nThat buying can push price higher, which raises delta again.\n\nThe loop is strongest near the strike, and near expiry, where gamma peaks.\n\nIt is a feedback story. It is not a schedule, and it does not have to finish.",
        accent=CORAL, ts=20, bs=17, align="center",
    )
    block(
        d, s, rx, 1.16, rw, 5.70, "Puts, and a falling price",
        "The other direction is the same gamma.\n\nDealers who are short puts, and a falling price, sell stock to hedge.\n\nThat selling can push price lower.\n\nThe upside version is the one called a squeeze.\nBoth are negative-gamma hedging.\n\nNeither one is a measured share count. The chart of gamma shows where the sensitivity is highest.",
        accent=TEAL, ts=20, bs=17, align="center",
    )


def slide_watch(d: Deck) -> None:
    s = d.chrome("Volatility", "VIX is the price of index volatility")
    (lx, lw), (rx, rw) = cols(2)
    block(
        d, s, lx, 1.16, lw, 5.70, "What VIX is",
        "VIX is the market's price for 30-day S&P 500 volatility.\n\nIt is implied volatility. It is not a print of what the index will do.\n\nA higher VIX means richer option premiums.\nThe same delta then hedges with more sensitivity to volatility.\n\nA lower VIX means cheaper premiums.\n\nThe gamma shape still peaks at the strike. VIX scales the premium.",
        accent=AMBER, ts=20, bs=17, align="center",
    )
    block(
        d, s, rx, 1.16, rw, 5.70, "High IV and low IV",
        "When implied volatility is high, a long option costs more, and a credit collected is larger.\n\nWhen implied volatility is low, the long option is cheaper, and the credit is smaller.\n\nIV can rise or fall with price held still. That sensitivity is vega.\n\nA fast drop in IV is crush. The crush chart is that drop, including a case where the stock rose and the call still fell.",
        accent=BLUE, ts=20, bs=17, align="center",
    )


def slide_bands(d: Deck) -> None:
    s = d.chrome("Markets", "A single stock and an index are different books")
    (lx, lw), (rx, rw) = cols(2)
    block(
        d, s, lx, 1.16, lw, 5.70, "A single stock",
        "Open interest can sit in size at a few strikes.\n\nThe wall is sharper. The stock can gap through it.\nA trading halt is possible.\n\nThe dealer hedge is in that one name.\nNews and earnings land on the name itself.\n\nThe gamma peak near expiry is the same shape. The book behind it is thinner.",
        accent=CORAL, ts=20, bs=17, align="center",
    )
    block(
        d, s, rx, 1.16, rw, 5.70, "An index",
        "The book is deeper. The hedge is spread across many names.\n\nWalls are usually wider, and less abrupt, than a single-name wall.\n\nSPX, SPY, and the futures are related. They are not the same open interest.\n\nA name can move while the index does not, and the reverse.\nThe October 16 example in this deck is SPY.",
        accent=TEAL, ts=20, bs=17, align="center",
    )


def slide_confirm(d: Deck) -> None:
    s = d.chrome("The model", "What GEX does not know")
    cards = [
        (AMBER, "A model, not a tape", "GEX estimates a dealer hedge from open interest. It is not the shares already bought or sold."),
        (CORAL, "The sign is assumed", "The model assigns options to dealers or to customers. Swap that assumption and the colors flip."),
        (BLUE, "OI is not the session", "Open interest is the open book. Volume can be opening or closing. The swing map and the day are different."),
        (TEAL, "It leaves things out", "It does not show the hedge already on, futures versus stock, or inventory you cannot see. A level is pressure, not a promise."),
    ]
    for i, (accent, title, body) in enumerate(cards):
        col = i % 2
        row = i // 2
        x = ML + col * 6.34
        y = 1.16 + row * 2.90
        block(d, s, x, y, 6.16, 2.74, title, body, accent=accent, ts=18, bs=16, align="center")


def slide_mas(d: Deck) -> None:
    s = d.chrome("Structures", "A naked short has no cap")
    (lx, lw), (rx, rw) = cols(2)
    block(
        d, s, lx, 1.16, lw, 4.40, "Short call",
        "The loss has no cap. The stock can rise without a limit.\n\nAssignment turns the short call into short stock, or into the cash to buy the shares.\n\nThe credit is the most the position can make.\nMargin is required. The broker can close the position.",
        accent=CORAL, ts=20, bs=17, align="center",
    )
    block(
        d, s, rx, 1.16, rw, 4.40, "Short put",
        "The loss grows as the stock falls.\n\nNear zero, the loss is about the strike times 100, minus the credit.\n\nAssignment turns the short put into long stock.\n\nSame facts: the loss is not capped on the way down, margin is required, and the broker can close it.",
        accent=AMBER, ts=20, bs=17, align="center",
    )
    d.panel(s, ML, 5.72, CONTENT_W, 1.12, accent=BLUE)
    in_box(
        d, s, ML + 0.22, 5.72, CONTENT_W - 0.44, 1.12,
        "A long option at a further strike caps that short. The cap is a spread. Without the long leg, the short is naked. American options can be assigned before expiry.",
        size=16, color=BODY, align="center", ml=0.08, mr=0.08,
    )


def slide_side(d: Deck) -> None:
    s = d.chrome("Structures", "The side is the view. The structure is the risk")
    block(
        d, s, ML, 1.16, 6.15, 5.70, "Upside",
        "A long call, or a call debit spread, is long gamma to the upside.\n\nA call credit spread, or a short call, is short gamma to the upside.\n\nThe first pays a debit. The second collects a credit.\n\nOnly the spread, or a covered short, has a defined loss.",
        accent=TEAL, ts=20, bs=18, align="center",
    )
    block(
        d, s, 6.76, 1.16, 6.15, 5.70, "Downside",
        "A long put, or a put debit spread, is long gamma to the downside.\n\nA put credit spread, or a short put, is short gamma to the downside.\n\nSame split: debit or credit.\n\nA naked short put is not capped until the stock is near zero.",
        accent=CORAL, ts=20, bs=18, align="center",
    )


def slide_fork(d: Deck) -> None:
    s = d.chrome("Structures", "Five ways the same view is held")
    rows = [
        (TEAL, "Long option", "Long gamma. Debit. Loss capped at the debit. Gain is open."),
        (BLUE, "Debit spread", "Long gamma, smaller. Debit. Loss capped. Gain capped at the short strike."),
        (AMBER, "Credit spread", "Short gamma. Credit. Gain capped at the credit. Loss capped at the width."),
        (PURPLE, "Condor or butterfly", "Short gamma around the walls or a strike. Gain if price sits. Loss defined at the wings."),
        (CORAL, "Naked short", "Short gamma. Credit. Gain capped. Loss is not capped. Margin is required."),
    ]
    top, height = 1.16, 5.70
    gap = 0.12
    row_h = (height - gap * (len(rows) - 1)) / len(rows)
    for i, (accent, name, line) in enumerate(rows):
        y = top + i * (row_h + gap)
        d.panel(s, ML, y, CONTENT_W, row_h, accent=accent)
        d.rect(s, ML, y + 0.055, 2.55, row_h - 0.055, NEST)
        d.text(s, ML, y + 0.055, 2.55, row_h - 0.055, title_case(name), size=15, bold=True, color=accent, align="center", anchor="middle", ml=0.08, mr=0.08, mt=0, mb=0)
        in_box(d, s, ML + 2.7, y, CONTENT_W - 2.9, row_h, line, size=16, color=TEXT, align="left", ml=0.1, mr=0.12)


def slide_single(d: Deck) -> None:
    s = d.chrome("Structures", "What a single long option is")
    points = [
        (TEAL, "Long gamma", "Delta increases as price moves in favor of the option."),
        (BLUE, "The debit", "The most the long option can lose is the premium paid."),
        (AMBER, "Theta and vega", "Time decay is paid every day. A drop in IV can cut the price even if the stock rises."),
    ]
    for i, (accent, title, body) in enumerate(points):
        y = 1.16 + i * 1.48
        d.panel(s, ML, y, CONTENT_W, 1.36, accent=accent)
        in_box(d, s, ML + 0.16, y, 3.45, 1.36, title, size=18, bold=True, color=accent, align="center", ml=0.06, mr=0.06)
        in_box(d, s, ML + 3.7, y, CONTENT_W - 3.9, 1.36, body, size=16, color=BODY, align="left", ml=0.08, mr=0.1)
    d.panel(s, ML, 5.62, CONTENT_W, 1.20, accent=CORAL)
    in_box(
        d, s, ML + 0.24, 5.62, CONTENT_W - 0.48, 1.20,
        "One contract is 100 shares of exposure, times delta. The shares are the exposure figure. They are not a stock position.",
        size=16, color=BODY, align="center", ml=0.08, mr=0.08,
    )


def slide_spread(d: Deck) -> None:
    s = d.chrome("Structures", "What a vertical spread is")
    d.panel(s, ML, 1.16, 7.7, 5.70, accent=BLUE)
    d.text(s, 0.66, 1.36, 7.2, 0.70, "Same expiry. One long strike.\nOne short strike, further out.", size=18, color=TEXT, align="center", ml=0, mr=0, mt=0, mb=0)
    d.rect(s, 0.75, 2.20, 3.15, 1.15, TEAL_DIM, line=TEAL, line_w=1.25)
    in_box(d, s, 0.85, 2.20, 2.95, 1.15, "Long the nearer\nstrike.", size=16, color=TEXT, align="center", ml=0.06, mr=0.06)
    d.arrow(s, 4.05, 2.58, 0.45, 0.30, BLUE, "right")
    d.rect(s, 4.65, 2.20, 3.15, 1.15, BLUE_DIM, line=BLUE, line_w=1.25)
    in_box(d, s, 4.75, 2.20, 2.95, 1.15, "Short the further\nstrike.", size=16, color=TEXT, align="center", ml=0.06, mr=0.06)
    in_box(
        d, s, 0.70, 3.55, 7.1, 3.05,
        "Debit spread: the long strike costs more. Max loss is the debit. Past the short strike the gain stops.\n\nCredit spread: the short strike brings in more. Max gain is the credit. Max loss is the width minus the credit.\n\nCall spread for the upside. Put spread for the downside.",
        size=16, color=BODY, align="center", ml=0.08, mr=0.1,
    )
    block(
        d, s, 8.30, 1.16, 4.61, 5.70, "The short leg",
        "It is covered by the long leg of the same expiry.\n\nAmerican options can be assigned early on that short leg.\n\nThe long leg is what makes the loss defined.\n\nWithout it, the short is naked.",
        accent=AMBER, ts=18, bs=16, align="center",
    )


def slide_payoff(d: Deck) -> None:
    s = d.chrome("Example", "Payoff at expiry, single versus spread")
    d.picture_in(s, d.charts["payoff"], ML, 1.14, 7.25, 5.20)
    call = money(FACTS["call_25"])
    spr = money(FACTS["spread_25"])
    gain = money(FACTS["max_gain_25"])
    block(
        d, s, 7.85, 1.20, 5.06, 4.55, "The example debits",
        f"21 days. IV 25%.\nSpot 100. Rate 4%.\nNo dividend.\n\nLong 100 call: {call}\n100/105 spread: {spr}\n\nSpread max gain:\n{gain} per share.\nAbout {float(gain)*100:.0f} dollars\na contract.",
        accent=TEAL, ts=16, bs=15,
    )
    d.text(
        s, ML, 6.55, CONTENT_W, 0.45,
        "EXAMPLE. Black-Scholes prices, then the expiry payoff of those debits. Per share. A contract is times 100. Not a forecast.",
        size=13, color=DIM, ml=0, mr=0, mt=0, mb=0,
    )


def slide_rich_iv(d: Deck) -> None:
    s = d.chrome("Example", "What rich IV does to the price")
    d.text(
        s, ML, 1.14, CONTENT_W, 0.40,
        "EXAMPLE. Same 21-day 100/105 call spread. Spot 100. Rate 4%. No dividend. Only IV changes.",
        size=14, color=DIM, ml=0, mr=0, mt=0, mb=0,
    )
    low = FACTS["rich"][0.18]
    high = FACTS["rich"][0.40]
    cols_iv = [
        ("IV 18%", low, TEAL),
        ("IV 40%", high, CORAL),
    ]
    # header row labels
    labels = [("ATM Call", "call"), ("Spread Debit", "spread"), ("Spread Max Gain", "max_gain")]
    for i, (name, key) in enumerate(labels):
        y = 2.20 + i * 1.40
        in_box(d, s, ML, y, 3.15, 1.22, name, size=16, bold=True, color=TEXT, align="center", ml=0.04, mr=0.06)
        for j, (title, data, accent) in enumerate(cols_iv):
            x = 3.75 + j * 4.55
            if i == 0:
                d.text(s, x, 1.62, 4.35, 0.40, title, size=16, bold=True, color=accent, align="center", ml=0, mr=0, mt=0, mb=0)
            d.panel(s, x, y, 4.35, 1.22, accent=accent)
            in_box(d, s, x, y, 4.35, 1.22, money(data[key]), size=28, bold=True, color=TEXT, align="center", ml=0.08, mr=0.08)
    d.text(
        s, ML, 6.40, CONTENT_W, 0.60,
        "The single call more than doubles. The spread debit rises less. Rich IV raises the single more than it raises the spread.",
        size=15, color=BODY, ml=0, mr=0, mt=0, mb=0,
    )


def slide_place(d: Deck) -> None:
    s = d.chrome("October 16, 2026", "How a structure sat on that map")
    d.panel(s, ML, 1.16, CONTENT_W, 0.85, accent=AMBER)
    in_box(
        d, s, ML + 0.22, 1.16, CONTENT_W - 0.44, 0.85,
        "A picture of the October 16 book. Not an order. IV for that session is not in the example.",
        size=16, color=TEXT, align="center", ml=0.08, mr=0.08,
    )
    block(
        d, s, ML, 2.18, 6.15, 4.55, "Through 760, still negative",
        "Spot stays under the flip.\n\nA long put, or a put debit spread, is long gamma to the downside.\n\nA put credit spread is short gamma under the put wall.\n\nA naked short put has no cap if the drop continues.",
        accent=CORAL, ts=18, bs=16, align="center",
    )
    block(
        d, s, 6.76, 2.18, 6.15, 4.55, "Through 775, across the flip",
        "That close crosses the call wall and the flip, about 774.5.\n\nThe dealer sign may have changed. Read the book again.\n\nA long call is long gamma above there.\nA call credit spread is short gamma at the wall.\nA naked short call has no cap if the rise continues.",
        accent=TEAL, ts=18, bs=16, align="center",
    )


def slide_delta_bands(d: Deck) -> None:
    s = d.chrome("Delta", "High delta, mid delta, low delta")
    d.picture_in(s, d.charts["delta"], ML, 1.16, 6.95, 5.14)
    band_rows = FACTS["bands"]
    # band order is 106, 100, 95. Present high, mid, low by strike descending delta.
    ordered = sorted(band_rows, key=lambda b: -b["delta"])
    names = ["High", "Mid", "Low"]
    gloss = [
        "Costs more. Moves like the stock.",
        "0.40 to 0.60. Reacts fast.",
        "Cheap. Needs a big move. Often expires worthless.",
    ]
    for i, ((y, bh), band, name, line) in enumerate(zip(bands(1.16, 5.14, 3, gap=0.14), ordered, names, gloss)):
        d.panel(s, 7.55, y, 5.36, bh)
        stack_box(
            d, s, 7.55, y, 5.36, bh,
            [
                (f"{name}    {band['delta']:.2f}    Strike {band['k']:.0f}", 15, True, TEXT),
                (f"{line}\nAbout {band['price']:.2f} a share. About {band['shares']:.0f} shares of exposure.", 14, False, BODY),
            ],
            align="center",
            gap=0.06,
        )
    d.text(
        s, ML, 6.62, CONTENT_W, 0.40,
        "EXAMPLE. 21 days, IV 25%, spot 100, rate 4%, no dividend. Exposure is delta times 100 shares, on a small move. That figure is not a share position.",
        size=12, color=DIM, ml=0, mr=0, mt=0, mb=0,
    )


def slide_itm(d: Deck) -> None:
    s = d.chrome("Delta", "Delta, ignoring the sign")
    d.panel(s, ML, 1.18, CONTENT_W, 1.35, accent=AMBER)
    d.text(
        s, ML + 0.24, 1.36, CONTENT_W - 0.48, 1.00,
        "Delta (ignoring the sign) is roughly the chance the option finishes in the money.",
        size=22, bold=True, color=TEXT, anchor="middle", ml=0, mr=0, mt=0, mb=0,
    )
    (lx, lw), (rx, rw) = cols(2)
    d.panel(s, lx, 2.68, lw, 2.45, accent=TEAL)
    stack_box(
        d, s, lx, 2.68, lw, 2.45,
        [("Call Delta +0.45", 20, True, TEAL), ("Ignore the plus.\nRead it as about a 45 percent chance.", 18, False, BODY)],
        gap=0.12,
    )
    d.panel(s, rx, 2.68, rw, 2.45, accent=CORAL)
    stack_box(
        d, s, rx, 2.68, rw, 2.45,
        [("Put Delta -0.45", 20, True, CORAL), ("Ignore the minus.\nSame read. About 45 percent.", 18, False, BODY)],
        gap=0.12,
    )
    d.panel(s, ML, 5.25, CONTENT_W, 1.52, accent=BLUE)
    d.text(
        s, ML + 0.24, 5.42, CONTENT_W - 0.48, 1.20,
        "Rough guide. Closer on short-dated options.\nIn the model, the chance of finishing in the money is a nearby number, not this one.\nUse it to plan. Do not treat it as exact.",
        size=16, color=BODY, ml=0, mr=0, mt=0, mb=0,
    )


def slide_theta(d: Deck) -> None:
    s = d.chrome("Theta", "Theta is decay")
    d.picture_in(s, d.charts["decay"], ML, 1.16, 7.45, 5.20)
    start = FACTS["decay"][0]
    end = FACTS["decay"][-1]
    block(
        d, s, 8.05, 1.16, CONTENT_W - (8.05 - ML), 5.20, "What the curves say",
        f"Spot held at 100.\nIV 25%. EXAMPLE.\n\nSingle call, 45 days\nto 1 day: {money(start[1])}\nto {money(end[1])}.\n\n100/105 spread:\n{money(start[2])} to\n{money(end[2])}.\n\nThe single gives up\nmore dollars. The curve\nsteepens into expiry.\n\nWith a day left, the 105\ncall is nearly worthless,\nso the two meet.",
        accent=BLUE, ts=16, bs=14,
    )
    d.text(
        s, ML, 6.58, CONTENT_W, 0.42,
        "A long option pays theta every day it is open. A short option collects that theta.",
        size=13, color=DIM, ml=0, mr=0, mt=0, mb=0,
    )


def slide_crush(d: Deck) -> None:
    s = d.chrome("IV", "IV crush is a fast drop in IV")
    d.picture_in(s, d.charts["crush"], ML, 1.16, 7.15, 5.20)
    a = money(FACTS["crush_start"])
    b = money(FACTS["crush_flat"])
    c = money(FACTS["crush_up"])
    block(
        d, s, 7.75, 1.16, CONTENT_W - (7.75 - ML), 5.20, "The three bars",
        f"14-day ATM call.\nEXAMPLE.\n\nStart, IV 50%:\n{a}\n\nSpot stays 100,\nIV falls to 25%:\n{b}\n\nSpot rises to 102,\nIV falls to 25%:\n{c}\n\nThe stock rose.\nThe call still fell.",
        accent=CORAL, ts=16, bs=15,
    )
    d.text(
        s, ML, 6.55, CONTENT_W, 0.48,
        "Often after an event. The option can lose value even if the stock stays flat, or rises. Black-Scholes, rate 4%, no dividend.",
        size=13, color=DIM, ml=0, mr=0, mt=0, mb=0,
    )


def slide_events(d: Deck) -> None:
    s = d.chrome("Events", "Earnings, and the 0DTE book")
    block(
        d, s, ML, 1.16, 6.15, 5.70, "Earnings",
        "IV is often bid up into the report, then it drops.\n\nThat drop is crush. The option can lose value even if the stock moves the right way.\n\nThe actual IV is on the chain. It is not assumed.\n\nA long option carries that vega. A spread carries less of it.",
        accent=AMBER, ts=22, bs=18, align="center",
    )
    block(
        d, s, 6.76, 1.16, 6.15, 5.70, "0DTE",
        "This is the one-day spike on the gamma chart.\n\nTheta is gone by the close.\n\nThe 0DTE column is that day's gamma.\n\nIt is not the weekly or monthly open interest, and it is not the swing wall.",
        accent=CORAL, ts=22, bs=18, align="center",
    )


def slide_stops(d: Deck) -> None:
    s = d.chrome("Risk", "Where the loss stops")
    cards = [
        (TEAL, "Long option", "The premium paid. The contract can go to zero. It cannot go below zero."),
        (BLUE, "Debit spread", "The debit. The gain stops at the distance between the strikes, minus that debit."),
        (AMBER, "Credit spread", "The width of the strikes, minus the credit. The gain stops at the credit."),
        (CORAL, "Naked short", "No defined stop in the option itself. A short call is open-ended. A short put runs toward the strike."),
    ]
    for i, (accent, title, body) in enumerate(cards):
        y = 1.16 + i * 1.42
        d.panel(s, ML, y, CONTENT_W, 1.30, accent=accent)
        in_box(d, s, ML + 0.16, y, 3.3, 1.30, title_case(title), size=18, bold=True, color=accent, align="center", ml=0.06, mr=0.08)
        in_box(d, s, ML + 3.6, y, CONTENT_W - 3.85, 1.30, body, size=16, color=BODY, align="left", ml=0.1, mr=0.12)


def slide_risk(d: Deck) -> None:
    s = d.chrome("Risk", "Margin and early assignment")
    (lx, lw), (rx, rw) = cols(2)
    block(
        d, s, lx, 1.16, lw, 5.70, "Margin",
        "A long option is paid in full. The debit is the cash outlay.\n\nA short option requires margin, because the loss may exceed the credit.\n\nA defined-risk spread margins the width, not an open-ended move.\n\nThe broker can close a position that no longer fits the margin.",
        accent=BLUE, ts=20, bs=17, align="center",
    )
    block(
        d, s, rx, 1.16, rw, 5.70, "Early assignment",
        "American options can be assigned before expiry.\n\nThe short leg of a spread is the one that can be assigned.\n\nThe long leg is the cover. It does not stop the assignment notice. It caps the stock position that assignment creates.\n\nA naked short has no long leg behind it.",
        accent=AMBER, ts=20, bs=17, align="center",
    )


def slide_flow(d: Deck) -> None:
    s = d.chrome("Flow", "Flow, dark pool, and the dollars")
    flow_cards = [
        (PURPLE, "Options flow", "Large prints at a wall are more size than prints in the middle of the range.\n\nSize is not the gamma map."),
        (BLUE, "Dark pool", "Large stock prints near a wall are size in the stock, at that price.\n\nSize is not the open-interest wall."),
        (AMBER, "A print", "A print does not replace the regime, the wall, or a close.\n\nIt is a number of contracts, or of shares."),
    ]
    for (x, w), (accent, title, body) in zip(cols(3), flow_cards):
        block(d, s, x, 1.16, w, 3.55, title, body, accent=accent, ts=18, bs=16, align="center")
    d.panel(s, ML, 4.90, CONTENT_W, 1.85, accent=TEAL)
    d.text(s, ML + 0.22, 5.05, CONTENT_W - 0.44, 0.36, "Dollar Premium. Price x Contracts x 100.", size=16, bold=True, color=TEXT, align="center", ml=0, mr=0, mt=0, mb=0)
    d.text(
        s, ML + 0.22, 5.48, CONTENT_W - 0.44, 1.05,
        "EXAMPLE.  2.00 x 500 x 100 = 100,000 dollars.    12.00 x 80 x 100 = 96,000 dollars.\nSimilar money. Very different if you only count contracts.",
        size=16, color=BODY, align="center", ml=0, mr=0, mt=0, mb=0,
    )


def slide_run16(d: Deck) -> None:
    s = d.chrome("October 16, 2026", "That book, in one pass")
    rows = [
        ("Levels", "Put wall 760. Call wall 775. Flip about 774.5. Spot about 763.2."),
        ("Regime", "Negative. Spot was under the flip. Most of 760 to 775 was too."),
        ("Range", "The 4h range was 760 to 775. It matched the walls."),
        ("760", "A close through 760 stayed under the flip. The downside hedge could add fuel."),
        ("775", "A hold above 775 crossed the call wall and the flip. That is a new read."),
        ("Structure", "Long gamma rides a run. Short gamma, including a naked short, is exposed to it."),
        ("Model", "IV for that day is not in the example. GEX is a model of this book, not a tape."),
    ]
    for i, (label, line) in enumerate(rows):
        y = 1.16 + i * 0.78
        d.panel(s, ML, y, CONTENT_W, 0.70)
        d.rect(s, ML, y, 1.55, 0.70, NEST)
        d.text(s, ML, y, 1.55, 0.70, label, size=14, bold=True, color=AMBER, align="center", anchor="middle", ml=0.04, mr=0.04, mt=0, mb=0)
        d.text(s, ML + 1.72, y, CONTENT_W - 1.9, 0.70, line, size=14, color=TEXT, anchor="middle", ml=0.08, mr=0.08, mt=0, mb=0)


def slide_check(d: Deck) -> None:
    s = d.chrome("Checklist", "Before the structure")
    groups = [
        ("Regime", [
            "Spot versus the flip, on the expiry you mean.",
            "Dealer positive: the hedge tends to fade.",
            "Dealer negative: the hedge can chase.",
        ]),
        ("Levels", [
            "Put wall and call wall, from open interest.",
            "The same story on a second public source.",
            "Customer gamma and dealer gamma are opposite signs.",
        ]),
        ("The model", [
            "GEX is an estimate from open interest.",
            "It is not the shares already hedged.",
            "0DTE gamma is not the swing wall.",
            "A wick is not a held close.",
        ]),
        ("Structure", [
            "Long gamma or short gamma, for your position.",
            "Defined loss, or a naked short with no cap.",
            "Debit, or credit. Theta is the other side of gamma.",
            "Dollar premium is price times contracts times 100.",
        ]),
    ]
    for i, (title, items) in enumerate(groups):
        col = i % 2
        row = i // 2
        x = ML + col * 6.34
        y = 1.16 + row * 2.85
        d.panel(s, x, y, 6.16, 2.72, accent=TEAL)
        d.text(s, x + 0.16, y + 0.10, 5.84, 0.34, title_case(title), size=16, bold=True, color=TEAL, align="center", ml=0, mr=0, mt=0, mb=0)
        for (yy, bh), line in zip(bands(y + 0.50, 2.10, len(items), gap=0.04), items):
            d.rect(s, x + 0.22, yy + (bh - 0.16) / 2, 0.16, 0.16, None, line=TEAL, line_w=1.25)
            in_box(d, s, x + 0.46, yy, 5.50, bh, line, size=14, color=BODY, align="left", ml=0.04, mr=0.08)


def slide_mistakes(d: Deck) -> None:
    s = d.chrome("Mistakes", "Easy ways to misread the book")
    items = [
        "Treating the flip as support.",
        "Reading the heatmap as put gamma.",
        "Looking for a put-wall line on the by-strike chart.",
        "Looking for a flip line on the by-strike chart.",
        "Setting a swing level from Directionalized Volume.",
        "Using the 0DTE column as the swing wall.",
        "Confusing dealer gamma with your own gamma.",
        "Treating GEX as shares already traded.",
        "Ignoring the customer-versus-dealer sign.",
        "Calling a wick a held close.",
        "Counting contracts and skipping the dollar premium.",
        "Forgetting a naked short has no cap.",
        "Averaging two expiries into a new strike.",
        "Reading a pin as a promise.",
    ]
    for i, line in enumerate(items):
        col = i // 7
        row = i % 7
        x = ML + col * 6.34
        y = 1.16 + row * 0.78
        d.panel(s, x, y, 6.16, 0.70)
        d.rect(s, x, y, 0.62, 0.70, NEST)
        d.text(s, x, y, 0.62, 0.70, f"{i+1:02d}", size=14, bold=True, color=CORAL, align="center", anchor="middle", ml=0, mr=0, mt=0, mb=0)
        d.text(s, x + 0.74, y, 5.25, 0.70, line, size=14, color=TEXT, anchor="middle", ml=0.06, mr=0.08, mt=0, mb=0)


def slide_gloss_a(d: Deck) -> None:
    s = d.chrome("Glossary", "The words")
    terms = [
        ("Delta", "Share count for the dealer hedge. For a long option, how it tracks the stock."),
        ("Gamma", "How fast delta changes as the stock price moves."),
        ("Long gamma", "Your delta increases as price moves your way. Long options. You pay theta."),
        ("Short gamma", "Your delta moves against you as price runs. Short options. You collect theta."),
        ("Theta", "Decay. What time takes from a long option, and what a short option collects."),
        ("Vega", "How much the option price changes when IV changes."),
        ("IV", "The volatility implied by the option's price. VIX is the index version."),
        ("GEX", "A model of dealer hedge from open interest. Not a tape of shares."),
    ]
    for i, (term, line) in enumerate(terms):
        y = 1.16 + i * 0.70
        d.rect(s, ML, y, 2.35, 0.62, NEST)
        d.text(s, ML, y, 2.35, 0.62, title_case(term), size=14, bold=True, color=AMBER, align="center", anchor="middle", ml=0.06, mr=0.06, mt=0, mb=0)
        d.text(s, ML + 2.55, y, 9.8, 0.62, line, size=15, color=BODY, anchor="middle", ml=0.08, mr=0.08, mt=0, mb=0)


def slide_gloss_b(d: Deck) -> None:
    s = d.chrome("Glossary", "The levels and the structures")
    terms = [
        ("Call Wall", "Heavy call open interest above spot. Often resistance."),
        ("Put Wall", "Heavy put open interest below spot. Often support. Longest purple bar under spot."),
        ("Flip", "Where net gamma changes sign. A behavior change. Not support."),
        ("Positive gamma", "Above the flip, for the dealer book. The hedge tends to fade. Price can range and pin."),
        ("Negative gamma", "Below the flip, for the dealer book. The hedge can chase. Price can extend."),
        ("Gamma squeeze", "Short-gamma hedging that buys into a rise, or sells into a drop, and can feed the move."),
        ("Credit spread", "Short a nearer strike, long a further one. Max gain is the credit. Max loss is the width minus the credit."),
        ("Iron condor", "A short put spread and a short call spread. Profit if price finishes between the short strikes."),
        ("Butterfly", "Long the wings, short two at the body. Profit if price finishes near the body."),
        ("Naked short", "A short option with no long leg covering it. The loss is not capped. Margin is required."),
        ("0DTE", "Expires the same day. Extreme gamma. Extreme theta. Not the swing open interest."),
        ("IV crush", "A fast drop in IV, often after an event. The option can lose value anyway."),
    ]
    left_terms = terms[:6]
    right_terms = terms[6:12]
    pairs = list(zip(left_terms, right_terms))
    row_top, row_h = 1.14, 0.90
    gap = 0.06
    for i, (left, right) in enumerate(pairs):
        y = row_top + i * (row_h + gap)
        for col, (term, line) in enumerate((left, right)):
            x = ML + col * 6.33
            d.rect(s, x, y, 6.16, row_h, CARD, line=LINE, line_w=1.0)
            d.text(s, x + 0.14, y + 0.06, 5.88, 0.28, title_case(term), size=14, bold=True, color=AMBER, ml=0, mr=0, mt=0, mb=0)
            d.text(s, x + 0.14, y + 0.36, 5.88, 0.48, line, size=13, color=BODY, anchor="middle", ml=0, mr=0, mt=0, mb=0)


def slide_one(d: Deck) -> None:
    s = d.chrome("Wrap-up", "The whole swing on one page")
    left = [
        "1  Mark put wall, call wall, flip, and spot.",
        "2  Name the expiry. 0DTE is not the swing book.",
        "3  Check the three levels on a second source.",
        "4  Dealer positive: the hedge tends to fade.",
        "5  Dealer negative: the hedge can chase.",
        "6  Your gamma is long or short, separately.",
        "7  A credit spread, condor, or butterfly is short gamma.",
        "8  A naked short has no cap, and it uses margin.",
        "9  GEX is a model. A wick is not a held close.",
    ]
    right = [
        "Flip: behavior change. Not support.",
        "Walls: call strike above, put strike below.",
        "Heatmap color: net gamma.",
        "Put wall: longest purple bar below spot.",
        "By-strike draws the call wall and spot only.",
        "Open interest: the swing. Volume: the day.",
        "VIX: the price of index volatility.",
        "Squeeze: the hedge can feed the move.",
        "Single names gap. Index books are deeper.",
        "Dollar premium: price x contracts x 100.",
        "Customer gamma flips the dealer colors.",
    ]
    (lx, lw), (rx, rw) = cols(2)
    d.panel(s, lx, 1.14, lw, 5.20, accent=TEAL)
    d.text(s, lx, 1.26, lw, 0.34, "Read It in This Order", size=16, bold=True, color=TEAL, align="center", ml=0, mr=0, mt=0, mb=0)
    for (y, bh), line in zip(bands(1.68, 4.48, len(left), gap=0.04), left):
        in_box(d, s, lx + 0.16, y, lw - 0.32, bh, line, size=15, color=BODY, align="left", ml=0.06, mr=0.06)
    d.panel(s, rx, 1.14, rw, 5.20, accent=AMBER)
    d.text(s, rx, 1.26, rw, 0.34, "Keep These Straight", size=16, bold=True, color=AMBER, align="center", ml=0, mr=0, mt=0, mb=0)
    for (y, bh), line in zip(bands(1.68, 4.48, len(right), gap=0.02), right):
        in_box(d, s, rx + 0.16, y, rw - 0.32, bh, line, size=14, color=BODY, align="left", ml=0.06, mr=0.06)
    d.panel(s, ML, 6.48, CONTENT_W, 0.50, accent=None)
    d.text(
        s, ML, 6.48, CONTENT_W, 0.50,
        "Dealers hedge. Gamma shows the push. Your gamma is a separate book.",
        size=15, bold=True, color=TEXT, align="center", anchor="middle", ml=0.1, mr=0.1, mt=0, mb=0,
    )


def finalize_package(path: Path) -> None:
    """Mark the slide size as 16:9 and embed Liberation Sans so the layout holds."""
    import zipfile
    from io import BytesIO

    regular = Path(FONT_REG).read_bytes()
    bold = Path(FONT_BOLD).read_bytes()
    buf = BytesIO()
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(buf, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/presentation.xml":
                text = data.decode("utf-8")
                text = text.replace('saveSubsetFonts="1"', 'embedTrueTypeFonts="1" saveSubsetFonts="0"', 1)
                text = text.replace('type="screen4x3"', 'type="screen16x9"', 1)
                font_xml = (
                    '<p:embeddedFontLst><p:embeddedFont>'
                    '<p:font typeface="Liberation Sans" pitchFamily="34" charset="0"/>'
                    '<p:regular r:id="rId71"/>'
                    '<p:bold r:id="rId72"/>'
                    "</p:embeddedFont></p:embeddedFontLst>"
                )
                text = text.replace("<p:defaultTextStyle>", font_xml + "<p:defaultTextStyle>", 1)
                data = text.encode("utf-8")
            elif item.filename == "ppt/_rels/presentation.xml.rels":
                text = data.decode("utf-8")
                extra = (
                    '<Relationship Id="rId71" '
                    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" '
                    'Target="fonts/font1.fntdata"/>'
                    '<Relationship Id="rId72" '
                    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" '
                    'Target="fonts/font2.fntdata"/>'
                )
                text = text.replace("</Relationships>", extra + "</Relationships>", 1)
                data = text.encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                if "fntdata" not in text:
                    text = text.replace(
                        "<Types ",
                        '<Types ',
                    )
                    text = text.replace(
                        "</Types>",
                        '<Default Extension="fntdata" ContentType="application/x-fontdata"/></Types>',
                        1,
                    )
                data = text.encode("utf-8")
            zout.writestr(item, data)
        zout.writestr("ppt/fonts/font1.fntdata", regular)
        zout.writestr("ppt/fonts/font2.fntdata", bold)
    path.write_bytes(buf.getvalue())


def render_previews(pptx_path: Path) -> None:
    out = ROOT / "previews"
    out.mkdir(parents=True, exist_ok=True)
    work = Path("/tmp/gamma-preview")
    if work.exists():
        shutil.rmtree(work)
    work.mkdir(parents=True)
    profile = Path("/tmp/lo-profile")
    profile.mkdir(parents=True, exist_ok=True)
    subprocess.check_call(
        [
            "soffice",
            "-env:UserInstallation=file:///tmp/lo-profile",
            "--headless",
            "--nolockcheck",
            "--nologo",
            "--nofirststartwizard",
            "--convert-to",
            "pdf",
            "--outdir",
            str(work),
            str(pptx_path),
        ]
    )
    pdf = work / "gamma_swing.pdf"
    subprocess.check_call(["pdftoppm", "-png", "-r", "130", str(pdf), str(work / "slide")])
    pngs = sorted(work.glob("slide-*.png"), key=lambda p: int(p.stem.split("-")[1]))
    for old in out.glob("slide_*.png"):
        old.unlink()
    for i, src in enumerate(pngs, start=1):
        shutil.copy(src, out / f"slide_{i:02d}.png")
    print(f"previews {len(pngs)} -> {out}")


def build() -> Deck:
    charts = build_charts()
    d = Deck()
    d.charts = charts
    slide_cover(d)
    slide_map(d)
    slide_tools(d)
    slide_core(d)
    slide_dealers(d)
    slide_delta(d)
    slide_gamma_def(d)
    slide_gamma_chart(d)
    slide_hedge_chart(d)
    slide_regimes(d)
    slide_positive(d)
    slide_negative(d)
    slide_hedge_sign(d)
    slide_flip(d)
    slide_flip_quiet(d)
    slide_flip_close(d)
    slide_walls(d)
    slide_walls_regime(d)
    slide_opex(d)
    slide_charm(d)
    slide_uw_rules(d)
    slide_heatmap(d)
    slide_bystrike(d)
    slide_oi_volume(d)
    slide_crosscheck(d)
    slide_month(d)
    slide_clocks(d)
    slide_align(d)
    slide_disagree(d)
    slide_spy_numbers(d)
    slide_spy_box(d)
    slide_spy_plan(d)
    slide_pos_play(d)
    slide_pos_candle(d)
    slide_neg_play(d)
    slide_neg_candle(d)
    slide_poke(d)
    slide_rebuild(d)
    slide_alerts(d)
    slide_watch(d)
    slide_bands(d)
    slide_confirm(d)
    slide_mas(d)
    slide_side(d)
    slide_fork(d)
    slide_single(d)
    slide_spread(d)
    slide_payoff(d)
    slide_rich_iv(d)
    slide_place(d)
    slide_delta_bands(d)
    slide_itm(d)
    slide_theta(d)
    slide_crush(d)
    slide_events(d)
    slide_stops(d)
    slide_risk(d)
    slide_flow(d)
    slide_run16(d)
    slide_check(d)
    slide_mistakes(d)
    slide_gloss_a(d)
    slide_gloss_b(d)
    slide_one(d)
    out = ROOT / "gamma_swing.pptx"
    d.save(out)
    finalize_package(out)
    copy_path = Path("/tmp/gamma_copy.txt")
    with copy_path.open("w", encoding="utf-8") as handle:
        current = None
        for rec in d.records:
            if rec["slide"] != current:
                current = rec["slide"]
                title = d.titles[current - 1]
                handle.write(f"\n===== {current:02d}  {title} =====\n")
            handle.write(rec["text"].replace("\n", " / ") + "\n")
    print(f"slides {d.n}")
    for i, title in enumerate(d.titles, start=1):
        print(f"{i:02d}  {title}")
    print(f"warnings {len(d.warnings)}")
    for warning in d.warnings:
        print("WARN", warning)
    print(f"wrote {out}")
    if "--preview" in sys.argv:
        render_previews(out)
    return d


if __name__ == "__main__":
    build()
